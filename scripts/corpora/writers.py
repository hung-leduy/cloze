"""The writers every scenario shares: xlsx, md, docx, pdf, the key, the surfaces.

A scenario module holds data (constants, rows, prose, its gold key); this module
turns that data into files. Scenario A's documents were written by one script
before this split and are pinned byte for byte
(``backend/tests/unit/test_corpus_scenarios.py``), so the writers make exactly
the library calls that script made, in the same order: an xlsx is a
``Workbook()`` whose active sheet is renamed and the rest created, a docx is
``add_heading`` / ``add_paragraph`` / ``add_table(rows=1)`` then ``add_row``.

Why the PDF is hand-written
---------------------------
The repository has ``pypdf`` for reading and no writer, and adding a dependency
to generate test fixtures is a poor trade. For ASCII text a base-14 font is
exactly the case where hand-writing a valid PDF is short and safe. It uses
Helvetica with ``WinAnsiEncoding``, which has no code points for ``ễ`` or ``ị``:
a Vietnamese PDF needs an embedded TrueType subset with a ToUnicode map, or a
converter such as LibreOffice. If these writers are ever pointed at the
deployment corpus, that is the part that breaks first.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
import zipfile
from collections.abc import Iterable, Sequence
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from openpyxl import Workbook

__all__ = [
    "FACT_KEY_NAME", "SURFACES_NAME", "Block", "Sheet", "backend_on_path", "build_key",
    "date_surfaces", "dump_key", "dump_surfaces", "flatten_surfaces", "money_surfaces",
    "normalise_ooxml", "pct", "readback", "write_docx", "write_markdown", "write_pdf",
    "write_xlsx",
]

# The driver (scripts/run_evaluation.py FACT_KEY_NAME) looks for the key under
# this name in every corpus directory, so B, C and D use it too despite the prefix.
FACT_KEY_NAME = "smoke_fact_key.json"

# Every constant string a scenario writes, beside its corpus, so the X4(a)
# generator-surface scan can read it from a run's recorded corpus path without
# importing this package. .json like the key, so no loader turns it into a
# Document (``ingest.SUFFIXES``; the pipeline filters its corpus by suffix).
SURFACES_NAME = "surfaces.json"

# Fixed epoch for the OOXML containers. Any constant does; this one is scenario
# A's own quarter end, so a reader who opens the properties dialog sees a date
# that belongs to the fiction rather than to whenever the file was made. Every
# scenario shares it, because A's bytes depend on it.
_EPOCH = (2026, 9, 30, 0, 0, 0)
_EPOCH_ISO = "2026-09-30T00:00:00Z"
_TIMESTAMP_RE = re.compile(
    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)")

_MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August",
           "September", "October", "November", "December")


def normalise_ooxml(path: Path) -> None:
    """Make an xlsx/docx byte-identical across runs.

    `openpyxl` and `python-docx` write a zip whose entry mtimes are "now", and
    openpyxl stamps `dcterms:created` into `docProps/core.xml`. Neither reaches
    `Document.text`, so the pipeline never notices — but `content_doc_id` is a
    SHA-256 over the *file bytes*, and that id is what the run manifest records.
    Left alone, regenerating the corpus renames every document in every
    manifest that referenced it, and two runs over identical content look like
    runs over different corpora.
    """
    with zipfile.ZipFile(path) as zf:
        entries = [(i, zf.read(i.filename)) for i in zf.infolist()]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename.endswith("core.xml"):
                data = _TIMESTAMP_RE.sub(
                    rb"\g<1>" + _EPOCH_ISO.encode() + rb"\g<2>", data)
            fixed = zipfile.ZipInfo(info.filename, date_time=_EPOCH)
            fixed.compress_type = info.compress_type
            fixed.external_attr = info.external_attr
            out.writestr(fixed, data)
    shutil.move(tmp, path)


# ------------------------------------------------------------------ writers

Sheet = tuple[str, Sequence[Sequence[Any]]]


def write_xlsx(path: Path, sheets: Sequence[Sheet]) -> None:
    """One sheet per (title, rows); the first row of each is its header."""
    wb = Workbook()
    for i, (title, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet(title)
        if i == 0:
            ws.title = title
        for row in rows:
            ws.append(list(row))
    wb.save(path)
    normalise_ooxml(path)


def write_markdown(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


# A docx is a sequence of blocks: ("h", text, level), ("p", text) or
# ("table", header, rows).
Block = tuple[Any, ...]


def write_docx(path: Path, blocks: Sequence[Block]) -> None:
    d = DocxDocument()
    for block in blocks:
        kind = block[0]
        if kind == "h":
            d.add_heading(block[1], level=block[2])
        elif kind == "p":
            d.add_paragraph(block[1])
        elif kind == "table":
            header, rows = block[1], block[2]
            t = d.add_table(rows=1, cols=len(header))
            for cell, text in zip(t.rows[0].cells, header, strict=True):
                cell.text = text
            for values in rows:
                for cell, text in zip(t.add_row().cells, values, strict=True):
                    cell.text = text
        else:
            raise ValueError(f"unknown docx block {kind!r}")
    d.save(path)
    normalise_ooxml(path)


def _esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf(path: Path, pages: list[list[str]], *, size: int = 11,
              leading: int = 15, left: int = 57, top: int = 770) -> None:
    """A minimal, valid PDF: one Helvetica text object per page.

    Object numbering is fixed: 1 catalog, 2 page tree, 3 font, then a content
    stream and a page object per page. `pypdf` reads the result, which is the
    only consumer that matters here.
    """
    contents: list[bytes] = []
    for lines in pages:
        stream = f"BT /F1 {size} Tf {leading} TL 1 0 0 1 {left} {top} Tm\n"
        stream += "".join(f"({_esc(ln)}) Tj T*\n" for ln in lines) + "ET"
        raw = stream.encode("latin-1")
        contents.append(b"<< /Length %d >>\nstream\n" % len(raw) + raw + b"\nendstream")

    kids = [4 + 2 * i + 1 for i in range(len(pages))]
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (b"<< /Type /Pages /Kids [" + b" ".join(b"%d 0 R" % k for k in kids)
         + b"] /Count %d >>" % len(pages)),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
        b"/Encoding /WinAnsiEncoding >>",
    ]
    for i, content in enumerate(contents):
        objects.append(content)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (4 + 2 * i))

    buf = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for n, obj in enumerate(objects, start=1):
        offsets.append(len(buf))
        buf += b"%d 0 obj\n" % n + obj + b"\nendobj\n"
    xref_at = len(buf)
    buf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        buf += b"%010d 00000 n \n" % off
    buf += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(objects) + 1, xref_at))
    path.write_bytes(bytes(buf))


# ----------------------------------------------------------------- surfaces

def money_surfaces(v: float) -> list[str]:
    """Every spelling of one amount the corpus or a report might use.

    Also the key's ``surfaces`` field for a numeric fact, so its output is part
    of scenario A's pinned key and must not change.
    """
    out = {f"{v:,.2f}", f"{v:.2f}"}
    if v == int(v):
        out |= {f"{int(v):,}", str(int(v))}
    else:
        out |= {f"{v:,}", repr(v)}
    return sorted(out)


def date_surfaces(d: date) -> list[str]:
    """A date as an xlsx cell reaches the loader, as ISO, and as prose writes it."""
    return [d.isoformat(), f"{d.isoformat()}T00:00:00",
            f"{d.day} {_MONTHS[d.month - 1]} {d.year}"]


def pct(spent: Decimal, authorised: Decimal) -> Decimal:
    """Share of authority spent, to one decimal place.

    The same rounding as ``fact_store._derive`` (Decimal quantize, half-even),
    so a key value and the R1 store's derived PCTSPENT cannot disagree by
    construction of the arithmetic.
    """
    return (spent / authorised * 100).quantize(Decimal("0.1"))


def flatten_surfaces(groups: dict[str, Iterable[str]]) -> list[str]:
    """Longest first, so a scan that stops at the first hit prefers the whole
    name ("Project Zephyr") to its fragment ("Zephyr")."""
    return sorted({s for v in groups.values() for s in v}, key=lambda s: (-len(s), s))


_SURFACES_NOTE = (
    "Every constant string this scenario writes into a document, by kind ('groups') "
    "and flattened longest first ('surfaces'). Some are short numerals ('340', "
    "'2027') and some are ordinary words (the 'labels' group, some referent "
    "fragments): a scan must match whole tokens and report the 'labels' group "
    "apart, or it will hit text the corpus never touched, such as hex slot ids "
    "and token counts."
)


def dump_surfaces(scenario_id: str, groups: dict[str, list[str]]) -> str:
    """The surfaces file's text, deterministic like the key's."""
    doc = {"scenario": scenario_id, "note": _SURFACES_NOTE,
           "groups": {k: list(v) for k, v in groups.items()},
           "surfaces": flatten_surfaces(groups)}
    return json.dumps(doc, indent=2, sort_keys=True) + "\n"


# ---------------------------------------------------------------------- key

NumericFact = tuple[str, str, float, bool, list[str]]
SurfaceFact = tuple[str, str, str, list[str]]
StringFact = tuple[str, str, str]


def build_key(*, generated_by: str, scenario: str,
              referents: dict[str, list[str]],
              attributes: dict[str, tuple[str, list[str]]],
              numeric_facts: Sequence[NumericFact],
              surface_facts: Sequence[SurfaceFact],
              string_facts: Sequence[StringFact],
              mask_patterns: list[str],
              extra: dict[str, Any] | None = None) -> dict:
    """The gold key as plain data, deterministic across runs.

    The schema is the one ``privdeep.eval.faithfulness.load_key`` reads, and
    for scenario A the output is pinned by sha256; ``extra`` adds top-level
    fields (B, C and D mark themselves provisional) without touching A.
    """
    refs = {k: sorted(set(v), key=lambda s: (-len(s), s)) for k, v in referents.items()}
    facts: list[dict] = []
    for subj, attr, value, derived, also in numeric_facts:
        facts.append({
            "fact_id": f"{subj}.{attr}",
            "subject": subj, "subjects": [subj, *also], "attribute": attr,
            "match": "numeric", "value": value, "derived": derived,
            "surfaces": money_surfaces(float(value)),
        })
    for subj, attr, value, surfaces in surface_facts:
        facts.append({
            "fact_id": f"{subj}.{attr}", "subject": subj, "subjects": [subj],
            "attribute": attr, "match": "surface", "value": value,
            "derived": False, "narrative": True,
            "surfaces": sorted(surfaces, key=lambda s: (-len(s), s)),
        })
    strings = [{"fact_id": f"{s}.{a}", "subject": s, "attribute": a, "value": v}
               for s, a, v in string_facts]
    key = {
        "generated_by": generated_by,
        "scenario": scenario,
        "referents": refs,
        "attributes": {k: {"label": lab, "lexicon": sorted(lex, key=lambda s: (-len(s), s))}
                       for k, (lab, lex) in attributes.items()},
        "facts": facts,
        "string_facts": strings,
        "mask_patterns": mask_patterns,
    }
    key.update(extra or {})
    return key


def dump_key(key: dict) -> str:
    return json.dumps(key, indent=2, sort_keys=True) + "\n"


# ---------------------------------------------------------------- read-back

def backend_on_path() -> None:
    src = str(Path(__file__).resolve().parents[2] / "backend" / "src")
    if src not in sys.path:
        sys.path.insert(0, src)


def readback(out: Path) -> int | None:
    """Load each file back through the real loaders and print what they saw.

    A corpus the pipeline cannot read is then a failure here rather than three
    minutes into a live run. The import is local: the generator is runnable
    from a checkout without the backend on the path, and only this check needs
    it. Returns the total character count, or None when the check was skipped.
    """
    backend_on_path()
    try:
        from privdeep.ingest import SUFFIXES, load
    except ImportError:
        return None
    total = 0
    for p in sorted(out.iterdir()):
        if p.suffix.lower() not in SUFFIXES:
            continue
        doc = load(p)
        total += len(doc.text)
        print(f"  {p.name:<24} {str(doc.kind):<16} {len(doc.text):>6} chars  "
              f"{len(doc.blocks):>3} blocks  {len(doc.cells):>3} cells")
    return total
