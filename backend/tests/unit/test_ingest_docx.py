"""Tests for the .docx loader. Fixtures are built in tmp_path with python-docx.

The property under test is the one the document model rests on: every Cell
and every block span slices the canonical text to exactly the string it
claims, and the text follows the body order of the Word file. The second
half pins the serialisation rules the module docstring states (cell
delimiter replacement, merged cells, content controls), because a silent
change to any of them moves every offset downstream.
"""

from __future__ import annotations

from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from privdeep.ingest.docx import load_docx
from privdeep.models.document import Cell, Document, DocumentKind

PARAS = [
    "Báo cáo quý ba của dự án Zephyr.",
    "Nguyễn Thị Lan nêu rủi ro trong buổi họp tháng sáu tại Hà Nội.",
    "Kết luận: tạm dừng thí điểm.",
]
ROWS = [
    ["Hạng mục", "Q3", "Ghi chú"],
    ["Doanh thu", "1.250.000", "giảm 12%\nso với Q2"],
]
TABLE_TEXT = (
    "## table: 1\n"
    "Hạng mục\tQ3\tGhi chú\n"
    "Doanh thu\t1.250.000\tgiảm 12% so với Q2\n"
    "\n"
)


def build_sample(path: Path) -> Path:
    """Three paragraphs with a 2x3 table between the second and the third."""
    d = docx.Document()
    d.add_paragraph(PARAS[0])
    d.add_paragraph(PARAS[1])
    t = d.add_table(rows=2, cols=3)
    for r, row in enumerate(ROWS):
        for c, value in enumerate(row):
            t.cell(r, c).text = value
    d.add_paragraph(PARAS[2])
    d.save(path)
    return path


def assert_offsets_hold(doc: Document) -> None:
    for cell in doc.cells:
        assert cell.span.slice(doc.text) == cell.value
    for label, span in doc.blocks:
        assert span.end <= len(doc.text), label


@pytest.fixture
def sample(tmp_path: Path) -> Document:
    return load_docx(build_sample(tmp_path / "sample.docx"))


class TestBodyOrder:
    def test_text_is_the_exact_body_serialisation(self, sample: Document):
        assert sample.text == PARAS[0] + "\n" + PARAS[1] + "\n" + TABLE_TEXT + PARAS[2] + "\n"

    def test_table_sits_between_paragraphs_two_and_three(self, sample: Document):
        i_para2 = sample.text.index(PARAS[1])
        i_table = sample.text.index("## table: 1")
        i_para3 = sample.text.index(PARAS[2])
        assert i_para2 < i_table < i_para3

    def test_kind_and_default_doc_id(self, sample: Document):
        assert sample.kind is DocumentKind.DOCX
        assert sample.doc_id == "sample"
        assert sample.source_path.name == "sample.docx"
        assert sample.metadata["paragraphs"] == "3"
        assert sample.metadata["tables"] == "1"
        assert len(sample.metadata["sha256"]) == 64

    def test_explicit_doc_id_is_kept(self, tmp_path: Path):
        doc = load_docx(build_sample(tmp_path / "s.docx"), doc_id="upload-7")
        assert doc.doc_id == "upload-7"


class TestCells:
    def test_every_cell_slices_to_its_value(self, sample: Document):
        assert len(sample.cells) == 6
        assert_offsets_hold(sample)

    def test_cells_carry_one_based_grid_coordinates(self, sample: Document):
        by_pos = {(c.row, c.col): c for c in sample.cells}
        assert set(by_pos) == {(r, c) for r in (1, 2) for c in (1, 2, 3)}
        assert all(c.sheet == "table1" for c in sample.cells)
        assert by_pos[(1, 1)].value == "Hạng mục"
        assert by_pos[(2, 2)].value == "1.250.000"

    def test_newline_in_cell_becomes_one_space(self, sample: Document):
        note = next(c for c in sample.cells if c.row == 2 and c.col == 3)
        assert note.value == "giảm 12% so với Q2"
        assert note.span.slice(sample.text) == note.value

    def test_no_cell_value_contains_a_delimiter(self, sample: Document):
        for cell in sample.cells:
            assert "\t" not in cell.value and "\n" not in cell.value

    def test_tab_in_cell_and_blank_cell(self, tmp_path: Path):
        d = docx.Document()
        t = d.add_table(rows=1, cols=3)
        t.cell(0, 0).text = "a\tb"
        t.cell(0, 2).text = "c"
        d.save(tmp_path / "t.docx")
        doc = load_docx(tmp_path / "t.docx")
        assert doc.text == "## table: 1\na b\t\tc\n\n"
        assert [(c.col, c.value) for c in doc.cells] == [(1, "a b"), (3, "c")]
        assert_offsets_hold(doc)

    def test_merged_cell_is_written_once_and_keeps_grid_width(self, tmp_path: Path):
        d = docx.Document()
        t = d.add_table(rows=2, cols=3)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "wide"
        t.cell(0, 2).text = "c"
        t.cell(1, 0).text = "x"
        t.cell(1, 1).text = "y"
        t.cell(1, 2).merge(t.cell(0, 2))  # vertical merge over column 3
        d.save(tmp_path / "m.docx")
        doc = load_docx(tmp_path / "m.docx")
        assert doc.text == "## table: 1\nwide\t\tc\nx\ty\t\n\n"
        assert [(c.row, c.col, c.value) for c in doc.cells] == [
            (1, 1, "wide"),
            (1, 3, "c"),
            (2, 1, "x"),
            (2, 2, "y"),
        ]
        assert_offsets_hold(doc)


class TestBlocks:
    def test_blocks_are_numbered_body_paragraphs_that_slice_correctly(self, sample: Document):
        assert [label for label, _ in sample.blocks] == ["para 1", "para 2", "para 3"]
        for (_, span), expected in zip(sample.blocks, PARAS, strict=True):
            assert span.slice(sample.text) == expected

    def test_third_block_starts_after_the_table(self, sample: Document):
        _, para3 = sample.blocks[2]
        assert para3.start == sample.text.index(TABLE_TEXT) + len(TABLE_TEXT)

    def test_empty_paragraph_keeps_its_newline_and_number(self, tmp_path: Path):
        d = docx.Document()
        d.add_paragraph("one")
        d.add_paragraph("")
        d.add_paragraph("three")
        d.save(tmp_path / "e.docx")
        doc = load_docx(tmp_path / "e.docx")
        assert doc.text == "one\n\nthree\n"
        assert [(label, sp.slice(doc.text)) for label, sp in doc.blocks] == [
            ("para 1", "one"),
            ("para 2", ""),
            ("para 3", "three"),
        ]

    def test_soft_line_break_stays_inside_the_paragraph(self, tmp_path: Path):
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("line one").add_break(WD_BREAK.LINE)
        p.add_run("line two")
        d.save(tmp_path / "b.docx")
        doc = load_docx(tmp_path / "b.docx")
        assert doc.text == "line one\nline two\n"
        assert doc.blocks[0][1].slice(doc.text) == "line one\nline two"

    def test_content_control_text_is_read_in_place(self, tmp_path: Path):
        d = docx.Document()
        d.add_paragraph("before")
        sdt = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        sdt.append(content)
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "inside control"
        r.append(t)
        p.append(r)
        content.append(p)
        d.element.body.sectPr.addprevious(sdt)
        d.add_paragraph("after")
        d.save(tmp_path / "c.docx")
        doc = load_docx(tmp_path / "c.docx")
        assert doc.text == "before\ninside control\nafter\n"
        assert [label for label, _ in doc.blocks] == ["para 1", "para 2", "para 3"]
        assert doc.blocks[1][1].slice(doc.text) == "inside control"


class TestDeterminism:
    def test_same_file_loads_identically(self, tmp_path: Path):
        path = build_sample(tmp_path / "d.docx")
        a, b = load_docx(path), load_docx(path)
        assert (a.text, a.cells, a.blocks, a.doc_id, a.metadata) == (
            b.text, b.cells, b.blocks, b.doc_id, b.metadata,
        )

    def test_same_content_in_a_fresh_build_yields_the_same_text(self, tmp_path: Path):
        a = load_docx(build_sample(tmp_path / "first.docx"))
        b = load_docx(build_sample(tmp_path / "second.docx"))
        assert a.text == b.text
        assert a.cells == b.cells
        assert a.blocks == b.blocks
        assert all(isinstance(c, Cell) for c in a.cells)
