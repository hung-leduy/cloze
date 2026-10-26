"""scripts/paper_format.py: the paper's format (the ICCIT instruction on the IEEE
conference template's two columns), and the fill pass of scripts/make_paper_tables.py
putting tables and figures in it."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from scripts import make_paper_tables as M
from scripts import paper_format as F

try:
    import docx
except ImportError:          # the paper's tools run under the system interpreter
    docx = None

needs_docx = pytest.mark.skipif(docx is None, reason="python-docx is not installed")
_REPO = Path(__file__).resolve().parents[3]
VALUES = {"alpha": "0.41", "count": "41"}


def test_numerals_and_letters():
    assert [F.roman(n) for n in (1, 4, 6, 8, 9, 14)] == ["I", "IV", "VI", "VIII", "IX", "XIV"]
    assert F.letter(1) == "A" and F.letter(3) == "C"
    assert F.section_label("4") == "IV" and F.section_label("6.2") == "VI-B"
    assert F.table_label("Table 3") == "Table III" and F.figure_label("Fig. 2") == "Figure 2"
    with pytest.raises(ValueError):
        F.section_label("4.2.1")
    with pytest.raises(ValueError):
        F.roman(0)


@pytest.mark.parametrize("text, expected", [
    ("as in §3 and (§6.2).", "as in Section III and (Section VI-B)."),
    ("see §1, §4.2, §5, §7 and §8 for", "see Sections I, IV-B, V, VII and VIII for"),
    ("Section 3 bounds the claim", "Section III bounds the claim"),
    ("Tables 2 and 3", "Tables II and III"),
    ("in Tables 2-3", "in Tables II-III"),
    ("(Table 3; Fig. 2)", "(Table III; Fig. 2)"),
    ("Table 3, and Table 1", "Table III, and Table I"),
    ("P1-P9, a 5-word run, «t4_A_r1_rba»", "P1-P9, a 5-word run, «t4_A_r1_rba»"),
])
def test_drafts_numbering_is_written_as_the_template_numbers(text, expected):
    assert F.ieee_labels(text) == expected


def _ieee_doc(paragraphs: list[str]):
    d = docx.Document()
    for el in list(d.element.body):
        if el.tag != F._qn("w:sectPr"):
            d.element.body.remove(el)
    F.apply(d)
    for t in paragraphs:
        d.add_paragraph(t, style=F.BODY)
    return d


def _cols(d) -> list[str]:
    return [s._sectPr.find(F._qn("w:cols")).get(F._qn("w:num")) for s in d.sections]


def _out(tmp_path: Path, values: dict[str, str]) -> Path:
    out = tmp_path / "paper"
    out.mkdir()
    (out / "numbers.json").write_text(json.dumps(values), encoding="utf-8")
    return out


@needs_docx
def test_apply_sets_the_templates_page_columns_and_styles():
    d = _ieee_doc([])
    sec = d.sections[-1]
    assert (sec.page_width, sec.page_height) == (docx.shared.Inches(8.5), docx.shared.Inches(11))
    # The ICCIT instruction's "at least one inch (3 cm) on all sides" (2026-10-25).
    assert sec.left_margin == sec.right_margin == docx.shared.Inches(1)
    assert sec.top_margin == sec.bottom_margin == docx.shared.Inches(1)
    assert F.TEXT_WIDTH_IN == 6.5 and 2 * F.COLUMN_WIDTH_IN + F.COLUMN_GAP / 1440 == 6.5
    assert _cols(d) == ["2"]
    assert sec._sectPr.find(F._qn("w:footerReference")) is None      # paginate() adds it
    sizes = {name: d.styles[name].font.size.pt for name in
             (F.BODY, F.TITLE, F.AUTHOR, F.ABSTRACT, F.KEYWORDS, F.TABLE_CAPTION,
              F.TABLE_NOTE, F.FIGURE_CAPTION, F.REFERENCE, "Heading 1", "Heading 2")}
    # The abstract and keywords at 10 pt and nothing below 8 pt (the author's choice of
    # 2026-10-25).
    assert sizes == {F.BODY: 10, F.TITLE: 14, F.AUTHOR: 11, F.ABSTRACT: 10, F.KEYWORDS: 10,
                     F.TABLE_CAPTION: 8, F.TABLE_NOTE: 8, F.FIGURE_CAPTION: 8,
                     F.REFERENCE: 8, "Heading 1": 10, "Heading 2": 10}
    # Times New Roman everywhere, with no theme font left to override it.
    rfonts = d.styles["Heading 1"].element.rPr.find(F._qn("w:rFonts"))
    assert rfonts.get(F._qn("w:ascii")) == "Times New Roman"
    assert not [a for a in rfonts.attrib if a.lower().endswith("theme")]


@needs_docx
def test_paginate_puts_the_running_head_and_page_numbers_on_every_page(tmp_path):
    """The ICCIT instruction's numbered pages and running head: a default header and
    footer on the first section, none of their own on the later ones (they inherit
    them), the template's unused footer dropped, and a second call a no-op."""
    d = docx.Document(str(_REPO / "paper" / "ICCIT2026_template.docx"))
    for el in list(d.element.body):
        if el.tag != F._qn("w:sectPr"):
            d.element.body.remove(el)
    F.apply(d)
    p = d.add_paragraph("Title", style=F.TITLE)
    F.end_section(d, p._p, 1, start=None)
    d.add_paragraph("Body.", style=F.BODY)
    F.paginate(d, F.RUNNING_HEAD)
    xml = (d.sections[0].header._element.xml, d.sections[0].footer._element.xml)
    F.paginate(d, F.RUNNING_HEAD)
    assert (d.sections[0].header._element.xml, d.sections[0].footer._element.xml) == xml
    head = d.sections[0].header.paragraphs
    assert [h.text for h in head] == [F.RUNNING_HEAD]
    assert head[0].alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    foot = d.sections[0].footer._element
    assert [t.text for t in foot.iter(F._qn("w:instrText"))] == [" PAGE "]
    assert not [t for t in foot.iter(F._qn("w:t")) if t.text]      # no numeral cached
    assert not d.sections[1]._sectPr.findall(F._qn("w:headerReference"))
    assert not d.sections[1]._sectPr.findall(F._qn("w:footerReference"))
    targets = sorted(r.target_ref for r in d.part.rels.values()
                     if r.reltype.endswith(("/header", "/footer")))
    assert targets == ["footer1.xml", "header1.xml"]
    out = tmp_path / "p.docx"
    d.save(str(out))
    assert docx.Document(str(out)).sections[1].header.is_linked_to_previous
    with pytest.raises(ValueError, match="50 characters"):
        F.paginate(d, "x" * 51)


def test_the_running_head_is_one_text_in_every_builder():
    from paper.tools import assemble_manuscript as A
    assert A.RUNNING_HEAD == F.RUNNING_HEAD and len(F.RUNNING_HEAD) <= 50
    skeleton = (_REPO / "paper" / "tools" / "make_skeleton.py").read_text(encoding="utf-8")
    assert f'RUNNING_HEAD = "{F.RUNNING_HEAD}"' in skeleton


@needs_docx
def test_a_table_spans_both_columns_with_the_templates_caption_and_notes(tmp_path):
    d = _ieee_doc(["Before.", "[TABLE 1 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["t1"], "After."])
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    out = _out(tmp_path, VALUES)
    spec = M.TableSpec("t1", "Table 1", "Faithfulness over «count» runs (§6.3).", ["A", "B"],
                       [["r1", "«alpha»"]], notes=["As in Table 1."], widths=[1.0, 2.0])
    M.write_tables_docx(src, [spec], out / "tables.docx", VALUES)
    M.fill(src, dst, out)
    f = docx.Document(str(dst))
    got = [(p.style.name, p.text.replace('\xa0', ' ')) for p in f.paragraphs]
    # Sentence case since 2026-10-19 (the author's choice), not the template's capitals.
    assert (F.TABLE_CAPTION, "Table I. Faithfulness over 41 runs (Section VI-C).") in got
    assert (F.TABLE_NOTE, "As in Table I.") in got
    widths = [int(g.get(F._qn("w:w"))) for g in f.tables[0]._tbl.tblGrid.iter(
        F._qn("w:gridCol"))]
    assert abs(sum(widths) - F.TEXT_WIDTH_IN * 1440) <= 2
    # Two columns, the table in one, two columns again.
    assert _cols(f) == ["2", "1", "2"]
    assert f.sections[1].start_type == docx.enum.section.WD_SECTION.CONTINUOUS


@needs_docx
def test_adjacent_blocks_share_one_single_column_section(tmp_path):
    d = _ieee_doc(["Before.", "[TABLE 1 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["t1"],
                   "[FIGURE 2 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["fig2"], "After."])
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    out = _out(tmp_path, VALUES)
    specs = [M.TableSpec("t1", "Table 1", "One.", ["A"], [["x"]], notes=["n1"])]
    M.write_tables_docx(src, specs, out / "tables.docx", VALUES)
    M.fill(src, dst, out)
    f = docx.Document(str(dst))
    assert _cols(f) == ["2", "1", "2"]
    assert [p.text for p in f.paragraphs][-1] == "After."


@needs_docx
def test_figure_one_gets_the_templates_caption_at_single_spacing(tmp_path):
    d = _ieee_doc(["Before.", "[FIGURE 1 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["fig1"],
                   "After."])
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    out = _out(tmp_path, VALUES)
    M.fill(src, dst, out)
    f = docx.Document(str(dst))
    cap = next(p for p in f.paragraphs if p.style.name == F.FIGURE_CAPTION)
    assert cap.text.startswith("Figure 1. ") and "§" not in cap.text
    # A title since 2026-10-19 (the author's request for short captions).
    assert cap.text == "Figure 1. Cloze on an illustrative budget row."
    pic = cap._p.getprevious()
    assert pic.xpath(".//pic:pic") and docx.text.paragraph.Paragraph(
        pic, None).paragraph_format.line_spacing == 1.0
    assert _cols(f) == ["2", "1", "2"]
