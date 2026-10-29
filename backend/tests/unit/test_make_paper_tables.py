"""PAP5a: numbers formatting, the placeholder fill pass, and the numeral checker.

The build itself runs under /usr/bin/python3 (python-docx 1.2.0, matplotlib
3.10.6); the project venv has python-docx 1.1.2 and no matplotlib. The fill and
checker tests need only python-docx and run in either interpreter; the figure
test is skipped cleanly where matplotlib is missing.

Every value here is synthetic and chosen to match no pilot output; key names
are neutral where the name adds nothing.
"""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

import pytest

from scripts import check_paper_numbers as C
from scripts import make_paper_tables as M

_REPO = Path(__file__).resolve().parents[3]


# ------------------------------------------------------------ formatting

def test_fmt_stat_summarises_runs_and_says_na_for_none():
    """Since the evaluation rewrite of 2026-10-15 a cell is mean [min–max], without
    the SD; fmt_stat_sd keeps the SD for the one key §6 quotes as mean ± SD."""
    three = {"n": 3, "mean": 0.41234, "stdev": 0.03111, "min": 0.3822, "max": 0.4431}
    assert M.fmt_stat(three) == "0.41 [0.38–0.44]"
    assert M.fmt_stat({"n": 1, "mean": 0.625, "stdev": None, "min": 0.625,
                       "max": 0.625}) == "0.62"
    assert M.fmt_stat({"n": 0, "mean": None, "stdev": None}) == "n/a"
    assert M.fmt_stat(None) == "n/a"
    assert M.fmt_stat(three, 1) == "0.4 [0.4–0.4]"
    assert M.fmt_stat_sd(three) == "0.41 ± 0.03 [0.38–0.44]"
    assert M.fmt_stat_sd({"n": 1, "mean": 0.625, "stdev": None}) == "0.62"
    assert M.fmt_stat_sd(None) == "n/a"


def test_a_cell_or_a_mean_over_fewer_runs_says_so():
    two = {"n": 2, "mean": 0.63, "stdev": 0.01, "min": 0.62, "max": 0.64}
    assert M._fmt_cell(two, 3) == "0.63 [0.62–0.64] (2 of 3 runs)"
    assert M._fmt_cell(two, 2) == "0.63 [0.62–0.64]"
    assert M._fmt_cell(None, 3) == "n/a"
    nb = M.Numbers()
    M._put_mean(nb, "alpha", two, 3, "e.json#/x")
    assert nb["alpha"] == "0.63" and nb["alpha_n"] == "2/3"
    assert nb.thin == {"alpha": (2, 3)}


def test_small_formats():
    assert M.fmt_fixed(-0.0001, 2) == "0.00"          # no negative zero
    assert M.fmt_fixed(-0.25, 2) == "-0.25"
    assert M.fmt_pct(0.12345) == "12.3%"
    assert M.fmt_pct(None) == "n/a"
    assert M.fmt_kn(7, 41) == "7/41"
    assert M.fmt_kn(None, 3) == "n/a"
    assert M.fmt_ci([0.1234, 0.3456]) == "[0.123, 0.346]"
    assert M.fmt_ci(None) == "n/a"
    assert M.fmt_num(None) == "n/a"
    assert M.fmt_range(["12.3%", "12.3%"]) == "12.3%"
    assert M.fmt_range(["20.1%", "9.0%"]) == "9.0%–20.1%"
    # A range that silently dropped a corpus would cover less than it says.
    assert M.fmt_range(["20.1%", "9.0%", "n/a"]) == "n/a"
    assert M.fmt_range([]) == "n/a"


def test_numbers_keep_one_value_per_key_with_its_source():
    nb = M.Numbers()
    nb.put("alpha", "0.41", "evaluation.json#/by_task/t", 0.412, "2 dp")
    nb.put("alpha", "0.41", "evaluation.json#/by_task/t", 0.412, "2 dp")
    assert nb.sources["alpha"]["source"].startswith("evaluation.json#/")
    with pytest.raises(ValueError, match="twice"):
        nb.put("alpha", "0.42", "elsewhere")
    with pytest.raises(ValueError, match="bad key"):
        nb.put("not a key", "1", "x")
    with pytest.raises(ValueError, match="guillemet"):
        nb.put("beta", "«alpha»", "x")


def test_json_pointer_escapes_task_ids_and_arm_names():
    src = M.Src("e.json", {"hypotheses": {"H4-D'": {"outcome": "holds"}},
                           "a/b": [10, {"c": 3}]})
    assert src.get("hypotheses", "H4-D'", "outcome") == (
        "holds", "e.json#/hypotheses/H4-D'/outcome")
    assert src.get("a/b", 1, "c") == (3, "e.json#/a~1b/1/c")
    assert src.get("missing", "x")[0] is None


def test_inputs_without_the_required_blocks_are_refused():
    with pytest.raises(SystemExit, match="hypotheses"):
        M.numbers_from_evaluation(M.Numbers(), M.Src("evaluation.json", {"arms": {}}))
    with pytest.raises(SystemExit, match="deployed_chain"):
        M.numbers_from_fp(M.Numbers(), M.Src("fp_report.json", {"linters": {}}))


# ------------------------------------------------------------ the aggregate

def _faith(sha: str, oracle: bool = True) -> dict:
    f = {"key_sha256": sha, "private": {"rba": 0.41}}
    if oracle:
        f["oracle"] = {"rba": 0.63}
    return f


def _evaluation(**over) -> dict:
    ev = {
        "hypotheses": {"corpora": {"A": "task-a"}, "registered_repeats": 3},
        "by_task": {"task-a": {"key_sha256": "k1", "stale_key_runs": [], "arms": {}}},
        "per_task": [{"task_id": "task-a", "arm": "r1", "run": "task-a__r1",
                      "faithfulness": _faith("k1")}],
    }
    ev.update(over)
    return ev


def test_a_run_scored_against_a_stale_key_stops_the_build():
    ev = _evaluation()
    ev["by_task"]["task-a"]["stale_key_runs"] = ["task-a__r1__rep2"]
    with pytest.raises(SystemExit, match=r"task-a \(private\): task-a__r1__rep2"):
        M.numbers_from_evaluation(M.Numbers(), M.Src("evaluation.json", ev))


def test_a_stale_key_on_the_oracle_side_stops_the_build_too():
    # The aggregate lists stale runs for the private side only; an oracle-only
    # scoring against another key is found in per_task.
    ev = _evaluation(per_task=[
        {"task_id": "task-a", "arm": "baseline", "run": "task-a__baseline",
         "faithfulness": {"key_sha256": "k0", "private": None, "oracle": {"rba": 0.63}}}])
    assert M.stale_key_runs(M.Src("e", ev)) == {"task-a (oracle)": ["task-a__baseline"]}
    with pytest.raises(SystemExit, match="rescore"):
        M.numbers_from_evaluation(M.Numbers(), M.Src("evaluation.json", ev))
    assert M.stale_key_runs(M.Src("e", _evaluation())) == {}


def _crossings() -> dict:
    return {"crossings": {
        "a_generator_surfaces": {
            "runs_scanned": ["t__baseline", "t__r1"],
            "runs_not_scanned": {"t__r1__rep2": "no egress_trace.jsonl"},
            "hits": [{"run": "t__r1", "arm": "r1", "count": 2}],
            "label_hits": [],
            "by_arm": {"baseline": {"runs": 1, "hits": 0, "label_hits": 0},
                       "r1": {"runs": 1, "hits": 2, "label_hits": 0}}},
        "b_extracted_entities": {
            "runs_scanned": ["t__baseline"],
            "runs_not_scanned": {"t__r1": "no entities.json or extractions.json"},
            "hits": [],
            "by_arm": {"baseline": {"hits": 0, "core_sent": 0, "cloud_authored": 0},
                       "r1": {"hits": 0, "core_sent": 0, "cloud_authored": 0}}},
        "c_reconstruction": {"hits": [], "distinct_entities": []}},
        "per_task": [], "reconstruction": {"by_arm": {}}}


def test_a_count_cell_prints_whole_minimum_and_maximum_without_decimals():
    # 2026-10-28 (critical review PRE10): '28.7 [20.0–37.0]' beside '17.2 [6–37]'.
    stat = {"n": 3, "mean": 28.7, "stdev": 8.5, "min": 20.0, "max": 37.0}
    assert M._fmt_count_cell(stat, 3) == "28.7 [20–37]"
    assert M._fmt_count_cell({**stat, "min": 20.5}, 3) == "28.7 [20.5–37.0]"
    assert M._fmt_count_cell({**stat, "n": 2}, 3) == "28.7 [20–37] (2 of 3 runs)"
    assert M._fmt_count_cell(None, 3) == "n/a"


def test_h3s_exclusions_are_counted_by_kind_and_n_a_when_the_block_has_none():
    # 2026-10-28 (critical review ANA6): the spans H3's scan excluded are counted, goal
    # and cloud-reply repeats apart, and how many of the latter are label words.
    ev = _crossings()
    ev["crossings"]["a_generator_surfaces"]["excluded"] = {
        "user-goal": [], "cloud-echo": [{"label_group": True}, {"label_group": False}]}
    nb = M.Numbers()
    M.numbers_from_crossings(nb, M.Src("evaluation.json", ev))
    assert (nb["h3_excluded_goal"], nb["h3_excluded_cloud_echo"],
            nb["h3_excluded_cloud_echo_label"]) == ("0", "2", "1")
    nb = M.Numbers()
    M.numbers_from_crossings(nb, M.Src("evaluation.json", _crossings()))
    assert nb["h3_excluded_goal"] == nb["h3_excluded_cloud_echo"] == "n/a"


def test_a_crossing_count_is_na_for_a_missing_arm_and_carries_its_scanned_runs():
    nb = M.Numbers()
    M.numbers_from_crossings(nb, M.Src("evaluation.json", _crossings()))
    # r5 has no entry: n/a, never a clean-looking 0.
    assert nb["cross_a_r5"] == "n/a" and nb["cross_b_core_r5"] == "n/a"
    assert nb["cross_a_scanned_r5"] == "n/a"
    assert nb["cross_a_r1"] == "2" and nb["cross_a_baseline"] == "0"
    # r1: one run scanned of two (rep2 had no trace); (b) scanned none of r1's.
    assert nb["cross_a_scanned_r1"] == "1/2" and nb["cross_a_scanned_baseline"] == "1/1"
    assert nb["cross_b_scanned_r1"] == "0/1"
    assert nb["cross_a_runs_not_scanned"] == "1" and nb["cross_b_runs_not_scanned"] == "1"
    assert "no egress_trace.jsonl" in nb["cross_a_not_scanned_reasons"]
    # Units: records at the top, occurrences per arm.
    assert nb["cross_a_hits"] == "1" and nb["cross_a_occurrences"] == "2"
    assert "occurrences" in nb.sources["cross_a_r1"]["format"]
    spec = M.table3(nb)
    arm_rows = [r for r in spec.rows if isinstance(r, list) and r[0] in M.ARM_ORDER]
    col = M.T3_ARM_HEAD.index("Generator-written constants and label words in core-sent "
                              "content (H3)")
    # H3 counts label surfaces too, so the H3 column gives both (claim-preservation
    # review of 2026-10-19: a column tagged H3 must not leave them out).
    assert {r[0]: r[col] for r in arm_rows} == {
        a: f"«cross_a_{a}» and «cross_a_label_{a}» in «cross_a_scanned_{a}» runs"
        for a in M.ARM_ORDER}
    # Core-sent and cloud-written entity strings in two columns since the readability
    # review of 2026-10-19 ("0 / 3" read as k/N beside the adversary's k/N).
    col = M.T3_ARM_HEAD.index("Extracted-entity strings in core-sent content")
    assert {r[0]: (r[col], r[col + 1]) for r in arm_rows} == {
        a: (f"«cross_b_core_{a}» in «cross_b_scanned_{a}» runs", f"«cross_b_cloud_{a}»")
        for a in M.ARM_ORDER}
    col = M.T3_ARM_HEAD.index("Adversary matches: extracted entities")
    assert {r[0]: (r[col], r[col + 1]) for r in arm_rows} == {
        a: (f"«g3_entities_{a}»", f"«g3_5grams_{a}»") for a in M.ARM_ORDER}


def _comp(outcome, lo, hi, short=(), reason=None, **extra):
    return {"test": "x", "outcome": outcome, "reason": reason,
            "short_of_registered_repeats": list(short),
            "lower": {"label": "lo", "runs": len(lo), "max": max(lo) if lo else None},
            "upper": {"label": "hi", "runs": len(hi), "min": min(hi) if hi else None},
            **extra}


def _hypotheses() -> dict:
    two = ["upper side: 2 pooled run(s)"]
    return {"hypotheses": {
        "registered_repeats": 3,
        "short_of_registered_repeats": {"r1 on A": 2},
        "H1": {"per_corpus": {"A": {"outcome": "holds", "components": {
            "rba": _comp("holds", [0.21, 0.24], [0.61, 0.63], two),
            "attribute_accuracy": _comp("holds", [0.31], [0.52, 0.55], two)}}},
               "holds_on": ["A"]},
        "H2": {"outcome": "fails", "components": {
            "coverage_stated": _comp("fails", [0.0], [0.0], floor=True)}},
        "H4-D": {"outcome": "not testable", "components": {
            "rba_gain_vs_B": _comp("not testable", [0.11], [], reason="no pooled r1 run on B")}},
        "H4-D'": {"outcome": "holds", "components": {"vs_A": _comp("holds", [0.1], [0.3])}},
        "H3": {"outcome": "not testable", "reason": "no r5 run on B", "hits": 0,
               "label_group_hits": 0, "runs_scanned": 4, "runs_not_scanned": {},
               "cells_uncovered": ["no r5 run on B"], "attempts_not_scanned": []}}}


def test_hypothesis_outcomes_carry_their_caveats():
    nb = M.Numbers()
    M._numbers_from_hypotheses(nb, M.Src("evaluation.json", _hypotheses()))
    assert nb["h1_A_outcome"] == "holds"
    assert nb["h1_A_short_of_repeats"] == "upper side: 2 pooled run(s)"
    assert nb["h1_A_rba_r1_min"] == "0.61" and nb["h1_A_rba_baseline_max"] == "0.24"
    assert nb["h1_A_rba_r1_runs"] == "2"
    assert nb["h_short_of_registered_repeats"] == "r1 on A: 2 of 3"
    # H2 is named by side, r1 below baseline, and its floor is exported.
    assert nb["h2_coverage_stated_r1_max"] == "0.00"
    assert nb["h2_coverage_stated_baseline_min"] == "0.00"
    assert nb["h2_coverage_stated_floor"] == "yes"
    assert nb["h4d_rba_gain_vs_B_D_max"] == "0.11" and nb["h4d_rba_gain_vs_B_B_min"] == "n/a"
    assert nb["h4d_reason"] == "no pooled r1 run on B"
    assert nb["h4dp_vs_A_A_min"] == "0.30" and nb["h4dp_short_of_repeats"] == "none"
    assert nb["h3_reason"] == "no r5 run on B" and nb["h3_runs_not_scanned"] == "0"
    # 'holds' over fewer runs than registered: a warning, or a refusal when the
    # skeleton quotes it.
    assert set(M.short_holds(nb)) == {"h1_A_outcome", "h1_A_rba_outcome",
                                      "h1_A_attr_outcome"}
    M._check_short_holds(nb, ["h4dp_outcome"], allow=False)
    with pytest.raises(SystemExit, match="h1_A_outcome"):
        M._check_short_holds(nb, ["h1_A_outcome"], allow=False)
    M._check_short_holds(nb, ["h1_A_outcome"], allow=True)


def test_a_reason_names_a_clause_shared_by_two_components_once():
    # H4-D's two RBA gains both compare against D's runs, so the aggregate gives
    # each component the same clause for D; the joined reason prints it once.
    d = "D side, no value for d__baseline: rba is null (empty pool)"
    h = _hypotheses()
    h["hypotheses"]["H4-D"]["components"] = {
        "rba_gain_vs_A": _comp("not testable", [], [0.2], reason=f"{d}; A side, no value "
                                                                 "for a__baseline: x"),
        "rba_gain_vs_B": _comp("not testable", [], [0.2], reason=d)}
    nb = M.Numbers()
    M._numbers_from_hypotheses(nb, M.Src("evaluation.json", h))
    assert nb["h4d_reason"] == f"{d}; A side, no value for a__baseline: x"


def _probe(cid: str) -> dict:
    ext = {"masked_exposure": {"masked_5gram_hit_share": 0.123,
                               "corpus_5grams_surviving_share": 0.2,
                               "generator_surfaces_surviving": {"distinct": 7,
                                                                "extractor_misses": []},
                               "refused_documents": 1, "documents": ["d1", "d2"]},
           "predicates": {"echo_share": 0.1, "whole_match_share": 0.0, "max_tokens": 9,
                          "n_predicates": 11},
           "r5_residue": {"referents_with_residue": 0, "referents_total": 4},
           "recall": {"referents": {"exact": 3, "n": 4}}}
    return {"corpus": {"id": cid}, "per_extraction": [ext],
            "predicate_control": {"expected_echo_share": 0.0},
            "variance": {"entity_jaccard": 0.87, "relation_jaccard": 0.83}, "g1": {}}


def test_pooled_masked_exposure_is_na_unless_every_measured_corpus_has_a_probe():
    corpora = {"A": "task-a", "C": "task-c"}
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", _probe("task-a"))], corpora, ["A", "C"])
    assert nb["maskdoc_5gram_pct_A"] == "12.3%"
    assert nb["maskdoc_5gram_pct"] == "n/a" and nb["maskdoc_surfaces_surviving"] == "n/a"
    assert nb["maskdoc_corpora"] == "A" and nb["maskdoc_corpora_missing"] == "C"
    assert "no probe for C" in nb.sources["maskdoc_5gram_pct"]["source"]
    assert nb["extractions_distinct_A"] == "n/a"          # absent, not 'None'
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", _probe("task-a"))], corpora, ["A"])
    assert nb["maskdoc_5gram_pct"] == "12.3%" and nb["maskdoc_corpora_missing"] == "none"


# ------------------------------------------------------------ PAP5b follow-ups

def _leak_eval(**over) -> dict:
    """Task a: r5 extracted, baseline and r1 (two repeats) reuse it. Task c: one
    extraction whose extracting run was not measured, read from a reuse run."""
    src_a = "/data/runs/task-a__r5/extractions.json"
    src_c = "/data/runs/task-c__r5/extractions.json"
    rows = [("task-a", "r5", "task-a__r5", "extracted"),
            ("task-a", "baseline", "task-a__baseline", src_a),
            ("task-a", "r1", "task-a__r1", src_a),
            ("task-a", "r1", "task-a__r1__rep2", src_a),
            ("task-c", "r1", "task-c__r1", src_c),
            ("task-c", "baseline", "task-c__baseline", src_c)]
    ev = {"hypotheses": {"corpora": {"A": "task-a", "C": "task-c"}},
          "by_task": {"task-a": {"arms": {"r5": {}}}, "task-c": {"arms": {"r1": {}}}},
          "per_task": [{"task_id": t, "arm": a, "run": r, "extraction_source": s}
                       for t, a, r, s in rows],
          "health_gate": {"outcomes": {
              r: {"relations_dropped_leak": 4 if t == "task-a" else 2,
                  "documents_all_relations_dropped": 1 if t == "task-a" else 0}
              for t, _, r, _ in rows}}}
    ev.update(over)
    return ev


def test_leak_gate_drops_are_read_once_per_extraction():
    nb = M.Numbers()
    M.numbers_from_leak_gate(nb, M.Src("evaluation.json", _leak_eval()))
    # Four runs mask task-a's one extraction; its 4 drops are counted once, not 16.
    assert nb["leak_gate_drops_A"] == "4" and nb["leak_gate_docs_emptied_A"] == "1"
    assert nb["leak_gate_extractions_A"] == "1"
    assert "task-a__r5 via task-a__r5" in nb.sources["leak_gate_drops_A"]["source"]
    # The extracting run of C was not measured: a reuse run carries its counters.
    assert nb["leak_gate_drops_C"] == "2"
    assert "task-c__r5 via task-c__baseline" in nb.sources["leak_gate_drops_C"]["source"]
    assert nb["leak_gate_drops_total"] == "6" and nb["leak_gate_docs_emptied_total"] == "1"


def test_an_extracting_run_older_than_the_counter_is_read_from_the_file_counters(capsys):
    """The extracting run's manifest predates the counter (None); a run that
    reused the file loaded its counters (0 for a file of that date, whose
    extractor aborted instead of dropping), so the extraction is read there."""
    ev = _leak_eval()
    out = ev["health_gate"]["outcomes"]
    out["task-a__r5"].update(relations_dropped_leak=None, documents_all_relations_dropped=None)
    for r in ("task-a__baseline", "task-a__r1", "task-a__r1__rep2"):
        out[r].update(relations_dropped_leak=0, documents_all_relations_dropped=0)
    nb = M.Numbers()
    M.numbers_from_leak_gate(nb, M.Src("evaluation.json", ev))
    assert nb["leak_gate_drops_A"] == "0"
    assert "task-a__r5 via task-a__baseline" in nb.sources["leak_gate_drops_A"]["source"]
    assert "disagree" not in capsys.readouterr().err       # None is not a disagreement
    out["task-a__r1"].update(relations_dropped_leak=3)
    M.numbers_from_leak_gate(M.Numbers(), M.Src("evaluation.json", ev))
    assert "disagree" in capsys.readouterr().err


def test_leak_gate_drops_are_na_when_an_extraction_or_a_counter_is_unknown():
    ev = _leak_eval()
    ev["per_task"][4]["extraction_source"] = "unrecorded"
    ev["per_task"][5]["extraction_source"] = "unrecorded"
    for m in ev["per_task"][:4]:                      # every manifest predates the counter
        ev["health_gate"]["outcomes"][m["run"]]["relations_dropped_leak"] = None
    nb = M.Numbers()
    M.numbers_from_leak_gate(nb, M.Src("evaluation.json", ev))
    assert nb["leak_gate_drops_A"] == "n/a"
    assert nb["leak_gate_drops_C"] == "n/a"          # no run says which extraction it masked
    assert "not recorded: task-c__r1, task-c__baseline" in \
        nb.sources["leak_gate_drops_C"]["source"]
    assert nb["leak_gate_drops_total"] == "n/a"
    # The outcome's own extraction_source is used when the measurement lacks one.
    ev = _leak_eval()
    for m in ev["per_task"]:
        ev["health_gate"]["outcomes"][m["run"]]["extraction_source"] = m.pop(
            "extraction_source")
    nb = M.Numbers()
    M.numbers_from_leak_gate(nb, M.Src("evaluation.json", ev))
    assert nb["leak_gate_drops_total"] == "6"


def test_a_reuse_of_a_reuse_is_counted_once_at_the_run_that_extracted():
    """Every run writes its own extractions.json, so an arm may reuse a run that
    itself reused. r5 extracted (3 drops), baseline reused r5, r1 reused
    baseline: one extraction, read once, not two."""
    rows = [("task-a", "r5", "task-a__r5", "extracted"),
            ("task-a", "baseline", "task-a__baseline", "/data/runs/task-a__r5/extractions.json"),
            ("task-a", "r1", "task-a__r1", "/data/runs/task-a__baseline/extractions.json")]
    ev = {"hypotheses": {"corpora": {"A": "task-a"}},
          "by_task": {"task-a": {"arms": {"r5": {}}}},
          "per_task": [{"task_id": t, "arm": a, "run": r, "extraction_source": s}
                       for t, a, r, s in rows],
          "health_gate": {"outcomes": {r: {"relations_dropped_leak": 3,
                                           "documents_all_relations_dropped": 1}
                                       for _, _, r, _ in rows}}}
    nb = M.Numbers()
    M.numbers_from_leak_gate(nb, M.Src("evaluation.json", ev))
    assert nb["leak_gate_extractions_A"] == "1"
    assert nb["leak_gate_drops_A"] == "3" and nb["leak_gate_docs_emptied_A"] == "1"
    assert nb.sources["leak_gate_drops_A"]["raw"] == {"task-a__r5": {
        "read_from": "task-a__r5", "relations_dropped_leak": 3,
        "documents_all_relations_dropped": 1}}
    # The chain is followed through the outcomes too, when the rows lack it.
    for m in ev["per_task"]:
        ev["health_gate"]["outcomes"][m["run"]]["extraction_source"] = m.pop(
            "extraction_source")
    nb = M.Numbers()
    M.numbers_from_leak_gate(nb, M.Src("evaluation.json", ev))
    assert nb["leak_gate_extractions_A"] == "1" and nb["leak_gate_drops_A"] == "3"
    # A named source whose own source is unknown ends the chain there, and a
    # cycle does not loop.
    assert M._root_extraction("x", {"x": "/r/y/extractions.json"}.get) == "y"
    cyc = {"x": "/r/y/extractions.json", "y": "/r/x/extractions.json"}
    assert M._root_extraction("x", cyc.get) in ("x", "y")


def test_counters_read_from_a_reuse_run_say_how_an_old_file_loads():
    """extraction_io: a file written before the counter existed loads it as 0,
    which is true of it; the sidecar says so wherever a reuse run is read."""
    nb = M.Numbers()
    M.numbers_from_leak_gate(nb, M.Src("evaluation.json", _leak_eval()))
    assert "loads it as 0" in nb.sources["leak_gate_drops_C"]["source"]
    assert "extraction_io.py" in nb.sources["leak_gate_drops_C"]["source"]
    assert "loads it as 0" not in nb.sources["leak_gate_drops_A"]["source"]


def test_table2_quotes_the_per_extraction_count_not_the_per_run_sum():
    nb = M.Numbers()
    for k, val in (("op_leak_gate_aborts", "0"), ("op_leak_gate_drops", "16"),
                   ("leak_gate_drops_total", "6"), ("leak_gate_docs_emptied_total", "1")):
        nb.put(k, val, "x")
    spec = M.table2(nb)
    rows = {r[0]: r[1] for r in spec.rows if isinstance(r, list)}
    # The per-extraction count, in its own row and named as no refusal (review
    # 2026-10-09, T2-leakgate-cell); since 2026-10-19 the documents left with no
    # relation and the aborts are the artifact's (their keys stay).
    assert rows["Relations dropped by the leak gate, once per extraction (not a refusal)"] == \
        "«leak_gate_drops_total»"
    assert spec.header[1] == "Refused (k/N)"
    assert not any("op_leak_gate_drops" in c for r in spec.rows if isinstance(r, list)
                   for c in r)


def test_oracle_truncated_tasks_is_a_key():
    ev = _evaluation(truncated_tasks={"ids": ["task-a", "task-x"], "note": "cap"})
    nb = M.Numbers()
    M.numbers_from_truncation(nb, M.Src("evaluation.json", ev))
    assert nb["oracle_truncated_tasks"] == "A, task-x" and nb["oracle_truncated_count"] == "2"
    assert nb.sources["oracle_truncated_tasks"]["source"].endswith("/truncated_tasks/ids")
    nb = M.Numbers()
    M.numbers_from_truncation(nb, M.Src("evaluation.json", _evaluation(
        truncated_tasks={"ids": []})))
    assert nb["oracle_truncated_tasks"] == "none" and nb["oracle_truncated_count"] == "0"


def test_an_aggregate_without_a_truncated_tasks_block_is_na_not_none(capsys):
    """A missing block says nothing about the cap; it must not read as 'the cap
    truncated no task'."""
    for ev in (_evaluation(), _evaluation(truncated_tasks={"note": "cap"}),
               _evaluation(truncated_tasks={"ids": None})):
        nb = M.Numbers()
        M.numbers_from_truncation(nb, M.Src("evaluation.json", ev))
        assert nb["oracle_truncated_tasks"] == "n/a"
        assert nb["oracle_truncated_count"] == "n/a"
        assert "truncated_tasks" in capsys.readouterr().err


def _g1_comparison(name: str, covers: list[str], kind: str, differs: str = "slot_counts",
                   stub_field: bool = True) -> dict:
    labels = [f"{c}:{kind}" for c in covers]
    comp = dict(name=name, corpora=labels,
                by_arm=dict((a, dict(differs=[differs], same=["documents"]))
                            for a in M.ARM_ORDER))
    if stub_field:
        comp["stub_corpora"] = labels if kind == "stub" else []
    return comp


def _g1_probe(cid: str, covers: list[str], kind: str = "saved") -> dict:
    pr = _probe(cid)
    name = M.G1_SAVED_COMPARISON if kind == "saved" else M.G1_STUB_COMPARISON
    pr["g1"] = dict(comparisons=[_g1_comparison(name, covers, kind)])
    return pr


def test_table3_carries_the_probe_figures_per_corpus_with_na_where_no_probe(capsys):
    corpora = {"A": "task-a", "C": "task-c"}
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", _g1_probe("task-a", ["task-a"]))],
                          corpora, ["A", "C"])
    spec = M.table3(nb, corpora, ["A", "C"])
    rows = {r[0]: r for r in spec.rows
            if isinstance(r, list) and len(r) == len(M.T3_CORPUS_HEAD)}
    # Table 2(b) since 2026-10-19: one row per corpus, the corpus by its letter
    # (change set 2026-10-10, TAB-III-note), A starred as in-sample; the shared
    # extraction's echo beside its control (R3, N7); r1's store figures n/a
    # without an evaluation.
    assert rows["A*"] == [
        "A*", "n/a", "n/a", "n/a", "«recall_referents_exact_A»",
        "«extraction_entity_jaccard_A» / «extraction_relation_jaccard_A»",
        "«pred_max_tokens_A»", "«pred_echo_pct_shared_A» («pred_control_echo_pct_shared_A»)"]
    assert rows["C"] == ["C"] + ["n/a"] * (len(M.T3_CORPUS_HEAD) - 1)
    assert nb["pred_max_tokens_A"] == "9" and nb["recall_referents_exact_A"] == "3/4"
    # G1 covers A only, and C is measured: no G1 rows, and the build says why.
    assert not any(k.startswith("g1_differs_") for k in nb.values)
    assert nb["g1_corpora"] == "A" and nb["g1_corpora_missing"] == "C"
    assert "does not cover corpus C" in capsys.readouterr().err
    assert not any(isinstance(r, tuple) and r[1].startswith("G1") for r in spec.rows)
    assert any("n/a: no probe or no value" in n for n in spec.notes)


def test_g1_keys_are_written_when_the_comparison_covers_every_measured_corpus():
    """The G1 keys feed Section VI-B's sentence; since the simulated review of
    2026-10-18 Table 3 carries no G1 rows."""
    corpora = {"A": "task-a", "C": "task-c"}
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json",
                                     _g1_probe("task-a", ["task-a", "task-c", "other"]))],
                          corpora, ["A", "C"])
    assert nb["g1_differs_r1"] == "slot_counts" and nb["g1_corpora_missing"] == "none"
    assert nb.sources["g1_differs_r1"]["source"] == (
        "probe/a.json#/g1/comparisons/0/by_arm/r1/differs")
    spec = M.table3(nb, corpora, ["A", "C"])
    assert not any(isinstance(r, tuple) and r[1].startswith("G1") for r in spec.rows)
    assert not any(isinstance(r, tuple) and r[0] == "wide" and r[1] in ("Differ", "Identical")
                   for r in spec.rows)
    # The fields in words (the _text keys), as §6 prints them; the identifier keys
    # stay in numbers.json.
    assert nb["g1_differs_r1_text"] == "slot counts"


def test_the_g1_fields_are_also_named_in_words_for_running_text():
    """Follow-up F11: the identifiers stay the keys Table 3 shows; the _text keys
    name the same fields in words, joined with 'and'."""
    corpora = {"A": "task-a", "C": "task-c"}
    pr = _g1_probe("task-a", ["task-a", "task-c"])
    for d in pr["g1"]["comparisons"][0]["by_arm"].values():
        d["differs"] = ["document_stats", "slot_counts", "relation_multiset"]
        d["same"] = ["attribute_schema", "documents"]
    pr["g1"]["comparisons"][0]["by_arm"]["baseline"]["same"] = []
    pr["g1"]["comparisons"][0]["by_arm"]["r5"]["same"] = ["new_field"]
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", pr)], corpora, ["A", "C"])
    assert nb["g1_differs_r1"] == "document_stats, slot_counts, relation_multiset"
    assert nb["g1_differs_r1_text"] == "document statistics, slot counts and relation set"
    assert nb["g1_same_r1_text"] == "attribute schema and document count"
    assert nb["g1_same_baseline"] == "none" and nb["g1_same_baseline_text"] == "none"
    assert nb["g1_same_r5_text"] == "new field"
    assert nb.sources["g1_differs_r1_text"]["source"] == (
        "probe/a.json#/g1/comparisons/0/by_arm/r1/differs")
    assert set(M.G1_FIELD_NAMES) == {"document_stats", "slot_counts", "relation_multiset",
                                     "attribute_schema", "documents"}


def test_only_a_one_rule_g1_comparison_is_used(capsys):
    """The mixed comparison sets a real extraction against stubs; it cannot
    stand in for the one-rule comparison the panel describes."""
    corpora = dict(A="task-a")
    pr = _g1_probe("task-a", ["task-a"])
    pr["g1"]["comparisons"][0]["name"] = "saved extraction where given, the stub otherwise"
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", pr)], corpora, ["A"])
    assert not any(k.startswith("g1_differs_") for k in nb.values)
    assert nb["g1_corpora_missing"] == "A" and nb["g1_extraction"] == "none"
    assert "no G1 comparison named" in capsys.readouterr().err


def test_the_saved_g1_comparison_is_taken_before_the_stub(capsys):
    corpora = dict(A="task-a", C="task-c")
    pr = _g1_probe("task-a", ["task-a", "task-c"], kind="stub")
    pr["g1"]["comparisons"].append(_g1_comparison(
        M.G1_SAVED_COMPARISON, ["task-a", "task-c"], "saved", differs="relation_multiset"))
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", pr)], corpora, ["A", "C"])
    assert nb["g1_extraction"] == "saved" and nb["g1_stub_corpora"] == "none"
    assert nb["g1_differs_r1"] == "relation_multiset"
    assert nb.sources["g1_differs_r1"]["source"] == (
        "probe/a.json#/g1/comparisons/1/by_arm/r1/differs")
    assert "stub" not in capsys.readouterr().err
    spec = M.table3(nb, corpora, ["A", "C"])
    # Table 3 has no G1 rows since the simulated review of 2026-10-18.
    assert not any(r[0] == "panel" and r[1].startswith("G1") for r in spec.rows)
    assert not any("g1_stub_corpora" in n for n in spec.notes)
    M._check_stub_g1(nb, allow=False)          # nothing to refuse


def test_a_saved_g1_comparison_short_of_a_corpus_falls_back_to_the_stub_named(capsys):
    """The saved comparison covers A only; the stub covers A and C. The keys come
    from the stub, which the keys and the build name (Table 3 has no G1 rows since
    the simulated review of 2026-10-18)."""
    corpora = dict(A="task-a", C="task-c")
    pr = _g1_probe("task-a", ["task-a", "task-c"], kind="stub")
    pr["g1"]["comparisons"].insert(0, _g1_comparison(M.G1_SAVED_COMPARISON, ["task-a"],
                                                     "saved"))
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", pr)], corpora, ["A", "C"])
    assert nb["g1_extraction"] == "stub" and nb["g1_stub_corpora"] == "A, C"
    assert nb["g1_corpora"] == "A, C" and nb["g1_differs_r1"] == "slot_counts"
    assert nb.sources["g1_stub_corpora"]["source"] == "probe/a.json#/g1/comparisons/1/stub_corpora"
    assert "rest on the stub extraction for corpus A, C" in capsys.readouterr().err
    spec = M.table3(nb, corpora, ["A", "C"])
    assert not any(r[0] == "panel" and r[1].startswith("G1") for r in spec.rows)
    with pytest.raises(SystemExit, match="stub extraction for corpus A, C"):
        M._check_stub_g1(nb, allow=False)
    M._check_stub_g1(nb, allow=True)
    assert "--allow-stub-g1" in capsys.readouterr().err


def test_a_probe_written_before_stub_corpora_still_names_its_stub_rows():
    corpora = dict(A="task-a")
    pr = _probe("task-a")
    pr["g1"] = dict(comparisons=[_g1_comparison(M.G1_STUB_COMPARISON, ["task-a", "other"],
                                                "stub", stub_field=False)])
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", pr)], corpora, ["A"])
    # A corpus with no letter is named by its id.
    assert nb["g1_stub_corpora"] == "A, other" and nb["g1_extraction"] == "stub"


def test_the_g1_comparison_names_match_the_probe():
    T = pytest.importorskip("scripts.tier1_probe")     # needs the backend's dependencies
    assert M.G1_COMPARISONS == (T.G1_SAVED, T.G1_STUB)


def _bound(commit: str, runs: dict[str, str | None], manifests: dict | None = None) -> dict:
    return {"git_commit": commit,
            "arms": {"r1": {"runs": sorted(runs)}},
            "per_task": [{"run": r, "provenance": {"git_commit": c} if c else None}
                         for r, c in runs.items()],
            "health_gate": {"outcomes": {r: {"git_commit": c}
                                         for r, c in (manifests or {}).items()}}}


def test_the_freeze_commit_must_be_the_aggregates_or_the_batchs(capsys):
    agg, batch = "a" * 40, "b" * 40
    ev = M.Src("evaluation.json", _bound(agg, {"t__r1": batch, "t__r1__rep2": batch}))
    meta = lambda fc: M.Src("meta.json", {"freeze_commit": fc})   # noqa: E731
    M.check_commit_binding(ev, meta(agg[:12]), allow=False)         # the aggregate's
    M.check_commit_binding(ev, meta(batch), allow=False)            # every pooled run's
    M.check_commit_binding(ev, None, allow=False)                   # nothing claimed
    with pytest.raises(SystemExit, match="allow-commit-mismatch"):
        M.check_commit_binding(ev, meta("c" * 12), allow=False)
    with pytest.raises(SystemExit, match="neither"):
        M.check_commit_binding(ev, meta("[PILOT DEMO: no freeze]"), allow=False)
    M.check_commit_binding(ev, meta("c" * 12), allow=True)
    assert "allowed: --allow-commit-mismatch" in capsys.readouterr().err
    # A batch split over two commits, or a run with none recorded, binds nothing.
    split = M.Src("e", _bound(agg, {"t__r1": batch, "t__r1__rep2": "d" * 40}))
    with pytest.raises(SystemExit, match="bbbbbbbbbbbb, dddddddddddd"):
        M.check_commit_binding(split, meta(batch), allow=False)
    gap = M.Src("e", _bound(agg, {"t__r1": batch, "t__r1__rep2": None}))
    with pytest.raises(SystemExit, match="not recorded for t__r1__rep2"):
        M.check_commit_binding(gap, meta(batch), allow=False)
    # Phase one's manifest commit, where the aggregate carries it, is the batch's.
    man = M.Src("e", _bound(agg, {"t__r1": None}, manifests={"t__r1": batch}))
    M.check_commit_binding(man, meta(batch[:7]), allow=False)
    # The aggregate's commit alone binds, but says the runs were measured elsewhere.
    capsys.readouterr()
    M.check_commit_binding(ev, meta(agg), allow=False)
    assert "not where the runs were measured" in capsys.readouterr().err
    same = M.Src("e", _bound(agg, {"t__r1": agg}))
    M.check_commit_binding(same, meta(agg), allow=False)
    assert "not where the runs were measured" not in capsys.readouterr().err


def _cal_rows() -> list[dict]:
    """The four rows §5's sentence states, deliberately not in the sentence's order."""
    return [{"id": "heldout_B_r1", "labeller": "author", "occurrences": 30,
             "agreement": {"k": 25, "n": 28}, "kappa": 0.6933, "blind": True,
             "in_sample": False},
            {"id": "insample_A_baseline", "labeller": "assistant", "occurrences": 43,
             "agreement": {"k": 7, "n": 41}, "kappa": 0.4167, "blind": False,
             "in_sample": True},
            {"id": "dev_A_oracle", "label": "ignored", "labeller": "author", "occurrences": 44,
             "agreement": {"k": 38, "n": 42}, "kappa": 0.7755, "blind": True,
             "in_sample": False},
            {"id": "dev_A_r1", "labeller": "author", "occurrences": 23,
             "agreement": {"k": 19, "n": 23}, "kappa": 0.7121, "blind": True,
             "in_sample": False}]


def _cal_numbers(rows: list[dict]) -> M.Numbers:
    nb = M.Numbers()
    M.numbers_from_calibration(nb, M.Src("calibration.json", {"rows": rows}))
    return nb


def test_calibration_rows_become_keys_bound_by_id_not_by_order():
    nb = _cal_numbers(_cal_rows())
    assert nb["cal_insample_A_baseline_agreement"] == "7/41"
    assert nb["cal_insample_A_baseline_occurrences"] == "43"
    assert nb["cal_insample_A_baseline_kappa"] == "0.42"
    assert nb["cal_insample_A_baseline_agreement_pct"] == "17%"
    assert nb["cal_heldout_B_r1_agreement"] == "25/28" and nb["cal_dev_A_r1_kappa"] == "0.71"
    # The same file in another order gives the same keys and values.
    again = _cal_numbers(list(reversed(_cal_rows())))
    assert again.values == nb.values
    # The exploratory row is absent from the file, so it gives no key.
    assert set(nb.values) == {f"cal_{rid}_{f}" for rid in M.CAL_ROWS
                              if rid not in M.CAL_EXPLORATORY for f in M.CAL_FIELDS}
    # No calibration table any more: the sentence carries the rows.
    assert "tcal" not in M.BLOCKS and not hasattr(M, "table_cal")


def test_a_missing_calibration_row_leaves_its_keys_unproduced_so_the_fill_refuses(tmp_path):
    """The held-out row does not exist until the batch: its keys are not written,
    so a manuscript quoting them cannot be filled with n/a."""
    rows = [r for r in _cal_rows() if r["id"] != "heldout_B_r1"]
    nb = _cal_numbers(rows)
    assert not [k for k in nb.values if k.startswith("cal_heldout_B_r1_")]
    assert nb["cal_dev_A_r1_agreement"] == "19/23"
    if docx is None:
        pytest.skip("python-docx is not installed")
    src = tmp_path / "src.docx"
    d = docx.Document()
    d.add_paragraph("Held out: «cal_heldout_B_r1_agreement»; development: "
                    "«cal_dev_A_r1_agreement».")
    d.save(str(src))
    out = tmp_path / "paper"
    out.mkdir()
    (out / "numbers.json").write_text(json.dumps(nb.values), encoding="utf-8")
    with pytest.raises(M.FillError, match="cal_heldout_B_r1_agreement"):
        M.fill(src, tmp_path / "dst.docx", out, allow_na=True)
    assert not (tmp_path / "dst.docx").exists()


@pytest.mark.parametrize(("change", "match"), [
    ({"labeller": None}, "must say who labelled it"),
    ({"labeller": "someone"}, "must say who labelled it"),
    ({"labeller": "assistant"}, "blind but labelled by assistant"),       # a blind row
    ({"blind": False}, "in_sample with labeller 'assistant'"),           # author, not blind
    ({"id": "dev_B_r1"}, "has id 'dev_B_r1'"),                             # unknown id
    ({"id": None}, "has id None"),                                        # no id
    ({"id": "dev_A_oracle"}, "appears twice"),                            # repeated id
    ({"in_sample": True}, "states"),                                      # flag the sentence denies
    ({"blind": 1}, "states"),                                             # not a boolean
    ({"agreement": {"k": 24, "n": 23}}, "0 <= k <= n <= occurrences"),
    ({"occurrences": None}, "0 <= k <= n <= occurrences"),
    ({"kappa": "0.7"}, "needs kappa"),
])
def test_a_calibration_row_the_sentence_would_misdescribe_is_refused(change, match):
    rows = _cal_rows()
    dev = next(r for r in rows if r["id"] == "dev_A_r1")
    dev.update(change)
    if change.get("labeller") is None and "labeller" in change:
        del dev["labeller"]
    with pytest.raises(SystemExit, match=match):
        _cal_numbers(rows)


def test_the_in_sample_row_must_be_the_assistants_non_blind_row():
    for change in ({"blind": True}, {"in_sample": False}, {"labeller": "author"}):
        rows = _cal_rows()
        next(r for r in rows if r["id"] == "insample_A_baseline").update(change)
        with pytest.raises(SystemExit):
            _cal_numbers(rows)
    missing = _cal_rows()
    del missing[0]["kappa"]
    with pytest.raises(SystemExit, match="needs kappa"):
        _cal_numbers(missing)
    undefined = _cal_rows()
    undefined[0]["kappa"] = None                     # both readers used one label
    assert _cal_numbers(undefined)["cal_heldout_B_r1_kappa"] == "n/a"


def test_the_setup_draft_states_each_calibration_row_as_cal_rows_does():
    """Since 2026-10-29 (the author's "quá nhiều thông số trong phần Hypothesis, Evaluation",
    choosing to keep only the numbers that answer §6's questions): 05_setup.md quotes the
    planned out-of-sample row (the held-out r1 report) by its agreement and kappa keys, says
    whose labels they are and what they were blind to, and names every other row of CAL_ROWS
    as the repository's, saying of each what CAL_ROWS says (in-sample, whose labels, blind
    or not). No other calibration key is in the prose. There is no table marker left."""
    text = (_REPO / "paper" / "sections" / "05_setup.md").read_text(encoding="utf-8")
    prose = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    assert "[TABLE (inline) PLACEHOLDER]" not in text
    keys = {k for k in M.PLACEHOLDER_RE.findall(prose) if k.startswith("cal_")}
    assert keys == {"cal_heldout_B_r1_agreement", "cal_heldout_B_r1_kappa"}
    paragraph = next(s for s in prose.split("\n") if "«cal_heldout_B_r1_kappa»" in s)
    sentence = next(s for s in re.split(r"(?<=\.)\s+(?=[A-Z])", paragraph)
                    if "«cal_heldout_B_r1_kappa»" in s)
    assert M.CAL_ROWS["heldout_B_r1"] == ("author", True, False)
    assert "first author's labels" in sentence and "blind to its output but not to the arm" \
        in sentence
    assert "held-out r1 report" in sentence and "at chance" in sentence
    assert "Cohen's κ" in sentence and "gold-valued occurrences" in sentence
    assert "attribute axis" in sentence and "not validated" in sentence
    # The other rows are named as the repository's, each as CAL_ROWS describes it.
    rest = paragraph[paragraph.index(sentence) + len(sentence):]
    assert "repository" in rest
    assert "development r1 and oracle reports" in rest and "in-sample for the one-time" in rest
    assert M.CAL_ROWS["insample_A_baseline"][1] is False
    assert "baseline report on corpus A" in rest and "non-blind" in rest and "AI assistant" in rest
    assert "heldout_B_baseline" in M.CAL_EXPLORATORY and "exploratory" in rest


def test_the_exploratory_baseline_row_is_the_authors_blind_held_out_row():
    assert M.CAL_ROWS["heldout_B_baseline"] == ("author", True, False)
    assert M.CAL_EXPLORATORY == {"heldout_B_baseline"}
    rows = [*_cal_rows(), {"id": "heldout_B_baseline", "labeller": "author", "blind": True,
                           "in_sample": False, "occurrences": 31,
                           "agreement": {"k": 13, "n": 29}, "kappa": 0.1234}]
    nb = _cal_numbers(rows)
    assert nb["cal_heldout_B_baseline_agreement"] == "13/29"
    assert nb["cal_heldout_B_baseline_kappa"] == "0.12"
    rows[-1]["in_sample"] = True
    with pytest.raises(SystemExit, match="states"):
        _cal_numbers(rows)


# ------------------------------------------------------- review analyses

def _spr(mean, sd=None, lo=None, hi=None, n=3) -> dict:
    return {"n": n, "mean": mean, "stdev": sd, "min": mean if lo is None else lo,
            "max": mean if hi is None else hi}


def _cal_review_row(rid: str, status: str = "labelled", **over) -> dict:
    who, blind, ins = M.CAL_ROWS[rid]
    row = {"id": rid, "labeller": who, "blind": blind, "in_sample": ins, "status": status}
    if status == "labelled":
        cats = ("MATCH", "MISMATCH", "UNRESOLVED")
        row.update(occurrences=26, agreement={"k": 19, "n": 25}, kappa=0.4321,
                   kappa_pe=0.5432, pabak=0.64, ac1=0.6789,
                   human={"MATCH": 17, "MISMATCH": 6, "UNRESOLVED": 2},
                   rule={"MATCH": 15, "MISMATCH": 7, "UNRESOLVED": 3},
                   confusion={h: {r: (13 if h == r == "MATCH" else 1) for r in cats}
                              for h in cats},
                   bootstrap={"ci95": [0.1234, 0.7777], "resamples": 400, "undefined": 3})
    row.update(over)
    return row


def _review(**over) -> dict:
    """A small synthetic review_analyses.json: two corpora, every section."""
    den = {"scanned": 30, "gold_shaped": 20, "referent": 5, "attribute": 18, "full": 4}
    cell = {"runs": 3, "sum": den,
            "rates": {f: _spr(0.25, 0.1, 0.1, 0.4) for f in ("referent", "attribute", "full")}}
    ra = {
        "instrument": "scripts/review_analyses.py",
        "inputs": {"evaluation.json": "e" * 64, "observation.json": "o" * 64,
                   "calibration.json": "c" * 64},
        "corpora": {"K": "corpus-k", "Q": "corpus-q"},
        "resolution": {
            "runs": [{"run": "corpus-k__r1", "letter": "K", "arm": "r1", "rep": 1,
                      "private": den, "oracle": den},
                     {"run": "corpus-k__baseline", "letter": "K", "arm": "baseline",
                      "rep": 1, "withheld": "its report is under blind labelling"}],
            "cells": {"K": {"r1": cell, "oracle": cell}, "Q": {"baseline": cell}},
            "arms": {"r1": cell, "oracle": cell, "baseline": cell}},
        "lexical_overlap": {
            "lexicons_identical_across_keys": True, "lexicon_phrases": 40,
            "codes": [{"code": "AMTA", "rba_attribute": "AUTHORISED"},
                      {"code": "AMTB", "rba_attribute": "SPENT"},
                      {"code": "NOTE", "rba_attribute": None}],
            "meanings_with_phrase": 2, "phrases_in_meanings": ["alpha", "beta gamma"],
            "money_codes": {"own_only": ["AMTA"], "mixed": [], "other_only": ["AMTB"],
                            "none": []},
            "r1_draft_requests": 16, "r1_draft_requests_with_meanings": 16,
            "usage": {"r1": {"runs": 3, "hits": 40, "shared": 10, "words": 2000},
                      "oracle": {"runs": 9, "hits": 0, "shared": 0, "words": 0}}},
        "coverage_split": {
            "facts": {"K": {"stated": 9, "derived": 3}, "Q": {"stated": 7, "derived": 2}},
            "cells": {"K": {"r1": {"runs": 3, "stated": _spr(0.4444, 0.05, 0.4, 0.5),
                                   "derived": _spr(0.0, 0.0), "stated_pooled": {"k": 12, "n": 27},
                                   "derived_pooled": {"k": 0, "n": 9}}}}},
        "without_share_spent": {
            "letter": "K", "facts": 10, "removed_facts": ["P.X", "Q.X"],
            "cells": {"r1": {"runs": 3, "rba": _spr(0.8123, 0.02, 0.79, 0.83),
                             "attribute_accuracy": _spr(0.9, 0.01, 0.89, 0.91),
                             "coverage_correct": _spr(0.7, 0.0)},
                      "baseline": {"runs": 3, "rba": _spr(0.5, n=1),
                                   "attribute_accuracy": _spr(0.5), "coverage_correct":
                                   _spr(0.1)}}},
        "binding": {
            "cells": {"K": {
                "baseline": {"runs": 3, "money_slots": [21], "gold_money_slots": [17, 18],
                             "gold_money_values": [9],
                             "attributes_per_referent": {"min": 2, "max": 3, "mean": 2.5,
                                                         "referents": 4},
                             "attribute": {"n": 40, "chance": 0.3333, "observed_k": 14,
                                           "observed": 0.35},
                             "full": {"n": 0, "chance": None, "observed_k": 0,
                                      "observed": None}},
                "r1": {"runs": 3, "referent_slots": [13], "money_slots": [0], "codes": 13,
                       "facts_in_store": [22]}}},
            "arms": {"baseline": {"attribute": {"n": 40, "chance": 0.3333, "observed_k": 14,
                                                "observed": 0.35},
                                  "full": {"n": 0, "chance": None, "observed_k": 0,
                                           "observed": None}}}},
        "calibration": {"bootstrap": {"resamples": 400, "seed": 1},
                        "rows": [_cal_review_row("dev_A_r1"),
                                 _cal_review_row("heldout_B_baseline", status="unlabelled")]},
        "corpus_stats": {
            L: {"documents": 4, "characters": ch, "words": w, "table_data_cells": tc,
                "referents": 5, "facts": st + dv,
                "facts_by_kind": {"table": tb, "prose": st - tb, "derived": dv}}
            for L, ch, w, tc, st, dv, tb in (("K", 3100, 470, 24, 9, 3, 6),
                                              ("Q", 3300, 455, 12, 7, 2, 0))},
        "maskdoc_without_lost": {
            "K": {"extractions": [
                {"file": "eval/runs/corpus-k__r5/extractions.json", "lost_documents": ["d4"],
                 "kept": {"surfaces_distinct": 4, "masked_5gram_hit_share": 0.0647}},
                {"file": "eval/probe/extractions/corpus-k/extraction_1.json",
                 "lost_documents": [],
                 "kept": {"surfaces_distinct": 5, "masked_5gram_hit_share": 0.105}}]},
            "Q": {"extractions": [
                {"file": "eval/probe/extractions/corpus-q/extraction_1.json",
                 "lost_documents": [],
                 "kept": {"surfaces_distinct": 2, "masked_5gram_hit_share": 0.126}}]}},
        "cost": {
            "arms": {"r1": {"deploy_usd": _spr(0.1234, 0.01, 0.099, 0.151),
                            "eval_usd": _spr(0.25, 0.02, 0.2, 0.29),
                            "seconds_reused": _spr(89.4, 9.0, 73.3, 107.8),
                            "seconds_extracted": {"n": 0, "mean": None}}},
            "all": {"deploy_usd": _spr(0.11, 0.01, 0.06, 0.15),
                    "eval_usd": _spr(0.26, 0.03, 0.18, 0.34),
                    "seconds_reused": _spr(84.8, 10.0, 49.9, 112.7),
                    "seconds_extracted": _spr(166.9, 20.0, 145.2, 195.3, n=4)},
            "totals": {"deploy_usd": 4.023, "eval_usd": 9.42, "deploy_input_tokens": 7000,
                       "deploy_output_tokens": 2600,
                       "eval_usd_by_role": {"oracle": 1.8, "judge": 4.3, "adversary": 3.32}}},
    }
    ra.update(over)
    return ra


def _review_numbers(ra: dict, nb: M.Numbers | None = None) -> M.Numbers:
    nb = nb or M.Numbers()
    M.numbers_from_review(nb, M.Src("review_analyses.json", ra))
    return nb


def test_review_analyses_become_keys_with_the_paper_formats():
    nb = _review_numbers(_review())
    # MR3a: counts, k/N over gold-valued occurrences, pooled %, per-run cells.
    assert nb["res_K_r1_scanned"] == "30" and nb["res_K_r1_gold"] == "20"
    assert nb["res_K_r1_both"] == "4/20" and nb["res_K_r1_attr_pct"] == "90.0%"
    assert nb["res_K_r1_ref_rate"] == "0.25 [0.10–0.40]"
    assert nb["res_oracle_both"] == "4/20" and nb["res_K_r1_rep1_attr"] == "18/20"
    assert "res_K_r1_rep1_attr_pct" not in nb.values
    # MR3b: lexical overlap, coverage split, A without the share-spent facts.
    assert nb["lex_meanings_with_phrase"] == "2/3" and nb["lex_money_codes"] == "2"
    assert nb["lex_phrases_in_meanings"] == "2/40"
    assert nb["lex_phrases_in_meanings_list"] == "“alpha” and “beta gamma”"
    assert nb["lex_money_codes_other_only_list"] == "AMTB"
    assert nb["lex_money_codes_mixed_list"] == "none"
    assert nb["lex_use_r1_shared"] == "10/40" and nb["lex_use_r1_shared_pct"] == "25.0%"
    assert nb["lex_use_r1_per_kword"] == "20.0" and nb["lex_use_oracle_shared_pct"] == "n/a"
    assert nb["lex_r1_requests_with_meanings"] == "16/16"
    assert nb["covsplit_K_r1_stated"] == "0.44 [0.40–0.50]"
    assert nb["covsplit_K_r1_stated_kn"] == "12/27"
    assert nb["noshare_K_r1_rba"] == "0.81 [0.79–0.83]"
    assert nb["noshare_K_baseline_rba"] == "0.50 (1 of 3 runs)"
    assert nb["noshare_rba_baseline_K"] == "0.50" and nb["noshare_rba_baseline_K_n"] == "1/3"
    assert nb["noshare_K_facts"] == "10" and nb["noshare_K_removed"] == "2"
    # MR4b: what each arm can bind, and the chance model.
    assert nb["bind_K_baseline_money_slots"] == "21"
    assert nb["bind_K_baseline_gold_slots"] == "17–18"
    assert nb["bind_K_baseline_attrs_per_referent"] == "2–3"
    assert nb["bind_K_r1_referent_slots"] == "13" and nb["bind_K_r1_facts"] == "22"
    assert nb["bind_r1_codes"] == "13"
    assert nb["chance_K_baseline_attr"] == "0.33" and nb["chance_K_baseline_attr_obs"] == "0.35"
    assert nb["chance_K_baseline_attr_kn"] == "14/40"
    assert nb["chance_baseline_full"] == "n/a" and nb["chance_baseline_full_kn"] == "0/0"
    # MR9a, MR9b, minor 8.
    assert nb["corpus_chars"] == "3100–3300" and nb["corpus_docs"] == "4"
    assert nb["corpus_facts_table_Q"] == "0" and nb["corpus_facts_stated_K"] == "9"
    assert nb["corpus_table_cells"] == "12–24"
    assert nb["maskdoc_surfaces_surviving_nolost_K"] == "4–5"
    assert nb["maskdoc_surfaces_surviving_shared_nolost_K"] == "4"
    assert nb["maskdoc_5gram_pct_shared_nolost_K"] == "6.5%"
    assert "maskdoc_surfaces_surviving_shared_nolost_Q" not in nb.values
    assert nb["maskdoc_surfaces_surviving_nolost"] == "2–5"
    assert nb["maskdoc_5gram_pct_nolost"] == "6.5%–12.6%"
    assert nb["maskdoc_lost_docs"] == "1" and nb["maskdoc_lost_docs_Q"] == "0"
    assert nb["cost_deploy_usd_r1"] == "0.12"
    assert nb["cost_deploy_usd_r1_range"] == "0.12 [0.10–0.15]"
    assert nb["cost_eval_usd"] == "0.26" and nb["cost_deploy_usd_total"] == "4.02"
    assert nb["cost_eval_usd_total_judge"] == "4.30"
    assert nb["cost_deploy_tokens_input_total"] == "7000"
    assert nb["latency_phase_one_s_r1"] == "89 [73–108]"
    assert nb["latency_phase_one_extract_s"] == "167 [145–195]"
    assert nb["latency_phase_one_extract_n"] == "4"
    # Every key says where it came from.
    assert all(nb.sources[k]["source"].startswith("review_analyses.json#/") for k in nb.values)


def test_a_run_under_blind_labelling_gets_no_per_run_key():
    nb = _review_numbers(_review())
    assert not [k for k in nb.values if k.startswith("res_K_baseline_rep")]
    assert nb["res_Q_baseline_gold"] == "20"          # its cell still counts it


def test_review_calibration_rows_give_statistics_and_an_unlabelled_row_gives_none():
    nb = _review_numbers(_review())
    assert nb["cal_dev_A_r1_agreement"] == "19/25" and nb["cal_dev_A_r1_kappa"] == "0.43"
    assert nb["cal_dev_A_r1_agreement_pct"] == "76%"
    assert nb["cal_dev_A_r1_human_match"] == "17/25"
    assert nb["cal_dev_A_r1_rule_unresolved"] == "3/25"
    assert nb["cal_dev_A_r1_cm_match_match"] == "13"
    assert nb["cal_dev_A_r1_cm_match_mismatch"] == "1"
    assert nb["cal_dev_A_r1_pabak"] == "0.64" and nb["cal_dev_A_r1_ac1"] == "0.68"
    assert nb["cal_dev_A_r1_kappa_pe"] == "0.54"
    assert nb["cal_dev_A_r1_kappa_ci"] == "[0.12, 0.78]"
    assert nb["cal_dev_A_r1_kappa_ci_undefined"] == "3/400"
    assert nb["cal_bootstrap_resamples"] == "400"
    assert not [k for k in nb.values if k.startswith("cal_heldout_B_baseline_")]
    # Labelled, the exploratory row gives every key.
    ra = _review()
    ra["calibration"]["rows"][1] = _cal_review_row("heldout_B_baseline")
    nb = _review_numbers(ra)
    assert nb["cal_heldout_B_baseline_agreement"] == "19/25"
    assert nb["cal_heldout_B_baseline_kappa_ci"] == "[0.12, 0.78]"


def test_review_calibration_figures_must_agree_with_calibration_json():
    rows = _cal_rows()
    dev = next(r for r in rows if r["id"] == "dev_A_r1")
    dev.update(occurrences=26, agreement={"k": 19, "n": 25}, kappa=0.4321)
    nb = _cal_numbers(rows)
    _review_numbers(_review(), nb)                     # the same figures: kept as they are
    assert nb.sources["cal_dev_A_r1_kappa"]["source"].startswith("calibration.json#/")
    dev["kappa"] = 0.5
    with pytest.raises(SystemExit, match="cal_dev_A_r1_kappa"):
        _review_numbers(_review(), _cal_numbers(rows))


@pytest.mark.parametrize(("change", "match"), [
    ({"instrument": "something else"}, "not scripts/review_analyses.py output"),
    ({"calibration": {"bootstrap": {"resamples": 1}, "rows": [
        _cal_review_row("dev_A_r1", blind=False)]}}, "in_sample with labeller"),
    ({"calibration": {"bootstrap": {"resamples": 1}, "rows": [
        _cal_review_row("dev_A_r1"), _cal_review_row("dev_A_r1")]}}, "appears twice"),
])
def test_review_analyses_that_misdescribe_themselves_are_refused(change, match):
    with pytest.raises(SystemExit, match=match):
        _review_numbers(_review(**change))


def test_corpus_counts_must_agree_with_the_coverage_split_and_the_aggregate():
    ra = _review()
    ra["coverage_split"]["facts"]["K"]["derived"] = 4
    with pytest.raises(SystemExit, match="stated and derived facts differ"):
        _review_numbers(ra)
    nb = M.Numbers()
    nb.put("table_cells_K", "25", "evaluation.json#/by_task/corpus-k/table_data_cells")
    with pytest.raises(SystemExit, match="table data cells"):
        _review_numbers(_review(), nb)


def test_review_analyses_from_other_inputs_are_refused(tmp_path):
    ev = tmp_path / "evaluation.json"
    ev.write_text("{}", encoding="utf-8")
    sha = hashlib.sha256(ev.read_bytes()).hexdigest()
    ra = M.Src("review_analyses.json", _review(inputs={"evaluation.json": sha}))
    M.check_review_inputs(ra, {"evaluation.json": ev, "calibration.json": None})
    ev.write_text("{ }", encoding="utf-8")
    with pytest.raises(SystemExit, match="another evaluation.json"):
        M.check_review_inputs(ra, {"evaluation.json": ev})


# ------------------------------------------- post-hoc analyses (scope §13)
# Added by the AI assistant on 2026-10-26 with the --posthoc reader (scope §13,
# signed by the author that day; S6 and S13 of the simulated review of 2026-10-26).

_PH_CORPORA = {"K": "corpus-k", "Q": "corpus-q"}


def _ph_cliff(n1, n2, delta):
    return {"n1": n1, "n2": n2, "greater": None if delta is None else 0,
            "less": None if delta is None else 0, "ties": None if delta is None else 0,
            "delta": delta}


def _posthoc(**over) -> dict:
    """A small synthetic posthoc.json: two corpora, both analyses."""
    ph = {
        "instrument": "scripts/posthoc_2026_10_26.py",
        "scope": "PRIVDEEP_ICCIT_SCOPE.md §13",
        "inputs": {"evaluation.json": "0" * 64},
        "reproduced": {"judge_mean_exgr_r1": "2.20", "attr_r1_K": "0.78"},
        "s6": {
            "K": {"task": "corpus-k", "n_runs": 9,
                  "range_over_runs": {"2-4": {"min": 22, "max": 23}},
                  "masked_documents": {"2-4": 493}},
            "Q": {"task": "corpus-q", "n_runs": 9,
                  "range_over_runs": {"2-4": {"min": 6, "max": 6}},
                  "masked_documents": {"2-4": 156}},
        },
        "s13": {
            "judged_quality": {
                "baseline": {"n": 12, "mean": 2.18, "sd": 0.266, "min": 1.6667, "max": 2.6667,
                             "ci95_bootstrap": [2.037, 2.3287]},
                "r5": {"n": 12, "mean": 2.236, "sd": 0.1743, "min": 2.0, "max": 2.5,
                       "ci95_bootstrap": [2.1435, 2.3287]},
                "r1": {"n": 12, "mean": 2.199, "sd": 0.3138, "min": 1.6667, "max": 2.7778,
                       "ci95_bootstrap": [2.0324, 2.3657]},
            },
            "bootstrap": {"resamples": 10000, "seed": 20261026, "level": 0.95},
            "cliffs_delta_r1_vs_baseline": {
                "K": {"attribute_accuracy": _ph_cliff(3, 3, 1.0),
                      "rba": _ph_cliff(3, 1, -0.3333)},
                "Q": {"attribute_accuracy": _ph_cliff(0, 3, None),
                      "rba": _ph_cliff(0, 0, None)},
            },
        },
    }
    ph.update(over)
    return ph


def _posthoc_numbers(ph: dict, nb: M.Numbers | None = None,
                     corpora: dict[str, str] | None = None) -> M.Numbers:
    nb = nb or M.Numbers()
    M.numbers_from_posthoc(nb, M.Src("posthoc.json", ph), corpora or _PH_CORPORA)
    return nb


def test_posthoc_analyses_become_keys_with_the_paper_formats():
    before = M.Numbers()
    before.put("judge_mean_exgr_r1", "2.20", "evaluation.json#/arms/r1/judge/mean")
    before.put("attr_r1_K", "0.78", "evaluation.json#/by_task/corpus-k")
    before.put("n_runs_r1", "12", "evaluation.json#/arms/r1/runs")
    kept = dict(before.values)
    nb = _posthoc_numbers(_posthoc(), before)
    # The keys that were there are untouched; every new key points into posthoc.json.
    assert {k: nb[k] for k in kept} == kept
    assert all(nb.sources[k]["source"].startswith("posthoc.json#/") for k in nb.values
               if k not in kept)
    # S6: lo–hi over runs (lo alone when they agree), the masked documents, ranges.
    assert nb["s6_core_ngrams_K"] == "22–23" and nb["s6_core_ngrams_Q"] == "6"
    assert nb["s6_masked_ngrams_K"] == "493" and nb["s6_masked_ngrams_Q"] == "156"
    assert nb["s6_core_ngrams_range"] == "6–23"
    assert nb["s6_masked_ngrams_range"] == "156–493"
    # S13: judged quality per arm at 2 dp, the interval as the kappa CI keys print one.
    assert nb["judge_sd_baseline"] == "0.27" and nb["judge_sd_r5"] == "0.17"
    assert nb["judge_sd_range"] == "0.17–0.31"
    assert nb["judge_minmax_r1"] == "1.67–2.78" and nb["judge_minmax_r5"] == "2.00–2.50"
    assert nb["judge_ci_baseline"] == "[2.04, 2.33]" and nb["judge_ci_r1"] == "[2.03, 2.37]"
    assert nb["judge_ci_resamples"] == "10000"
    # Cliff's delta: 2 dp, n/a where a side has no run, each with its runs.
    assert nb["cliff_attr_K"] == "1.00" and nb["cliff_rba_K"] == "-0.33"
    assert nb["cliff_attr_K_n_r1"] == "3" and nb["cliff_attr_K_n_baseline"] == "3"
    assert nb["cliff_rba_K_n_baseline"] == "1"
    assert nb["cliff_attr_Q"] == "n/a" and nb["cliff_attr_Q_n_r1"] == "0"
    assert nb["cliff_rba_Q"] == "n/a" and nb["cliff_rba_Q_n_baseline"] == "0"
    assert nb["cliff_attr_runs"] == "3"
    assert nb.sources["cliff_attr_K"]["source"] == (
        "posthoc.json#/s13/cliffs_delta_r1_vs_baseline/K/attribute_accuracy/delta")


def _ph_changed(path: tuple, value) -> dict:
    ph = _posthoc()
    cur = ph
    for p in path[:-1]:
        cur = cur[p]
    cur[path[-1]] = value
    return ph


@pytest.mark.parametrize(("ph", "match"), [
    (_posthoc(instrument="scripts/review_analyses.py"), "not scripts/posthoc_2026_10_26.py"),
    (_ph_changed(("s6", "Q", "task"), "corpus-z"), "corpus Q is 'corpus-z'"),
    (_ph_changed(("s6", "K", "range_over_runs", "2-4", "min"), None), "expected a count"),
    (_ph_changed(("s13", "cliffs_delta_r1_vs_baseline", "Q", "rba", "delta"), 0.5),
     "a delta needs a run with a value on each side"),
    (_ph_changed(("s13", "cliffs_delta_r1_vs_baseline", "K", "rba", "delta"), 1.5),
     "expected a number in"),
    (_ph_changed(("s13", "judged_quality", "r1", "ci95_bootstrap"), [2.0]),
     "expected \\[lo, hi\\]"),
])
def test_posthoc_analyses_that_misdescribe_themselves_are_refused(ph, match):
    with pytest.raises(SystemExit, match=match):
        _posthoc_numbers(ph)


def _ph_without(path: tuple) -> dict:
    ph = _posthoc()
    cur = ph
    for p in path[:-1]:
        cur = cur[p]
    del cur[path[-1]]
    return ph


@pytest.mark.parametrize(("ph", "match"), [
    (_ph_without(("s6", "Q")), "s6 covers \\['K'\\]"),
    (_ph_without(("s13", "cliffs_delta_r1_vs_baseline", "Q")), "Cliff's delta covers"),
    (_posthoc(reproduced={}), "no reproduced frozen keys"),
])
def test_posthoc_analyses_that_cover_less_than_the_aggregate_are_refused(ph, match):
    """Final check of 2026-10-26: a file missing a corpus would make the ranges over
    corpora cover fewer silently, and one with no reproduced keys skips the tie to the
    frozen build."""
    with pytest.raises(SystemExit, match=match):
        _posthoc_numbers(ph)


def test_posthoc_analyses_from_other_data_are_refused(tmp_path):
    # A frozen key the analyses reproduced must be this build's value.
    nb = M.Numbers()
    nb.put("attr_r1_K", "0.77", "evaluation.json#/by_task/corpus-k")
    with pytest.raises(SystemExit, match="attr_r1_K"):
        _posthoc_numbers(_posthoc(), nb)
    # So must the arm's run count and mean.
    nb = M.Numbers()
    nb.put("n_runs_r5", "11", "evaluation.json#/arms/r5/runs")
    with pytest.raises(SystemExit, match="pools 11 r5 runs"):
        _posthoc_numbers(_posthoc(), nb)
    # And the evaluation.json the script read, by sha256, as for review_analyses.json.
    ev = tmp_path / "evaluation.json"
    ev.write_text("{}", encoding="utf-8")
    sha = hashlib.sha256(ev.read_bytes()).hexdigest()
    ph = M.Src("posthoc.json", _posthoc(inputs={"evaluation.json": sha}))
    M.check_review_inputs(ph, {"evaluation.json": ev}, M.POSTHOC_INSTRUMENT)
    ev.write_text("{ }", encoding="utf-8")
    with pytest.raises(SystemExit, match="rerun scripts/posthoc_2026_10_26.py"):
        M.check_review_inputs(ph, {"evaluation.json": ev}, M.POSTHOC_INSTRUMENT)


def test_posthoc_is_an_optional_input_of_the_build(capsys):
    with pytest.raises(SystemExit) as exc:
        M.main(["--help"])
    assert exc.value.code == 0
    assert "--posthoc POSTHOC_JSON" in capsys.readouterr().out


# ------------------------------------------------------------ fill pass

try:
    import docx
except ImportError:          # pragma: no cover - python-docx is in both interpreters
    docx = None
needs_docx = pytest.mark.skipif(docx is None, reason="python-docx is not installed")

VALUES = {"alpha": "0.41", "beta": "0.63", "count": "41"}


def _doc_with_split_placeholder(path: Path) -> None:
    d = docx.Document()
    p = d.add_paragraph()
    r1 = p.add_run("R1 scores ")
    r2 = p.add_run("«al")
    r2.bold = True
    p.add_run("ph")
    p.add_run("a» and «beta», unchanged tail.")
    d.add_paragraph("An author sentence with no placeholder stays as typed.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "«count» runs"
    t.cell(0, 1).text = "plain"
    d.sections[0].header.paragraphs[0].text = "Header «count»"
    d.add_paragraph("[TABLE 1 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["t1"]
                    + " The author's own sentence.")
    d.add_paragraph("[FIGURE 3 PLACEHOLDER] Wire excerpt.")
    assert r1.text
    d.save(str(path))


def _out(tmp_path: Path, values: dict[str, str]) -> Path:
    out = tmp_path / "out"
    out.mkdir()
    (out / "numbers.json").write_text(json.dumps(values), encoding="utf-8")
    return out


def _value_bookmarks(d) -> list[str]:
    from docx.oxml.ns import qn
    return [b.get(qn("w:name")) for b in d.element.body.iter(qn("w:bookmarkStart"))
            if (b.get(qn("w:name")) or "").startswith("pv_")]


@needs_docx
def test_fill_replaces_a_placeholder_split_across_runs_and_tags_it(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    _doc_with_split_placeholder(src)
    out = _out(tmp_path, VALUES)
    r = M.fill(src, dst, out)
    d = docx.Document(str(dst))
    p = d.paragraphs[0]
    assert p.text == "R1 scores 0.41 and 0.63, unchanged tail."
    # The value is a run of its own, with the format of the run that held «.
    assert [run.text for run in p.runs if run.bold] == ["0.41"]
    assert d.paragraphs[1].text == "An author sentence with no placeholder stays as typed."
    assert d.tables[0].cell(0, 0).text == "41 runs"
    assert d.sections[0].header.paragraphs[0].text == "Header 41"
    assert r["replaced"] == 4 and r["values_tagged"] == 4
    # Each value sits between a bookmark named after its key, unique in the file.
    names = _value_bookmarks(d)
    assert [n.rsplit("_", 1)[0] for n in names] == [C.value_tag("alpha"), C.value_tag("beta"),
                                                    C.value_tag("count")]
    assert len(set(names)) == len(names)
    # Nothing to insert: the block placeholders are left exactly as they were.
    assert sorted(r["placeholders_left"]) == ["[FIGURE 3 PLACEHOLDER]",
                                              "[TABLE 1 PLACEHOLDER]"]
    assert any(p.text.startswith("[TABLE 1 PLACEHOLDER]") for p in d.paragraphs)
    # The source is not modified: it keeps its placeholders for the next refill.
    assert "«beta»" in "".join(p.text for p in docx.Document(str(src)).paragraphs)


@needs_docx
def test_fill_refuses_a_placeholder_without_a_value_and_writes_nothing(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    _doc_with_split_placeholder(src)
    out = _out(tmp_path, {"alpha": "0.41", "count": "41"})
    with pytest.raises(M.FillError, match="beta"):
        M.fill(src, dst, out)
    assert not dst.exists()


@needs_docx
def test_fill_inserts_fig3_and_its_caption_not_the_strip_plot(tmp_path):
    # Fig. 3 (Fig. 4 for part of 2026-10-19) is the per-corpus binding figure
    # (fig3_binding.png);
    # the strip plot an earlier build left is the artifact's and is never inserted
    # (plan Part D AD33). The caption's keys are checked before anything is written.
    from paper.tools.page_estimate import _blank_png
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d = docx.Document()
    d.add_paragraph("R1 scores «alpha».")
    d.add_paragraph(f"{M.BLOCKS['fig3']} {M.BLOCK_DESCRIPTIONS['fig3']}")
    d.save(str(src))
    out = _out(tmp_path, VALUES)
    (out / "fig3_rba_strip.png").write_bytes(_blank_png(6, 2))
    (out / M.FIG_RESULTS_CAPTION).write_text("Fig. 3. Over «no_such_key» runs.\n", encoding="utf-8")
    r = M.fill(src, dst, out)
    assert r["inserted"] == [] and r["placeholders_left"] == [M.BLOCKS["fig3"]]
    (out / M.FIG_RESULTS_PNG).write_bytes(_blank_png(6, 2))
    with pytest.raises(M.FillError, match="no_such_key"):
        M.fill(src, dst, out)
    (out / M.FIG_RESULTS_CAPTION).write_text("Fig. 3. Scores at «alpha».\n", encoding="utf-8")
    r = M.fill(src, dst, out)
    assert r["inserted"] == ["fig3"] and r["placeholders_left"] == []
    from docx.oxml.ns import qn
    filled = docx.Document(str(dst))
    assert next(filled.element.body.iter(qn("w:drawing")), None) is not None
    texts = [p.text for p in filled.paragraphs]
    assert texts[0] == "R1 scores 0.41." and texts[-1].endswith("Scores at 0.41.")
    assert not any(t.startswith("[FIGURE 3") for t in texts)


@needs_docx
def test_fill_refuses_an_unmatched_guillemet(tmp_path):
    d = docx.Document()
    d.add_paragraph("A broken «alpha placeholder.")
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    with pytest.raises(M.FillError, match="unmatched"):
        M.fill(src, dst, _out(tmp_path, VALUES))
    assert not dst.exists()


@needs_docx
def test_fill_refuses_na_in_running_text_unless_allowed(tmp_path):
    d = docx.Document()
    d.add_paragraph("R1 raises accuracy to «alpha».")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "«alpha»"
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    out = _out(tmp_path, {"alpha": "n/a"})
    with pytest.raises(M.FillError, match="n/a in running text for: alpha"):
        M.fill(src, dst, out)
    assert not dst.exists()
    r = M.fill(src, dst, out, allow_na=True)
    assert r["na_in_prose"] == ["alpha"]
    res = C.check(dst, out / "numbers.json")
    assert [f.kind for f in res.fails] == ["na-in-prose", "empty"]


def _add_zip_entry(path: Path, name: str, xml: str) -> None:
    with zipfile.ZipFile(path, "a") as z:
        z.writestr(name, xml)


_FOOTNOTES = ('<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
              '2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>{}</w:t></w:r></w:p>'
              '</w:footnote></w:footnotes>')


@needs_docx
def test_fill_refuses_a_placeholder_in_a_part_it_cannot_edit(tmp_path):
    d = docx.Document()
    d.add_paragraph("Body «alpha».")
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    _add_zip_entry(src, "word/footnotes.xml", _FOOTNOTES.format("See «beta»."))
    with pytest.raises(M.FillError, match="/word/footnotes.xml"):
        M.fill(src, dst, _out(tmp_path, VALUES))
    assert not dst.exists()


@needs_docx
def test_fill_inserts_tables_keeps_author_text_and_writes_a_banner(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    _doc_with_split_placeholder(src)
    out = _out(tmp_path, VALUES)
    spec = M.TableSpec("t1", "Table 1", "Faithfulness over «count» runs.", ["Arm", "RBA"],
                       [("panel", "Corpus A"), ["r1", "«alpha»"]], notes=["A note."])
    M.write_tables_docx(src, [spec], out / "tables.docx", VALUES)
    r = M.fill(src, dst, out, banner=M.BANNER_PILOT)
    assert r["inserted"] == ["t1"] and r["text_kept"] == ["t1"]
    d = docx.Document(str(dst))
    texts = [p.text.replace('\xa0', ' ') for p in d.paragraphs]
    assert not any("[TABLE 1 PLACEHOLDER]" in t or "Faithfulness and utility" in t
                   for t in texts)
    # The skeleton's description went with the marker; the author's sentence stayed.
    assert "The author's own sentence." in texts
    assert "Table 1. Faithfulness over 41 runs." in texts and "A note." in texts
    rows = [[c.text for c in row.cells] for row in d.tables[-1].rows]
    assert rows[0] == ["Arm", "RBA"] and rows[-1] == ["r1", "0.41"]
    assert texts[0] == M.BANNER_PILOT
    assert M.BANNER_PILOT in "".join(p.text for p in d.sections[0].header.paragraphs)
    # Table values are tagged too, and every tag in the filled file is unique:
    # alpha, beta, the cell's count, the caption's count and the table's alpha
    # in the body (the header's count is in its own part).
    names = _value_bookmarks(d)
    assert len(names) == len(set(names)) == 5
    from docx.oxml.ns import qn
    ids = [b.get(qn("w:id")) for b in d.element.body.iter(qn("w:bookmarkStart"))]
    assert len(ids) == len(set(ids))


@needs_docx
def test_a_banner_keeps_the_running_head_that_later_sections_inherit(tmp_path):
    """On a paginated source (paper_format.paginate), the banner goes into the one
    header every section inherits, beside the running head: a later section gets no
    header of its own, which would lack the running head (verification of
    2026-10-25)."""
    from scripts import paper_format as F
    d = docx.Document(str(_REPO / "paper" / "ICCIT2026_template.docx"))
    for el in list(d.element.body):
        if el.tag != F._qn("w:sectPr"):
            d.element.body.remove(el)
    F.apply(d)
    p = d.add_paragraph("Title", style=F.TITLE)
    F.end_section(d, p._p, 1, start=None)
    d.add_paragraph("Body.", style=F.BODY)
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    M.fill(src, dst, _out(tmp_path, VALUES), banner=M.BANNER_PILOT)
    f = docx.Document(str(dst))
    assert [p.text for p in f.sections[0].header.paragraphs] == [F.RUNNING_HEAD,
                                                                M.BANNER_PILOT]
    assert all(s.header.is_linked_to_previous for s in f.sections[1:])
    assert all(s.footer.is_linked_to_previous for s in f.sections[1:])


@needs_docx
def test_a_marker_paragraph_with_only_the_description_is_removed(tmp_path):
    d = docx.Document()
    d.add_paragraph("[TABLE 1 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["t1"])
    d.add_paragraph("After.")
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    out = _out(tmp_path, VALUES)
    M.write_tables_docx(src, [M.TableSpec("t1", "Table 1", "Cap.", ["A"], [["x"]])],
                        out / "tables.docx", VALUES)
    r = M.fill(src, dst, out)
    assert r["text_kept"] == []
    assert [p.text.replace('\xa0', ' ') for p in docx.Document(str(dst)).paragraphs] == [
        "Table 1. Cap.", "After."]


@needs_docx
def test_placeholder_keys_finds_split_keys(tmp_path):
    src = tmp_path / "src.docx"
    _doc_with_split_placeholder(src)
    keys = M.placeholder_keys(docx.Document(str(src)))
    assert set(keys) == {"alpha", "beta", "count"}


# The keys the two tables carried before the writing fallback folded them:
# the what-crosses table (per arm, corpora pooled; per corpus from the probe;
# G1 per arm) and the faithfulness table (per corpus and arm).
_FORMER_T3_ARM = ("obs_{a}_slots_adv", "obs_{a}_slots_wire",
                  "obs_{a}_relations", "obs_{a}_fact_refs", "obs_{a}_pruned",
                  "obs_{a}_req_bytes", "obs_{a}_resp_bytes", "obs_{a}_plan_input",
                  "obs_{a}_plan_output", "obs_{a}_draft_input", "obs_{a}_draft_output",
                  "obs_{a}_usd", "obs_{a}_cost_complete", "cross_a_{a}", "cross_a_label_{a}",
                  "cross_a_scanned_{a}", "cross_b_core_{a}", "cross_b_cloud_{a}",
                  "cross_b_scanned_{a}", "g3_entities_{a}", "g3_5grams_{a}")
_FORMER_T3_CORPUS = ("pred_max_tokens_{L}", "pred_echo_pct_{L}", "pred_control_echo_pct_{L}",
                     "recall_referents_exact_{L}", "extraction_entity_jaccard_{L}",
                     "extraction_relation_jaccard_{L}")
_FORMER_T4 = ("t4_{L}_{a}_rba", "t4_{L}_{a}_attr", "t4_{L}_{a}_covcorrect",
              "t4_{L}_{a}_covstated", "t4_{L}_{a}_pruned", "t4_{L}_{a}_factsstruct")
# The run counts: in numbers.json still, but out of the table since the evaluation
# rewrite of 2026-10-15 (the note states «n_registered_repeats» per private arm;
# the pooled cells say how many runs were scanned, k/N).
# The G1 keys, identifiers and words, left the table with its G1 rows (simulated
# review of 2026-10-18); Section VI-B quotes the _text keys.
_RUN_KEYS = ("t4_{L}_{a}_runs", "obs_{a}_runs", "g1_differs_{a}", "g1_same_{a}",
             "g1_differs_{a}_text", "g1_same_{a}_text")
_TRAFFIC = ("slots_wire", "relations", "fact_refs", "req_bytes", "resp_bytes", "plan_input",
            "plan_output", "draft_input", "draft_output")


def _folded_numbers() -> M.Numbers:
    nb = M.Numbers()
    keys = [k.format(a=a, L=L) for a in M.ARM_ORDER for L in "AC"
            for k in (*_FORMER_T3_ARM, *_FORMER_T3_CORPUS, *_FORMER_T4, *_RUN_KEYS)]
    keys += [f"t4_{L}_oracle_{m}" for L in "AC" for m in ("runs", "rba", "attr", "covcorrect",
                                                          "covstated")]
    keys += ["table_cells_A", "table_cells_C", "min_rba_pool", "g1_stub_corpora",
             "h3_outcome", "h3_reason", "h3_runs_not_scanned_list", "n_registered_repeats"]
    for k in dict.fromkeys(keys):
        nb.put(k, "none" if k == "g1_stub_corpora" else "3" if k.endswith(("_runs", "repeats"))
               else "7", "x")
    return nb


def _spec_keys(spec: M.TableSpec) -> set[str]:
    cells = [c for r in spec.rows for c in (r[1:] if isinstance(r, tuple) else r)]
    return set(M.PLACEHOLDER_RE.findall(" ".join([*cells, *spec.notes, spec.caption])))


def test_table2_keeps_the_crossing_probe_and_store_keys_and_leaves_the_scores_to_fig4():
    """Table 2 (formerly Table 3) since 2026-10-19: (a) one row per arm, A-D
    pooled, each measurement a column; (b) one row per corpus, the shared
    extraction and r1's store, a second table under the caption. No per-arm score
    (Fig. 4 draws them); left for the artifact, their keys kept: the label-word
    counts, the per-arm cost and pooled pruning, the all-extraction echo ranges;
    the default leaves out the traffic figures (plan Part D AD32)."""
    nb = _folded_numbers()
    spec = M.table3(nb, {"A": "task-a", "C": "task-c"}, ["A", "C"], traffic=True)
    assert (spec.bid, spec.label) == ("t2", "Table 2")
    assert spec.header == list(M.T3_ARM_HEAD) and len(spec.widths) == len(spec.header)
    assert spec.widths2 == list(M.T3_CORPUS_WIDTHS)
    got = _spec_keys(spec)
    want = {k.format(a=a) for a in M.ARM_ORDER
            for k in ("obs_{a}_slots_adv", "cross_a_{a}", "cross_a_label_{a}",
                      "cross_a_scanned_{a}",
                      "cross_b_core_{a}", "cross_b_cloud_{a}", "cross_b_scanned_{a}",
                      "g3_entities_{a}", "g3_5grams_{a}")}
    want |= {k.format(L=L) for L in "AC"
             for k in ("pred_max_tokens_{L}", "recall_referents_exact_{L}",
                       "extraction_entity_jaccard_{L}", "extraction_relation_jaccard_{L}",
                       "table_cells_{L}", "r1_store_cells_{L}", "t4_{L}_r1_pruned")}
    assert want <= got, sorted(want - got)
    gone = {f"t4_{L}_{a}_{m}" for L in "AC" for a in (*M.ARM_ORDER, "oracle")
            for m in ("rba", "attr", "covcorrect", "covstated")}
    gone |= {f"obs_{a}_{m}" for a in M.ARM_ORDER for m in ("usd", "cost_complete", "pruned")}
    gone |= {f"pred_echo_pct_{L}" for L in "AC"} | {f"pred_control_echo_pct_{L}" for L in "AC"}
    assert not gone & got, sorted(gone & got)
    runs = {k.format(a=a, L=L) for a in M.ARM_ORDER for L in "AC" for k in _RUN_KEYS}
    assert not runs & got and runs <= set(nb.values)
    head = next(r for r in spec.rows if isinstance(r, tuple) and r[0] == "head")
    assert list(head[1:]) == list(M.T3_CORPUS_HEAD)
    arm_rows = [r for r in spec.rows[:spec.rows.index(head)] if isinstance(r, list)]
    assert [r[0] for r in arm_rows] == list(M.ARM_ORDER)
    assert all(len(r) == len(spec.header) for r in arm_rows)
    corpus = [r for r in spec.rows[spec.rows.index(head) + 1:] if isinstance(r, list)]
    assert [r[0] for r in corpus] == ["A*", "C"]
    assert all(len(r) == len(M.T3_CORPUS_HEAD) for r in corpus)
    lean = _spec_keys(M.table3(nb, {"A": "task-a", "C": "task-c"}, ["A", "C"]))
    assert got - lean == {f"obs_{a}_{t}" for a in M.ARM_ORDER for t in _TRAFFIC}
    assert all(f"obs_{a}_{t}" in nb.values for a in M.ARM_ORDER for t in _TRAFFIC)


def test_three_tables_three_figures_and_their_markers_are_in_the_skeleton():
    """Since 2026-10-19 (the author's requests) three tables (the monitor's
    refusals, what crosses, the hypotheses) and three figures (the example figure,
    which merged the architecture and mechanism figures; the wire listing; the
    binding figure), each with its own marker and description in the skeleton."""
    # Since later on 2026-10-21 (the author's request for a short evaluation): one
    # table (the hypotheses) and three figures (the example, the listing, the
    # results).
    assert M.BLOCKS == {"t1": "[TABLE 1 PLACEHOLDER]",
                        **{f"fig{n}": f"[FIGURE {n} PLACEHOLDER]" for n in (1, 2, 3)}}
    assert M.BLOCK_DESCRIPTIONS["t1"] == "Hypotheses: registered rule and outcome."
    assert C.PAPER_TABLES == 1 and C.PAPER_FIGURES == 3
    skel = _REPO / "paper" / "PrivDeep_ICCIT2026.docx"
    if docx is None or not skel.exists():
        pytest.skip("python-docx or the skeleton is missing")
    texts = [p.text for p in docx.Document(str(skel)).paragraphs]
    markers = [t for t in texts if C.BLOCK_PLACEHOLDER.match(t)]
    for bid, marker in M.BLOCKS.items():
        assert [t for t in markers if t.startswith(marker)] == [
            f"{marker} {M.BLOCK_DESCRIPTIONS[bid]}"], marker
    assert not any(t.startswith("[TABLE 2") for t in markers)
    # The skeleton's markers and the checker's labels agree: Tables 1 to
    # PAPER_TABLES and Figs 1 to PAPER_FIGURES, and no unnumbered table since the
    # calibration table became a sentence (plan Part D AD32).
    tables = sorted(int(m) for t in markers for m in re.findall(r"^\[TABLE (\d+) ", t))
    figures = sorted(int(m) for t in markers for m in re.findall(r"^\[FIGURE (\d+) ", t))
    assert tables == list(range(1, C.PAPER_TABLES + 1))
    assert figures == list(range(1, C.PAPER_FIGURES + 1))
    assert len(markers) == len(tables) + len(figures)
    # Every marker is a block the fill inserts, Fig. 1 included since it is drawn.
    assert set(M.BLOCKS.values()) == {re.match(r"^\[[^\]]*\]", t).group(0) for t in markers}


@needs_docx
def test_a_head_row_is_bold_with_a_rule_under_it(tmp_path):
    src = tmp_path / "s.docx"
    docx.Document().save(str(src))
    spec = M.TableSpec("t3", "Table 3", "Cap.", ["Arm", "RBA"],
                       [["r1", "«alpha»"], ("head", "Arm", "Cost"), ["r1", "«count»"]])
    M.write_tables_docx(src, [spec], tmp_path / "t.docx", VALUES)
    t = docx.Document(str(tmp_path / "t.docx")).tables[0]
    row = t.rows[2]
    assert [c.text for c in row.cells] == ["Arm", "Cost"]
    assert all(r.bold for c in row.cells for r in c.paragraphs[0].runs)
    assert all("w:bottom" in c._tc.xml for c in row.cells)
    with pytest.raises(ValueError, match="head row of 1 cells"):
        M.write_tables_docx(src, [M.TableSpec("t3", "Table 3", "C.", ["A", "B"],
                                              [("head", "A")])], tmp_path / "u.docx", VALUES)


# ------------------------------------------------------------ Table 1

def test_table1_condenses_every_property_row_of_the_threat_model():
    spec = M.table1_from_threat_model(_REPO / "THREAT_MODEL.md")
    rows = [r for r in spec.rows if isinstance(r, list)]
    ids = [r[0] for r in rows]
    # The MX6 row shows no id: MX6 is a plan identifier the paper never defines
    # (review 2026-10-09, T1-mx6-label).
    assert ids == [f"P{i}" for i in range(1, 10)] + ["", "G2"]
    assert all(len(r) == len(spec.header) for r in rows)
    # The caption names no repository file: "§4" there read as the paper's §4.
    assert "THREAT_MODEL" not in spec.caption and "§" not in spec.caption
    # No figure from the threat model's grade text reaches the table: its only
    # numerals are P2's threshold and shingle lengths, structural constants
    # ("5-word", "5-8-gram").
    paras = [C.Para(cell, "Table Paragraph", "t1", in_table=True)
             for r in rows for cell in r]
    res = C.check_paragraphs(paras, {})
    assert res.fails == [] and res.numerals == 3
    assert [n for r in rows for c in r for n, _, _ in C.numerals(c)] == ["5", "5", "8"]


def test_every_property_row_has_a_mechanism_in_words(tmp_path):
    """Change set 2026-10-10, TAB-I: the Mechanism column is a short gloss per
    P-row (T1_MECHANISM), not module names; every P-row of THREAT_MODEL.md §4 has
    one, and a P-row without one stops the build."""
    tm = (_REPO / "THREAT_MODEL.md").read_text(encoding="utf-8")
    pids = re.findall(r"^\|\s*(P\d+)\s*\|", tm, re.M)
    assert pids and set(pids) == set(M.T1_MECHANISM)
    spec = M.table1_from_threat_model(_REPO / "THREAT_MODEL.md")
    mechs = {r[0]: r[2] for r in spec.rows if isinstance(r, list)}
    for pid, words in M.T1_MECHANISM.items():
        assert mechs[pid] == words
        assert not re.search(r"\.py\b|/|\b\w+\.\w+\b", words), words
    assert mechs[""] == "Scripted drafter in the running r1 pipeline"
    assert mechs["G2"] == "Payload of typed relations, slot ids and coarse statistics"
    props = {r[0]: r[1] for r in spec.rows if isinstance(r, list)}
    # The threat model's internal names are in the paper's terms.
    assert props["P1"] == "Nothing on T's internal network can reach the internet"
    # P2's threshold in the text's form, "5-word", which the numeral check admits.
    assert props["P2"].startswith("No contiguous 5-word or longer run of corpus words leaves")
    assert "C_user" not in props["P2"] and "five" not in props["P2"]
    assert not any("privdeep-internal" in c or "C_user" in c
                   for r in spec.rows if isinstance(r, list) for c in r)
    # A new P-row without a mechanism stops the table.
    extra = tm.replace("| P9 |", "| P10 | A new property | `x.py` | **Built** |\n| P9 |", 1)
    assert extra != tm
    moved = tmp_path / "THREAT_MODEL.md"
    moved.write_text(extra, encoding="utf-8")
    with pytest.raises(SystemExit, match="P10 has no mechanism"):
        M.table1_from_threat_model(moved)


def test_every_property_row_is_pinned_by_a_test_the_document_cites():
    """PAP5b: 'Pinned by' was a dash for P3, P5, P6, P7 and P9 because §4 cited no
    test for them; it is read from the document, so the document now cites them.
    Since plan Part D AD32 each file is named by its stem, and each stem names
    exactly one existing file in the repository."""
    spec = M.table1_from_threat_model(_REPO / "THREAT_MODEL.md")
    pinned = {r[0]: r[4] for r in spec.rows if isinstance(r, list)}
    assert "—" not in pinned.values()
    assert "test_canary_matcher" in pinned["P3"]
    assert "test_pipeline_exfil" in pinned["P5"] and "test_slot_reintegrator" in pinned["P5"]
    assert pinned["P6"] == "test_threat_model_claims"
    # Every fact of a cell stays; only the paths became stems.
    assert pinned["P7"] == "test_citation_grounder (unit tests; not exercised)"
    assert pinned["P9"] == "test_fact_store; test_r1_facts"
    assert pinned[""] == "test_pipeline_exfil"                      # the MX6 row
    assert pinned["G2"] == "test_payload_invariance"
    assert pinned["P1"] == "verify_partition"
    assert pinned["P2"] == "test_shingle_filter; test_monitor_bypasses"
    assert "test_proxy" in pinned["P3"]
    assert not any("/" in c or re.search(r"\.(?:py|sh)\b", c) for c in pinned.values())
    index = M._stem_index(_REPO)
    stems = {w for c in pinned.values() for w in re.findall(r"\b[a-z]\w*_\w+\b", c)}
    assert stems and all(len(index[w]) == 1 for w in stems), {w: index.get(w) for w in stems}
    # A measurement script records a rate; it does not pin the property.
    assert not any(index[w][0].startswith("scripts/") and index[w][0].endswith(".py")
                   for w in stems)
    for w in stems:
        assert (_REPO / index[w][0]).exists()
    assert "file stem" in " ".join(spec.notes)
    props = {r[0]: r[1] for r in spec.rows if isinstance(r, list)}
    # THREAT_MODEL §2's caveat, in the change set's words (TAB-I); the build still
    # checks that §2 states "Quantisation is not invariance".
    assert "unless the rounded character count changes" in props["G2"]
    grades = {r[0]: r[3] for r in spec.rows if isinstance(r, list)}
    assert grades[""] == "Defence in depth, not a claim about the evaluated runs"
    # The MX6 row names itself as the test Sections III and VI-A cite, and no
    # row gives a repository path as its mechanism (review 2026-10-09, T1-mx6-label).
    assert props[""].startswith("Adversarial-drafter test: ")
    assert "second threat source" not in props[""]
    mechs = {r[0]: r[2] for r in spec.rows if isinstance(r, list)}
    assert not any("/" in m or ".py" in m for m in mechs.values())


def test_a_one_word_grade_keeps_the_qualifier_that_bounds_it():
    """Review 2026-10-09, T1-grades-truncated: the bold lead alone dropped P3's
    'for the exact token' and P9's 'given the extraction'. A longer lead, or one
    that ends its sentence, stands alone."""
    spec = M.table1_from_threat_model(_REPO / "THREAT_MODEL.md")
    grades = {r[0]: r[3] for r in spec.rows if isinstance(r, list)}
    assert grades["P3"] == "Deterministic, for the exact token"
    assert grades["P9"] == "Deterministic given the extraction"
    assert grades["P1"] == "Structural"
    assert grades["P2"] == "Deterministic for verbatim spans inside one stream"
    assert grades["P8"] == 'Deterministic for "no audit line, no delivery"'
    assert M._grade("**Built.** More text: here") == "Built"
    assert M._grade("**Exact**; then a clause") == "Exact"
    assert M._grade("no bold lead: here") == "no bold lead"


def test_the_stems_table1_names_are_unique_in_the_repository():
    """A stem stands for its file only while no other .py or .sh file shares it:
    the build refuses a duplicate stem, so a future test file with the same name
    fails it instead of making Table 1 ambiguous."""
    spec = M.table1_from_threat_model(_REPO / "THREAT_MODEL.md")
    index = M._stem_index(_REPO)
    cells = " ".join(r[4] for r in spec.rows if isinstance(r, list))
    named = re.findall(r"\b(?:test_\w+|verify_partition)\b", cells)
    assert len(named) >= 11
    for stem in named:
        assert len(index[stem]) == 1, (stem, index[stem])
    # The index skips hidden directories (a checkout's worktrees, .git, .venv).
    assert not any(p.startswith(".") or "/." in p for ps in index.values() for p in ps)
    dup = {**index, "test_proxy": [*index["test_proxy"], "backend/tests/unit/test_proxy.py"]}
    with pytest.raises(SystemExit, match="test_proxy"):
        M._as_stems("gateway/tests/test_proxy.py", dup)
    with pytest.raises(SystemExit, match="none"):
        M._as_stems("gateway/tests/test_no_such_file.py", index)
    assert M._as_stems("gateway/tests/test_proxy.py (unit tests; not exercised)", index) == \
        "test_proxy (unit tests; not exercised)"


def test_a_duplicate_stem_in_the_tree_stops_table1(tmp_path):
    for d in ("gateway/tests", "backend/tests", ".hidden/w/gateway", "build", "eval"):
        (tmp_path / d).mkdir(parents=True)
    for f in ("gateway/tests/test_proxy.py", "backend/tests/test_proxy.py",
              ".hidden/w/gateway/test_other.py", "build/test_other.py", "eval/test_other.py"):
        (tmp_path / f).write_text("", encoding="utf-8")
    index = M._stem_index(tmp_path)
    # Hidden directories and the top-level build and data directories are not searched.
    assert index == {"test_proxy": ["backend/tests/test_proxy.py",
                                    "gateway/tests/test_proxy.py"]}
    with pytest.raises(SystemExit, match="exactly one"):
        M._as_stems("gateway/tests/test_proxy.py", index)


def test_table1_refuses_when_the_threat_model_stops_stating_a_test(tmp_path):
    tm = (_REPO / "THREAT_MODEL.md").read_text(encoding="utf-8")
    moved = tmp_path / "THREAT_MODEL.md"
    moved.write_text(tm.replace("MX6 is defence in depth", "MX6 is a test"), encoding="utf-8")
    with pytest.raises(SystemExit, match="MX6"):
        M.table1_from_threat_model(moved)
    # The G2 row quotes §2's quantisation caveat, so it stops if §2 drops it.
    moved.write_text(tm.replace("Quantisation is not\n  invariance", "Quantisation is")
                     .replace("Quantisation is not invariance", "Quantisation is"),
                     encoding="utf-8")
    with pytest.raises(SystemExit, match="G2"):
        M.table1_from_threat_model(moved)


# ------------------------------------------------------------ Fig. 2

def _run_dir(tmp_path: Path, sentence: str, resolved: str) -> Path:
    run = tmp_path / "task-a__r1"
    run.mkdir()
    req = ("SLOTS:\n- ⟦SLOT:PROJECT_0a1b2c3d⟧  [PROJECT]  facts: ⟦FACT:0a1b2c3d.<ATTR>⟧\n"
           "- ⟦SLOT:PERSON_4e5f6a7b⟧  [PERSON]  facts: ⟦FACT:4e5f6a7b.<ATTR>⟧\n\n"
           "ATTRIBUTES (format ⟦FACT:<referent hex8>.<ATTR>⟧):\n"
           "- AUTHAMT: amount authorised (budget authority)\n"
           "- SPENTAMT: amount spent or committed\n- DATESTRT: start date\n\nEND")
    trace = [{"stage": "plan", "messages": [{"content": "plan"}], "response_text": ""},
             {"stage": "draft:s1", "messages": [{"content": req}], "refused": None,
              "response_text": "Intro sentence here. " + sentence + " Closing words."}]
    (run / "egress_trace.jsonl").write_text(
        "\n".join(json.dumps(e, ensure_ascii=False) for e in trace), encoding="utf-8")
    (run / "report.md").write_text("# R\n\n## S\n\nIntro sentence here. " + resolved
                                   + " Closing words.\n", encoding="utf-8")
    return run


SENT = ("⟦SLOT:PROJECT_0a1b2c3d⟧ holds ⟦FACT:0a1b2c3d.AUTHAMT⟧ and has spent "
        "⟦FACT:0a1b2c3d.SPENTAMT⟧.")
RESOLVED = "Project Heron holds 410,000.00 USD and has spent 63.0%."


def test_wire_excerpt_takes_the_resolved_sentence_and_checks_the_wire(tmp_path):
    run = _run_dir(tmp_path, SENT, RESOLVED)
    ex = M.wire_excerpt(run, surfaces=["410,000.00", "Heron"], entities=["Project Heron"])
    assert ex.cloud_sentence == SENT
    assert ex.resolution == RESOLVED            # the decimal is not cut at its point
    assert "- AUTHAMT: amount authorised (budget authority)" in ex.sent_lines
    assert not any("DATESTRT" in ln for ln in ex.sent_lines)   # unused codes elided
    assert not any("PERSON_4e5f6a7b" in ln for ln in ex.sent_lines)
    text = ex.render()
    assert "RECEIVED" in text and "LOCAL" in text


def test_the_full_listing_is_the_excerpt_without_its_header_and_each_line_is_a_key(tmp_path):
    ex = M.wire_excerpt(_run_dir(tmp_path, SENT, RESOLVED), surfaces=["Heron"],
                        entities=["Project Heron"])
    listing = ex.listing()
    assert not any(ln.startswith("#") for ln in listing.splitlines())
    assert ex.render().endswith(listing)            # the excerpt is header + listing
    assert ex.render().startswith("# Fig. 2 wire excerpt")
    nb = M.Numbers()
    M.numbers_from_fig2(nb, ex, full=True)
    lines = [ln for ln in listing.splitlines() if ln]
    assert nb["fig2_lines"] == str(len(lines))
    assert [nb[f"fig2_listing_{i:02d}"] for i in range(1, len(lines) + 1)] == lines
    assert nb["fig2_listing_01"].startswith("SENT (drafter request")
    assert nb["fig2_run"] == "task-a__r1" and nb["fig2_exchange"] == "1"
    assert nb["fig2_surfaces_checked"] == "1" and nb["fig2_entities_checked"] == "1"
    # Every key of the caption template is produced.
    assert set(M.PLACEHOLDER_RE.findall(M.FIG2_CAPTION_FULL)) <= set(nb.values)
    assert M.fig2_caption(ex, full=True) == M.FIG2_CAPTION_FULL


# The drafter request's ATTRIBUTES line as section_drafter._attribute_block
# writes it: longer than one line of the listing.
LONG_ATTR = ("ATTRIBUTES (any placeholder of category PERSON, ORG, PROJECT, LOCATION may have "
             "these; format ⟦FACT:<referent hex8>.<ATTR>⟧):")
SLOT_LINES = ["- ⟦SLOT:PROJECT_0a1b2c3d⟧  [PROJECT]  facts: ⟦FACT:0a1b2c3d.<ATTR>⟧",
              "- ⟦SLOT:PERSON_4e5f6a7b⟧  [PERSON]  facts: ⟦FACT:4e5f6a7b.<ATTR>⟧"]
CODE_LINES = ["- AUTHAMT: amount authorised (budget authority)",
              "- SPENTAMT: amount spent or committed"]
# Too long for the listing: with its resolution it would take more than the
# four rendered lines left beside the ATTRIBUTES and slot lines.
LONG_SENT = ("By the end of the reporting period the programme lead ⟦SLOT:PERSON_4e5f6a7b⟧ "
             "confirmed that ⟦SLOT:PROJECT_0a1b2c3d⟧ holds an authority of "
             "⟦FACT:0a1b2c3d.AUTHAMT⟧ and has already spent ⟦FACT:0a1b2c3d.SPENTAMT⟧ of it "
             "on the works that the steering group approved at its spring meeting.")
LONG_RESOLVED = ("By the end of the reporting period the programme lead Alice Brown confirmed "
                 "that Project Heron holds an authority of 410,000.00 USD and has already "
                 "spent 63.0% of it on the works that the steering group approved at its "
                 "spring meeting.")


def _excerpt(sentence: str, resolved: str, attr: str = LONG_ATTR) -> M.Excerpt:
    m = M._resolution_match(resolved, sentence)
    assert m is not None and m.group(0) == resolved
    return M.Excerpt("task-a__r1", 1, "draft:s1", [], sentence, resolved, 1, 1, {},
                     slot_lines=SLOT_LINES, attr_line=attr, code_lines=CODE_LINES,
                     value_ends=[m.end(g) - m.start() for g in range(1, m.re.groups + 1)])


def test_the_fallback_listing_has_six_lines_that_show_the_three_things_verbatim():
    """The writing fallback's Fig. 2: at most six rendered lines, showing the
    ATTRIBUTES line, the cloud's sentence with its fact references and the local
    resolution, each verbatim up to a visible ellipsis."""
    ex = _excerpt(SENT, RESOLVED)
    short = ex.short()
    lines = short["lines"]
    assert len(lines) <= M.FIG2_SHORT_LINES and short["rows"] <= M.FIG2_SHORT_LINES
    assert lines == [
        "SENT     " + SLOT_LINES[0],                  # the referent's line, its fact stem
        "         ATTRIBUTES (any placeholder of category PERSON, ORG, PROJECT, LOCATION "
        "may have these; format …",
        "         " + CODE_LINES[0],
        "         " + CODE_LINES[1],
        "RECEIVED " + SENT,
        "LOCAL    " + RESOLVED]
    # Every line fits the listing's width; a cut one is a prefix of its source.
    assert all(len(ln) <= M.FIG2_WIDTH for ln in lines)
    assert LONG_ATTR.startswith(lines[1].strip()[:-len(M.FIG2_CUT)])
    assert short["cut"] and short["slot"] and short["codes"] == 2
    cap = M.fig2_caption(ex)
    # A title since 2026-10-19 (the author's request for short captions); the
    # listing's labels name its parts, and §4 states its check.
    # The cut-line convention is §4's text since the review of 2026-10-19.
    assert cap == M.FIG2_CAPTION_SHORT and "is cut" not in cap
    # Nothing to cut: the caption does not say that anything is.
    assert "is cut" not in M.fig2_caption(_excerpt(SENT, RESOLVED, attr="ATTRIBUTES (x):"))


def test_a_reply_too_long_for_the_listing_is_cut_after_the_same_token_in_both():
    ex = _excerpt(LONG_SENT, LONG_RESOLVED)
    assert not ex.reply_fits()
    short = ex.short()
    recv = next(ln for ln in short["lines"] if ln.startswith("RECEIVED"))[M.FIG2_LABEL_W:]
    local = next(ln for ln in short["lines"] if ln.startswith("LOCAL"))[M.FIG2_LABEL_W:]
    assert recv.endswith("⟧" + M.FIG2_CUT) and local.endswith(M.FIG2_CUT)
    recv, local = recv[:-len(M.FIG2_CUT)], local[:-len(M.FIG2_CUT)]
    assert LONG_SENT.startswith(recv) and LONG_RESOLVED.startswith(local)
    # Cut after the same reference, the latest at which both fit: the resolution
    # ends in that reference's value.
    assert recv.endswith("⟦FACT:0a1b2c3d.AUTHAMT⟧") and local.endswith("410,000.00 USD")
    assert short["rows"] <= M.FIG2_SHORT_LINES and short["cut"]


def test_the_excerpt_prefers_a_sentence_whose_reply_fits_the_listing(tmp_path):
    run = _run_dir(tmp_path, LONG_SENT + " " + SENT, LONG_RESOLVED + " " + RESOLVED)
    ex = M.wire_excerpt(run, surfaces=["Heron"], entities=["Project Heron"])
    assert ex.cloud_sentence == SENT and ex.reply_fits()
    assert ex.attr_line.startswith("ATTRIBUTES") and ex.code_lines == CODE_LINES
    # With only the long sentence there, it is taken, and cut.
    other = tmp_path / "other"
    other.mkdir()
    ex = M.wire_excerpt(_run_dir(other, LONG_SENT, LONG_RESOLVED), surfaces=["Heron"],
                        entities=["Project Heron"])
    assert ex.cloud_sentence == LONG_SENT and not ex.reply_fits()


# Three fact references: too long for the listing, and a later one that fits.
LONG3_SENT = LONG_SENT.replace("of it on", "of it since ⟦FACT:0a1b2c3d.DATESTRT⟧ on")
LONG3_RESOLVED = LONG_RESOLVED.replace("of it on", "of it since 2024-01-05 on")
SENT3 = ("⟦SLOT:PROJECT_0a1b2c3d⟧ holds ⟦FACT:0a1b2c3d.AUTHAMT⟧, has spent "
         "⟦FACT:0a1b2c3d.SPENTAMT⟧ and started ⟦FACT:0a1b2c3d.DATESTRT⟧.")
RESOLVED3 = "Project Heron holds 410,000.00 USD, has spent 63.0% and started 2024-01-05."


def test_the_preference_never_trades_fact_references_for_fit(tmp_path):
    """Review of b38664c: the preference took a later sentence with fewer
    references (a name and a place instead of amounts). It now keeps the first
    sentence's number of references, and a sentence with fewer is never taken
    over it."""
    run = _run_dir(tmp_path, LONG3_SENT + " " + SENT, LONG3_RESOLVED + " " + RESOLVED)
    ex = M.wire_excerpt(run, surfaces=["Heron"], entities=["Project Heron"])
    assert ex.cloud_sentence == LONG3_SENT and not ex.reply_fits()
    # A later sentence with as many references that fits is preferred.
    other = tmp_path / "other"
    other.mkdir()
    run = _run_dir(other, LONG3_SENT + " " + SENT + " " + SENT3,
                   LONG3_RESOLVED + " " + RESOLVED + " " + RESOLVED3)
    ex = M.wire_excerpt(run, surfaces=["Heron"], entities=["Project Heron"])
    assert ex.cloud_sentence == SENT3 and ex.reply_fits()


def test_the_full_listing_takes_the_first_sentence_as_before_the_fallback(tmp_path):
    """--fig2-full reproduces the listing of the build before the fallback: no
    preference for a sentence that fits the short listing."""
    run = _run_dir(tmp_path, LONG_SENT + " " + SENT, LONG_RESOLVED + " " + RESOLVED)
    ex = M.wire_excerpt(run, surfaces=["Heron"], entities=["Project Heron"],
                        prefer_fit=False)
    assert ex.cloud_sentence == LONG_SENT
    assert M.wire_excerpt(run, surfaces=["Heron"], entities=["Project Heron"]
                          ).cloud_sentence == SENT


def test_the_caption_says_when_the_listing_shows_fewer_code_lines_than_the_sentence_uses():
    """Review of b38664c: with three codes and room for one, the caption read
    'the code line the sentence uses'. It now gives both counts, as keys."""
    m = M._resolution_match(RESOLVED3, SENT3)
    codes = [*CODE_LINES, "- DATESTRT: start date"]
    ex = M.Excerpt("task-a__r1", 1, "draft:s1", [], SENT3, RESOLVED3, 1, 1, {},
                   slot_lines=SLOT_LINES, attr_line=LONG_ATTR, code_lines=codes,
                   value_ends=[m.end(g) - m.start() for g in range(1, m.re.groups + 1)])
    short = ex.short()
    assert ex.reply_fits() and short["rows"] <= M.FIG2_SHORT_LINES
    assert (short["codes"], short["codes_used"]) == (1, 3)
    assert sum(1 for ln in short["lines"] if ln.strip().startswith("- ") and ":" in ln
               and "⟦SLOT" not in ln) == 1
    cap = M.fig2_caption(ex)
    assert "«fig2_codes_shown» of the «fig2_codes_used» code lines the sentence uses" in cap
    assert "the code line the sentence uses" not in cap
    nb = M.Numbers()
    M.numbers_from_fig2(nb, ex)
    assert (nb["fig2_codes_shown"], nb["fig2_codes_used"]) == ("1", "3")
    assert set(M.PLACEHOLDER_RE.findall(cap)) <= set(nb.values)
    # Every code line shown: the title alone (since 2026-10-19), with no counts.
    full = _excerpt(SENT, RESOLVED)
    assert (full.short()["codes"], full.short()["codes_used"]) == (2, 2)
    assert M.fig2_caption(full).startswith(M.FIG2_CAPTION_SHORT)
    assert "fig2_codes_shown" not in M.fig2_caption(full)


def test_the_fallback_listing_is_the_default_and_each_line_is_a_key(tmp_path):
    ex = M.wire_excerpt(_run_dir(tmp_path, SENT, RESOLVED), surfaces=["Heron"],
                        entities=["Project Heron"])
    nb = M.Numbers()
    M.numbers_from_fig2(nb, ex)
    lines = ex.short_lines()
    assert nb["fig2_lines"] == str(len(lines)) and len(lines) <= M.FIG2_SHORT_LINES
    assert [nb[f"fig2_listing_{i:02d}"] for i in range(1, len(lines) + 1)] == lines
    assert set(M.PLACEHOLDER_RE.findall(M.fig2_caption(ex))) <= set(nb.values)
    # The resolved figures are inside a listing line, so a tagged value.
    assert nb[f"fig2_listing_{len(lines):02d}"] == "LOCAL    " + RESOLVED


def _fig2_out(tmp_path: Path, full: bool = False) -> tuple[Path, dict[str, str]]:
    ex = M.wire_excerpt(_run_dir(tmp_path, SENT, RESOLVED), surfaces=["Heron"],
                        entities=["Project Heron"])
    nb = M.Numbers()
    M.numbers_from_fig2(nb, ex, full=full)
    out = _out(tmp_path, nb.values)
    (out / "fig2_listing.txt").write_text("\n".join(M.fig2_lines_of(ex, full)) + "\n",
                                          encoding="utf-8")
    (out / "fig2_caption.txt").write_text(M.fig2_caption(ex, full) + "\n", encoding="utf-8")
    return out, nb.values


@needs_docx
def test_fill_pastes_the_listing_as_a_tagged_monospaced_block(tmp_path):
    out, values = _fig2_out(tmp_path)
    d = docx.Document()
    d.add_paragraph("Before.")
    d.add_paragraph("[FIGURE 2 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["fig2"])
    d.add_paragraph("After.")
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    r = M.fill(src, dst, out)
    assert r["inserted"] == ["fig2"] and r["placeholders_left"] == []
    f = docx.Document(str(dst))
    texts = [p.text.replace('\xa0', ' ') for p in f.paragraphs]
    listing = (out / "fig2_listing.txt").read_text(encoding="utf-8").rstrip("\n").split("\n")
    assert texts[1:1 + len(listing)] == listing
    # The caption names the run's corpus, not its run directory (change set
    # 2026-10-10, FIG-2); without a corpus map, the task id stands for it.
    assert texts[1 + len(listing)].startswith("Fig. 2. Wire excerpt of arm r1 from a registered "
                                              "run on corpus task-a.")
    assert texts[-1] == "After." and not any("PLACEHOLDER" in t for t in texts)
    # Every listing line keeps with the next, the last with the caption below it, so a
    # page break cannot part the listing from its caption (2026-10-28: the caption
    # opened a page alone after the text above it shortened).
    assert all(p.paragraph_format.keep_with_next for p in f.paragraphs[1:1 + len(listing)])
    local = next(p for p in f.paragraphs if p.text == "LOCAL    " + RESOLVED)
    assert {run.font.name for run in local.runs} == {M.FIG2_FONT}
    # A wrapped line continues under the text, not under the label column; the
    # text keeps FIG2_BAND_PAD_PT from both column edges, so the band's borders
    # stay inside the margins.
    indent = M.FIG2_LABEL_W * M.FIG2_PT * M.FIG2_ADVANCE
    assert round(local.paragraph_format.left_indent.pt, 1) == round(
        indent + M.FIG2_BAND_PAD_PT, 1)
    assert round(local.paragraph_format.first_line_indent.pt, 1) == -round(indent, 1)
    assert round(local.paragraph_format.right_indent.pt, 1) == M.FIG2_BAND_PAD_PT
    # The resolved figures are tagged values, so the checker passes them; typed
    # as plain text they would fail.
    res = C.check(dst, out / "numbers.json")
    assert not [x for x in res.fails if x.kind in ("numeral", "stale-value", "unknown-value")]
    assert res.values >= len([ln for ln in listing if ln])


@needs_docx
def test_fill_pastes_the_full_listing_with_no_hanging_indent(tmp_path):
    out, _ = _fig2_out(tmp_path, full=True)
    d = docx.Document()
    d.add_paragraph("[FIGURE 2 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["fig2"])
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    M.fill(src, dst, out)
    f = docx.Document(str(dst))
    local = next(p for p in f.paragraphs if p.text == RESOLVED)
    assert local.paragraph_format.left_indent is None
    assert f.paragraphs[0].text.startswith("SENT (drafter request")
    assert f.paragraphs[-1].text.replace('\xa0', ' ') == M.FIG2_CAPTION_FULL.replace(
        "«fig2_corpus»", "task-a").replace(
        "«fig2_surfaces_checked»", "1").replace("«fig2_entities_checked»", "1")


@needs_docx
def test_fill_refuses_a_listing_from_another_build(tmp_path):
    out, values = _fig2_out(tmp_path)
    (out / "numbers.json").write_text(json.dumps(
        {**values, "fig2_listing_05": values["fig2_listing_05"] + " (edited)"}),
        encoding="utf-8")
    d = docx.Document()
    d.add_paragraph("[FIGURE 2 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["fig2"])
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    with pytest.raises(M.FillError, match="fig2_listing_05"):
        M.fill(src, dst, out)
    assert not dst.exists()


@needs_docx
def test_a_wide_row_spans_the_value_columns(tmp_path):
    src = tmp_path / "s.docx"
    docx.Document().save(str(src))
    spec = M.TableSpec("t3", "Table 3", "Cap.", ["", "baseline", "r5", "r1"],
                       [["Runs", "«count»", "«count»", "«count»"],
                        ("wide", "Predicate length", "A: «count»; C: n/a")])
    M.write_tables_docx(src, [spec], tmp_path / "t.docx", VALUES)
    t = docx.Document(str(tmp_path / "t.docx")).tables[0]
    last = t.rows[-1]
    assert [c.text for c in last.cells] == ["Predicate length"] + ["A: 41; C: n/a"] * 3
    assert last.cells[1]._tc is last.cells[3]._tc          # one merged cell


def test_wire_excerpt_refuses_when_the_wire_carries_a_surface(tmp_path):
    leaky = SENT.replace("holds", "at Leeds holds")
    run = _run_dir(tmp_path, leaky, RESOLVED.replace("holds", "at Leeds holds"))
    with pytest.raises(SystemExit, match="free of generator surfaces"):
        M.wire_excerpt(run, surfaces=["Leeds"], entities=[])


# ------------------------------------------------------------ numeral checker

def _para(text: str, *values: tuple[str, str], style: str = "Body Text",
          in_table: bool = False) -> C.Para:
    """A paragraph whose value spans are the first stand-alone occurrence of each
    (key, text)."""
    spans = []
    for key, val in values:
        a = re.search(r"(?<![\w.])" + re.escape(val) + r"(?![\w])", text).start()
        spans.append(C.Span(a, a + len(val), "value", f"{C.value_tag(key)}_{len(spans) + 1}"))
    return C.Para(text, style, "p", spans, in_table)


def test_numerals_are_not_identifiers_and_include_leading_decimals():
    got = [n for n, _, _ in C.numerals(
        "R5 and P1, sha256 8ca0d45e80a3, 5-grams, 7/41 runs, 12.3% and 0.41, 600dpi, F03, "
        "(.81), v1.2")]
    assert got == ["5", "7", "41", "12.3", "0.41", "600", ".81"]


def test_checker_passes_tagged_values_and_structural_numerals():
    numbers = {"alpha": "0.41 ± 0.03 [0.38–0.44]", "kn": "7/41"}
    paras = [
        C.Para("6.2 What crosses", "Heading 2", "p0"),
        _para("RBA 0.41 ± 0.03 [0.38–0.44]; refused 7/41 as in [1, 2].",
              ("alpha", "0.41 ± 0.03 [0.38–0.44]"), ("kn", "7/41")),
        C.Para("See Table 1, Fig. 2 and Section 6.2, on a 1-5 scale, 5-grams, 13 codes, "
               "10 pages, 600 dpi, Wilson 95% intervals, Law No. 91/2025/QH15, "
               "Decree 356/2025/ND-CP, Art. 4(5)/5(1)(c).",
               "Body Text", "p2"),
        C.Para("References", "Heading 2", "p3"),
        C.Para("[1] A. Author, Title. J. Phys. 106, 2429 (2012).", "Body Text", "p4"),
        C.Para("[2] B. Author, Other. 2019.", "Body Text", "p5"),
    ]
    res = C.check_paragraphs(paras, numbers)
    assert res.fails == [] and res.values == 2 and res.numerals > 0


def test_a_hand_typed_numeral_equal_to_a_numbers_value_fails():
    numbers = {"alpha": "0.41", "count": "41"}
    paras = [_para("R1 scores 0.41 over 41 runs.", ("alpha", "0.41"), ("count", "41")),
             C.Para("Typed by hand: 0.41 and n=41, over .41 of runs.", "Body Text", "p2")]
    res = C.check_paragraphs(paras, numbers)
    assert [(f.kind, f.numeral) for f in res.fails] == [
        ("numeral", "0.41"), ("numeral", "41"), ("numeral", ".41")]


def test_a_stale_or_unknown_value_span_fails():
    numbers = {"alpha": "0.42"}
    res = C.check_paragraphs([_para("R1 scores 0.41.", ("alpha", "0.41"))], numbers)
    assert [(f.kind, f.numeral) for f in res.fails] == [("stale-value", "0.41")]
    res = C.check_paragraphs([_para("R1 scores 0.41.", ("gone", "0.41"))], numbers)
    assert [f.kind for f in res.fails] == ["unknown-value"]


def test_a_negative_value_prints_with_the_minus_sign_and_still_checks():
    # S13(5) of the simulated review of 2026-10-26 (scope §13.2(c)): presentation only;
    # numbers.json keeps the ASCII hyphen-minus, the fill prints U+2212, the check compares the
    # printed form. A hyphen after a letter, digit, point or bracket is not a sign.
    assert C.presented("-0.01") == "−0.01"
    assert C.presented("[-0.07, 0.08]") == "[−0.07, 0.08]"
    for v in ("2026-10-07", "claude-sonnet-5", "0009-0009-2412-9566", "2.18–2.24", "5/22"):
        assert C.presented(v) == v
    numbers = {"k": "-0.01"}
    assert not C.check_paragraphs([_para("κ −0.01.", ("k", "−0.01"))], numbers).fails
    stale = C.check_paragraphs([_para("κ -0.01.", ("k", "-0.01"))], numbers).fails
    assert [f.kind for f in stale] == ["stale-value"]


@needs_docx
def test_fill_prints_a_negative_value_with_the_minus_sign_and_the_check_passes_it(tmp_path):
    # S13(5), 2026-10-26 (the AI assistant, at the author's request): the fill-level half of
    # the test above, on a numbers.json of its own.
    d = docx.Document()
    d.add_paragraph("The rule's held-out κ is «k», interval «ci».")
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d.save(str(src))
    out = _out(tmp_path, {"k": "-0.01", "ci": "[-0.07, 0.08]"})
    M.fill(src, dst, out)
    text = docx.Document(str(dst)).paragraphs[0].text
    assert text == "The rule's held-out κ is −0.01, interval [−0.07, 0.08]."
    assert C.check(dst, out / "numbers.json").fails == []


def test_structural_numerals_are_bound_to_their_context():
    paras = [C.Para("References", "Heading 2", "r"),
             C.Para("[1] One reference.", "Body Text", "r1")]
    body = [C.Para("Over 95% of runs, see [3–7] and Table 12, Fig. 9, Section 9, "
                   "with 13 runs, 5 corpora and 8 arms.", "Body Text", "p")]
    res = C.check_paragraphs(body + paras, {})
    assert [f.numeral for f in res.fails] == ["95", "3", "7", "12", "9", "9", "13", "5", "8"]
    ok = [C.Para("Wilson 95% CI; 95% confidence; [1]; Table 1; Fig. 2.", "Body Text", "p")]
    assert C.check_paragraphs(ok + paras, {}).fails == []
    # Three tables and three figures since the example figure merged the
    # architecture and mechanism figures (2026-10-19): a "Table 4" or "Fig. 4" in
    # the text is not a label of this paper.
    # Since later on 2026-10-21, one table and three figures.
    assert C.PAPER_TABLES == 1 and C.PAPER_FIGURES == 3
    ok3 = [C.Para("See Table 1, and Figs. 1-3.", "Body Text", "p")]
    assert C.check_paragraphs(ok3 + paras, {}).fails == []
    gone = [C.Para("See Table 2 and Tables 1-2, and Fig. 4.", "Body Text", "p")]
    assert [f.numeral for f in C.check_paragraphs(gone + paras, {}).fails] == [
        "2", "1", "2", "4"]


def test_the_statute_rule_allows_the_cited_instruments_by_their_full_numbers():
    ok = C.Para("No legal assessment under Law No. 91/2025/QH15, Law 91/2025/QH15, "
                "Law on Personal Data Protection No. 91/2025/QH15, Decree 356/2025/ND-CP, "
                "Decree No. 356/2025/NĐ-CP, Regulation (EU) 2016/679, Art. 4(5)/5(1)(c) "
                "or Article 5.", "Body Text", "p")
    assert C.check_paragraphs([ok], {}).fails == []


def test_a_number_year_pair_outside_a_cited_statute_fails():
    # The same digits pass only inside the instrument's full number: a bare fraction,
    # the dropped Decree 13/2023, another number or year, or the Law's number under
    # "Decree" are all hand-typed numerals. After "Art." only the article number
    # passes: a year joined to it by a slash is not a paragraph-qualified article.
    bad = C.Para("Over 91/2025 runs; Decree 13/2023/ND-CP; Law No. 92/2025/QH15; "
                 "Decree 356/2024/ND-CP; Decree 91/2025/QH15; Law 91/2025; "
                 "Art. 13/2023; Article 91/2025.", "Body Text", "p")
    assert [f.numeral for f in C.check_paragraphs([bad], {}).fails] == [
        "91", "2025", "13", "2023", "92", "2025", "356", "2024", "91", "2025", "91", "2025",
        "2023", "2025"]


def test_na_in_prose_fails_and_in_a_table_passes():
    numbers = {"alpha": "n/a", "count": "7"}
    res = C.check_paragraphs([_para("R1 raises it to n/a over 7.", ("alpha", "n/a"),
                                    ("count", "7"))], numbers)
    assert [f.kind for f in res.fails] == ["na-in-prose"]
    res = C.check_paragraphs([_para("n/a over 7", ("alpha", "n/a"), ("count", "7"),
                                    in_table=True)], numbers)
    assert res.fails == []


def test_a_document_with_no_numeral_fails():
    res = C.check_paragraphs([C.Para("No figures here.", "Body Text", "p")], {})
    assert [f.kind for f in res.fails] == ["empty"]


def test_number_words_are_warnings():
    res = C.check_paragraphs([C.Para("Over three corpora, 1-5 scale.", "Body Text", "p")], {})
    assert res.fails == [] and [f.numeral for f in res.warns] == ["three"]


def test_edit_notes_fail_unless_allowed_even_in_references():
    paras = [C.Para("[EDIT: pilot figures 0.12 and 0.34]", "Body Text", "p1"),
             C.Para("References", "Heading 2", "p2"),
             C.Para("[EDIT: Page budget: 0.5 + 0.7 = 1.2 pages]", "Body Text", "p3")]
    res = C.check_paragraphs(paras, {})
    assert [f.numeral for f in res.fails] == ["0.12", "0.34", "0.5", "0.7", "1.2"]
    assert res.warns == []
    res = C.check_paragraphs(paras, {}, allow_edit_notes=True)
    assert res.fails == [] and len(res.warns) == 5


@needs_docx
def test_checker_reads_a_filled_docx(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    _doc_with_split_placeholder(src)
    out = _out(tmp_path, VALUES)
    M.fill(src, dst, out)
    # Nothing was inserted for the two block placeholders: failures when strict,
    # warnings while drafting; their "2" and "3" are labels, not numerals.
    res = C.check(dst, out / "numbers.json")
    assert [f.kind for f in res.fails] == ["block-placeholder", "block-placeholder"]
    assert res.values == 4
    res = C.check(dst, out / "numbers.json", allow_edit_notes=True)
    assert res.fails == [] and sorted(f.kind for f in res.warns) == [
        "block-placeholder", "block-placeholder"]
    # A figure typed by hand fails, even one equal to a numbers.json value.
    d = docx.Document(str(dst))
    d.add_paragraph("Typed by hand: 0.41 and 0.889.")
    d.save(str(dst))
    res = C.check(dst, out / "numbers.json", allow_edit_notes=True)
    assert [f.numeral for f in res.fails] == ["0.41", "0.889"]
    # A numbers.json that moved on makes every earlier value stale.
    (out / "numbers.json").write_text(json.dumps({**VALUES, "alpha": "0.42"}),
                                      encoding="utf-8")
    res = C.check(dst, out / "numbers.json", allow_edit_notes=True)
    assert ("stale-value", "0.41") in [(f.kind, f.numeral) for f in res.fails]


@needs_docx
def test_checker_reads_footnotes_and_skips_page_fields(tmp_path):
    d = docx.Document()
    d.add_paragraph("Body with Table 1.")
    path = tmp_path / "f.docx"
    d.save(str(path))
    _add_zip_entry(path, "word/footnotes.xml", _FOOTNOTES.format("A typed 0.41."))
    footer = ('<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
              '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGE '
              '</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
              '<w:r><w:t>11</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
              '</w:ftr>')
    _add_zip_entry(path, "word/footer9.xml", footer)
    res = C.check_paragraphs(C.docx_paragraphs(path), {})
    assert [(f.numeral, f.where) for f in res.fails] == [("0.41", "footnotes.xml:p0")]


@needs_docx
def test_checker_reads_a_non_breaking_hyphen_as_a_hyphen(tmp_path):
    """Word writes a non-breaking hyphen as <w:noBreakHyphen/>, not as text; read
    as nothing, "5-grams" would become "5grams" and fail as a stray numeral."""
    from docx.oxml import OxmlElement
    d = docx.Document()
    r = d.add_paragraph().add_run("Word 5")
    r._r.append(OxmlElement("w:noBreakHyphen"))
    t = OxmlElement("w:t")
    t.text = "grams and a typed 0.41."
    r._r.append(t)
    path = tmp_path / "h.docx"
    d.save(str(path))
    paras = C.docx_paragraphs(path)
    assert paras[0].text == "Word 5-grams and a typed 0.41."
    assert [f.numeral for f in C.check_paragraphs(paras, {}).fails] == ["0.41"]


def test_structural_constants_match_the_code():
    gw = (_REPO / "gateway" / "main.py").read_text(encoding="utf-8")
    assert re.search(r'SHINGLE_N_MIN", "5"', gw) and "5" in C.STRUCTURAL_CONSTANTS
    assert re.search(r'SHINGLE_N_MAX", "8"', gw) and "8" in C.STRUCTURAL_CONSTANTS
    assert re.search(r"^_CARRY_MAX_PATHS = 64$", gw, re.M)
    from privdeep.models.facts import ATTRIBUTE_CODES
    assert str(len(ATTRIBUTE_CODES)) == "13" and "13" in C.STRUCTURAL_CONSTANTS
    from scripts.real_corpora import tab as T
    assert (T.MIN_ANNOTATORS, T.DOCS_PER_COLLECTION, T.MAX_TEXT_CHARS) == (4, 4, 8000)
    assert "4" in C.STRUCTURAL_CONSTANTS and "8,000" in C.STRUCTURAL_CONSTANTS
    batch = (_REPO / "scripts" / "freeze_batch_real.sh").read_text(encoding="utf-8")
    assert re.search(r"^CAP_USD=15$", batch, re.M) and "15" in C.STRUCTURAL_CONSTANTS
    for n, const in C.STRUCTURAL_CONSTANTS.items():
        assert n in re.findall(r"\d+(?:,\d{3})*", const.pattern.pattern) and const.source
    venue = (_REPO / "paper" / "VENUE_FACTS.md")
    if venue.exists():   # paper/ is not yet committed on every branch
        v = venue.read_text(encoding="utf-8")
        assert "10 pages" in v and "600 dpi" in v and "150 words" in v


# ------------------------------------------------------------ the strip plot (artifact)

def test_fig3_is_written_at_600_dpi_with_an_empty_corpus_drawn(tmp_path):
    # Fig. 3 until the author removed it from the manuscript (plan Part D AD33);
    # the build still draws it for the artifact.
    pytest.importorskip("matplotlib")
    data = {"A": {"task": "task-a", "oracle": (0.55, 0.66),
                  "points": {"baseline": [("a", 0.21)], "r5": [("b", None)],
                             "r1": [("c", 0.61), ("d", 1.0)]}},
            "B": {"task": "task-b", "oracle": None,
                  "points": {"baseline": [], "r5": [], "r1": []}}}
    png, svg = M.make_fig3(data, tmp_path)
    text = svg.read_text(encoding="utf-8")
    assert text.lstrip().startswith("<?xml") and "no runs" in text
    from PIL import Image
    with Image.open(png) as im:
        assert round(im.info["dpi"][0]) == M.FIG_DPI == 600


# ------------------------------------------------------------ Fig. 1
# The concept figure the fill inserts (FIG1_PNG); then the tests of
# fig1_architecture.py, the former Fig. 1, kept in the repository for the artifact.

FIG1 = _REPO / "paper" / "figures"


def _fig1_module():
    from paper.figures import fig1_architecture
    return fig1_architecture


def _png_header(data: bytes) -> tuple[int, int, float]:
    """(width px, height px, dpi) of a PNG, from IHDR and pHYs; standard library only,
    so this runs in the venv, which has no Pillow."""
    import struct
    assert data[:8] == b"\x89PNG\r\n\x1a\n"
    w, h = struct.unpack(">II", data[16:24])
    pos, dpi = 8, 0.0
    while pos < len(data):
        n, kind = struct.unpack(">I4s", data[pos:pos + 8])
        if kind == b"pHYs":
            x, _y, unit = struct.unpack(">IIB", data[pos + 8:pos + 17])
            assert unit == 1                      # pixels per metre
            dpi = x * 0.0254
        pos += 12 + n
    return w, h, dpi


def _fig1_dir(tmp_path: Path, caption: str) -> Path:
    from paper.tools.page_estimate import _blank_png
    d = tmp_path / "figures"
    d.mkdir()
    (d / M.FIG1_PNG).write_bytes(_blank_png(30, 9))
    (d / M.FIG1_CAPTION).write_text(caption + "\n", encoding="utf-8")
    return d


def _fig1_doc(path: Path, heading: bool = False) -> None:
    d = docx.Document()
    if heading:                 # the section the caption points to, as in the manuscript
        d.add_heading("3 Threat Model", level=1)
    d.add_paragraph("Before, at «alpha».")
    d.add_paragraph("[FIGURE 1 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["fig1"])
    d.add_paragraph("After.")
    d.save(str(path))


@needs_docx
def test_fill_inserts_fig1_once_at_its_marker_with_its_caption(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    _fig1_doc(src)
    out = _out(tmp_path, VALUES)
    fig1 = _fig1_dir(tmp_path, "Fig. 1. Architecture, drawn.")
    r = M.fill(src, dst, out, fig1_dir=fig1)
    assert r["inserted"] == ["fig1"] and r["placeholders_left"] == [] and r["text_kept"] == []
    f = docx.Document(str(dst))
    # The marker and the skeleton's description are gone; the picture and the
    # caption stand where they were.
    assert [p.text.replace('\xa0', ' ') for p in f.paragraphs] == ["Before, at 0.41.", "",
                                              "Fig. 1. Architecture, drawn.", "After."]
    assert len(f.inline_shapes) == 1
    assert f.inline_shapes[0].width.inches == pytest.approx(M.FIG1_WIDTH_IN)
    # The picture is kept with its caption, and "Fig. 1." is bold, as Fig. 2's is.
    assert f.paragraphs[1].paragraph_format.keep_with_next
    assert [run.text.replace("\xa0", " ") for run in f.paragraphs[2].runs
            if run.bold] == ["Fig. 1."]
    # A refill of the filled file finds no marker and inserts nothing more.
    again = tmp_path / "again.docx"
    assert M.fill(dst, again, out, fig1_dir=fig1)["inserted"] == []
    assert len(docx.Document(str(again)).inline_shapes) == 1


@needs_docx
def test_fill_leaves_the_fig1_marker_until_the_figure_and_its_caption_exist(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    _fig1_doc(src)
    out = _out(tmp_path, VALUES)
    fig1 = _fig1_dir(tmp_path, "Fig. 1. Architecture.")
    (fig1 / M.FIG1_CAPTION).unlink()
    r = M.fill(src, dst, out, fig1_dir=fig1)
    assert r["inserted"] == [] and r["placeholders_left"] == ["[FIGURE 1 PLACEHOLDER]"]
    assert len(docx.Document(str(dst)).inline_shapes) == 0
    # A caption key without a value is refused like any other placeholder.
    (fig1 / M.FIG1_CAPTION).write_text("Fig. 1. Over «no_such_key» runs.\n", encoding="utf-8")
    with pytest.raises(M.FillError, match="no_such_key"):
        M.fill(src, tmp_path / "dst2.docx", out, fig1_dir=fig1)
    assert not (tmp_path / "dst2.docx").exists()


def test_no_break_space_joins_a_label_and_its_number():
    """Pre-submission audit of 2026-10-25: "(Fig." / "3(b)" and "(Table" / "I)" broke
    across lines; the fill puts a no-break space after "Fig.", "Table" and
    "Section" before a number or a Roman numeral, and nowhere else."""
    from docx.oxml.ns import qn  # noqa: PLC0415
    from lxml import etree  # noqa: PLC0415
    body = etree.Element(qn("w:body"))
    t = etree.SubElement(etree.SubElement(etree.SubElement(body, qn("w:p")), qn("w:r")),
                         qn("w:t"))
    t.text = "See Fig. 3(b), Table I and Section VI-B; the Table of a Section stays."
    assert M._nbsp_labels(body) == 3
    assert t.text == ("See Fig.\u00a03(b), Table\u00a0I and Section\u00a0VI-B; the Table of a "
                      "Section stays.")

def test_fig1_is_the_example_figure_at_600_dpi_and_the_text_width():
    """Since 2026-10-19 Fig. 1 is the TikZ example figure (the author's request:
    figures shown by example, short captions), put in at the template's text
    width; the concept and mechanism figures stay in the repository for the
    artifact and the fill no longer reads them."""
    from paper.tools import page_estimate as P
    from scripts import paper_format
    assert (M.FIG1_PNG, M.FIG1_CAPTION) == ("fig_example.png", "fig_example_caption.txt")
    assert M.FIG1_WIDTH_IN == paper_format.TEXT_WIDTH_IN == P.FIG1_WIDTH_IN == 6.5
    assert P.FIG1_PNG == FIG1 / M.FIG1_PNG and P.FIG1_CAPTION_FILE == FIG1 / M.FIG1_CAPTION
    assert (FIG1 / "fig_example.tex").is_file()
    assert "fig_example" in (FIG1 / "build_figs.sh").read_text(encoding="utf-8")
    w, h, dpi = _png_header((FIG1 / M.FIG1_PNG).read_bytes())
    # Placed FIG1_WIDTH_IN wide, it has 600 pixels per inch or more
    # (paper/VENUE_FACTS.md), and says so itself; page_estimate's stand-in has its height.
    assert w / M.FIG1_WIDTH_IN >= M.FIG_DPI == 600
    assert round(dpi) >= 600
    assert h * M.FIG1_WIDTH_IN / w == pytest.approx(P.FIG1_HEIGHT_IN, abs=0.01)
    caption = (FIG1 / M.FIG1_CAPTION).read_text(encoding="utf-8")
    assert caption.startswith("Fig. 1. ") and caption.endswith("\n") and "«" not in caption
    assert caption.strip() == P.FIG1_CAPTION
    # The caption is a title (the author's request of 2026-10-19) and names no
    # internal label (review MR10); what never crosses is drawn in the figure.
    assert "privdeep-internal" not in caption and "paper-freeze" not in caption
    assert caption.strip() == "Fig. 1. Cloze on an illustrative budget row."
    # Since later on 2026-10-21 the plain figure says what stays in T in its footer.
    assert "stay in T" in (FIG1 / "fig_example.tex").read_text(encoding="utf-8")
    # 1.5 MB since 2026-10-19: the example figure replaces two figures that together
    # weighed about 1.3 MB (fig_concept.png and fig_mechanism.png).
    assert (FIG1 / M.FIG1_PNG).stat().st_size < 1_500_000


@needs_docx
def test_the_drawn_fig1_and_its_caption_pass_the_numeral_check(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    _fig1_doc(src, heading=True)
    out = _out(tmp_path, VALUES)
    r = M.fill(src, dst, out)           # the repository's own paper/figures
    assert r["inserted"] == ["fig1"]
    f = docx.Document(str(dst))
    shape = f.inline_shapes[0]
    assert shape.width.inches == pytest.approx(6.5)
    assert shape.height.inches == pytest.approx(1.38, abs=0.01)
    # No distance from the text: LibreOffice would otherwise offset the picture past
    # the right margin (ICCIT compliance check of 2026-10-25). Space above it keeps its
    # grey band off the text above.
    assert [shape._inline.get(f"dist{s}") for s in "TBLR"] == ["0"] * 4
    pic = next(p for p in f.paragraphs if p._p.findall(".//" + M._qn("wp:inline")))
    assert pic.paragraph_format.space_before.pt == M.FIG1_SPACE_BEFORE_PT
    caption = (FIG1 / M.FIG1_CAPTION).read_text(encoding="utf-8").strip()
    assert caption in [p.text.replace('\xa0', ' ') for p in f.paragraphs]
    res = C.check(dst, out / "numbers.json")
    assert res.fails == [] and res.warns == [], [x.line() for x in res.fails + res.warns]


def test_fig1_follows_the_compose_file():
    """The boxes are docker-compose.yml's: the services on each network, the
    network with no route out, and the one service mounting each volume."""
    import yaml
    F = _fig1_module()
    compose = yaml.safe_load((_REPO / "docker-compose.yml").read_text(encoding="utf-8"))
    on: dict[str, set[str]] = {}
    mounts: dict[str, set[str]] = {}
    for name, svc in compose["services"].items():
        nets = svc.get("networks") or []
        for n in (nets if isinstance(nets, list) else list(nets)):
            on.setdefault(n, set()).add(name)
        for v in svc.get("volumes") or []:
            src = v.split(":")[0] if isinstance(v, str) else v.get("source")
            if src in (compose.get("volumes") or {}):
                mounts.setdefault(src, set()).add(name)
    assert {n: set(s) for n, s in F.TOPOLOGY["networks"].items()} == on
    assert {n for n, c in compose["networks"].items() if c.get("internal")} == set(
        F.TOPOLOGY["internal"])
    assert {n: set(s) for n, s in F.TOPOLOGY["volumes"].items()} == mounts
    # The one service on both networks is the one the figure draws as M.
    assert set.intersection(*on.values()) == {"reference-monitor"}
    assert "reference-monitor" in F.LABELS["m"]


def test_every_word_of_fig1_has_a_source_that_exists():
    F = _fig1_module()
    assert set(F.LABELS) == set(F.SOURCES)
    for key, src in F.SOURCES.items():
        for path in re.findall(r"([\w./-]+\.(?:py|yml|md))", src):
            here = [_REPO / path, _REPO / "backend/src/privdeep/tier2_cloud" / path]
            assert any(p.exists() for p in here), (key, path)
    # The caption carries no measured number: only the figure's label and §3.
    assert re.findall(r"\d+", F.CAPTION) == ["1", "3"]


def test_fig1_text_is_legible_inside_the_canvas_and_never_overlaps():
    pytest.importorskip("matplotlib")
    import matplotlib.pyplot as plt
    from matplotlib.text import Text
    F = _fig1_module()
    fig = F.draw()
    try:
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        texts = [t for t in fig.findobj(Text) if t.get_text()]
        assert min(t.get_fontsize() for t in texts) >= F.MIN_PT == 6.5
        boxes = [t.get_window_extent(r) for t in texts]
        for t, b in zip(texts, boxes, strict=True):
            assert fig.bbox.x0 <= b.x0 and b.x1 <= fig.bbox.x1, t.get_text()
            assert fig.bbox.y0 <= b.y0 and b.y1 <= fig.bbox.y1, t.get_text()
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                assert not boxes[i].overlaps(boxes[j]), (texts[i].get_text(),
                                                         texts[j].get_text())
        # Every line of LABELS is drawn, once, and nothing else is.
        assert sorted(t.get_text() for t in texts) == sorted(
            ln for lines in F.LABELS.values() for ln in lines)
        assert tuple(fig.get_size_inches()) == (F.WIDTH_IN, F.HEIGHT_IN) == (6.0, 1.8)
    finally:
        plt.close(fig)


def test_the_fig1_source_regenerates_identical_bytes(tmp_path):
    pytest.importorskip("matplotlib")
    F = _fig1_module()
    written = F.write(tmp_path)
    for path in written:
        assert path.read_bytes() == (FIG1 / path.name).read_bytes(), path.name
    again = F.write(tmp_path / "again")
    assert [p.read_bytes() for p in again] == [p.read_bytes() for p in written]


def test_fig1_draws_t_around_its_slm_and_core_and_every_label_in_its_box():
    """T holds the SLM (THREAT_MODEL.md §1, §3 Parties), so T's box holds the
    local-slm and privdeep-core boxes; every box's words stand inside it, and no
    text straddles a box's edge (the networks excepted: M stands on both)."""
    pytest.importorskip("matplotlib")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    from matplotlib.text import Text
    F = _fig1_module()
    fig = F.draw()
    try:
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        box = {p.get_gid(): p.get_window_extent(r) for p in fig.findobj(FancyBboxPatch)
               if p.get_gid()}
        assert set(box) == {"internal", "external", "t", "slm", "core", "tier1", "mapping",
                            "facts", "tier3", "m", "c"}

        def inside(inner, outer):
            return (outer.x0 <= inner.x0 and inner.x1 <= outer.x1
                    and outer.y0 <= inner.y0 and inner.y1 <= outer.y1)

        for part in ("slm", "core"):
            assert inside(box[part], box["t"]), part
        for part in ("tier1", "mapping", "facts", "tier3"):
            assert inside(box[part], box["core"]), part
        assert inside(box["t"], box["internal"])
        assert not box["m"].overlaps(box["t"]) and not box["c"].overlaps(box["external"])
        texts = {t.get_text(): t.get_window_extent(r) for t in fig.findobj(Text)
                 if t.get_text()}
        for key, b in box.items():
            for line in F.LABELS[key]:
                assert inside(texts[line], b), (key, line)
        for line, tb in texts.items():
            for key, b in box.items():
                if key not in ("internal", "external") and tb.overlaps(b):
                    assert inside(tb, b), (line, key)
        # Tier 3 is §4.3's: the scrub is the drafter's (§4.2), and "not recorded"
        # is the literal an absent attribute renders, so it is quoted.
        assert "scrub" not in " ".join(ln for v in F.LABELS.values() for ln in v)
        tier3 = " ".join(F.LABELS["tier3"])
        assert "“not recorded”" in tier3 and "refusal" in tier3
    finally:
        plt.close(fig)


def _settings_children(path: Path) -> list[str]:
    """The local names of word/settings.xml's child elements, in order."""
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/settings.xml"))
    return [etree.QName(c).localname for c in root if isinstance(c.tag, str)]


@needs_docx
def test_fill_marks_the_filled_document_do_not_compress_pictures(tmp_path):
    """Word's "Do not compress images in file" is set once, in its place in
    CT_Settings, so that Word does not compress Fig. 1 to its default resolution."""
    out = _out(tmp_path, VALUES)
    fig1 = _fig1_dir(tmp_path, "Fig. 1. Architecture, drawn.")
    # python-docx's default template, which already carries the setting, and the
    # manuscript skeleton's settings, which do not, with the test's own body (the
    # skeleton's author fields have no value here).
    src = tmp_path / "src.docx"
    _fig1_doc(src)
    assert _settings_children(src).count("doNotAutoCompressPictures") == 1
    skeleton = tmp_path / "skeleton.docx"
    d = docx.Document(str(_REPO / "paper" / "PrivDeep_ICCIT2026.docx"))
    body = d.element.body
    for el in list(body):
        if el.tag != M._qn("w:sectPr"):
            body.remove(el)
    d.add_paragraph("[FIGURE 1 PLACEHOLDER] " + M.BLOCK_DESCRIPTIONS["fig1"])
    d.save(str(skeleton))
    assert "doNotAutoCompressPictures" not in _settings_children(skeleton)
    after = {t.split(":")[1] for t in M._AFTER_NO_COMPRESS}
    before = {"zoom", "proofState", "defaultTabStop", "footnotePr", "endnotePr", "compat",
              "docVars", "rsids", "mathPr", "themeFontLang", "clrSchemeMapping"}
    for i, s in enumerate((src, skeleton)):
        dst = tmp_path / f"dst{i}.docx"
        M.fill(s, dst, out, fig1_dir=fig1)
        kids = _settings_children(dst)
        assert kids.count("doNotAutoCompressPictures") == 1, s
        at = kids.index("doNotAutoCompressPictures")
        assert not after & set(kids[:at]) and not before & set(kids[at + 1:]), (s, kids)
        again = tmp_path / f"again{i}.docx"
        M.fill(dst, again, out, fig1_dir=fig1)
        assert _settings_children(again) == kids
    # In the skeleton it stands between clrSchemeMapping and shapeDefaults.
    kids = _settings_children(tmp_path / "dst1.docx")
    at = kids.index("doNotAutoCompressPictures")
    assert kids[at - 1:at + 2] == ["clrSchemeMapping", "doNotAutoCompressPictures",
                                   "shapeDefaults"]
    # With no successor, it follows the sequence and precedes the w14/w15 extensions.
    d = docx.Document(str(skeleton))
    settings = d.settings.element
    for t in M._AFTER_NO_COMPRESS:
        for el in settings.findall(M._qn(t)):
            settings.remove(el)
    M._keep_pictures_uncompressed(d)
    M._keep_pictures_uncompressed(d)
    names = [c.tag.split("}")[1] for c in settings]
    assert names[-3:] == ["doNotAutoCompressPictures", "docId", "docId"], names
    assert names[-4] == "clrSchemeMapping" and names.count("doNotAutoCompressPictures") == 1


# ------------------------------------------------------------ claims review of 2026-10-09

def _manifest(runs_dir: Path, run: str, facts: dict) -> None:
    (runs_dir / run).mkdir(parents=True)
    (runs_dir / run / "manifest.json").write_text(json.dumps({"facts": facts}), encoding="utf-8")


def test_a_fact_store_that_counted_no_structural_fact_reads_zero_not_na(tmp_path, capsys):
    """T3-facts-structural-na: the fact store's Counter writes no
    facts_structural key for a count of zero, and the aggregate reads the absent
    key as no value. The build reads the manifests: facts_total recorded and no
    counter is a measured 0; no facts block at all is no fact store."""
    runs = ["t__r1", "t__r1__rep2", "t__r1__rep3"]
    _manifest(tmp_path, runs[0], {"fact_dropped_unknown_header": 2, "facts_total": 0})
    _manifest(tmp_path, runs[1], {"facts_total": 0})
    _manifest(tmp_path, runs[2], {"facts_structural": 3, "facts_total": 5})
    agg = {"n": 1, "mean": 3.0, "stdev": None, "min": 3, "max": 3, "values": [3]}
    stat, src, fmt = M._facts_structural(agg, "e.json#/x", "f", runs, tmp_path)
    assert stat["values"] == [0, 0, 3] and stat["n"] == 3
    assert M._fmt_cell(stat, 3, 1) == "1.0 [0.0–3.0]"
    assert "0 for t__r1, t__r1__rep2" in src and "e.json#/x" in src
    assert "fact_store.py" in fmt
    # An arm without a fact store keeps the aggregate's empty cell.
    _manifest(tmp_path, "t__baseline", {})
    empty = {"n": 0, "mean": None, "stdev": None, "min": None, "max": None, "values": []}
    assert M._facts_structural(empty, "e#/b", "f", ["t__baseline"], tmp_path) == (
        empty, "e#/b", "f")
    # Manifests that are not the ones the aggregate read stop the build.
    with pytest.raises(SystemExit, match="not the counters"):
        M._facts_structural({**agg, "values": [4]}, "e#/x", "f", runs, tmp_path)
    # A run with no manifest leaves the aggregate's figure, and says so.
    assert M._facts_structural(agg, "e#/x", "f", ["t__r9"], tmp_path)[0] is agg
    assert "no manifest for t__r9" in capsys.readouterr().err


def _t3_numbers() -> M.Numbers:
    nb = M.Numbers()
    for a in (*M.ARM_ORDER, "r1prefix", "oracle"):
        for m in ("runs", "rba", "attr", "covcorrect", "covstated", "pruned", "factsstruct"):
            if a in ("r1prefix", "oracle") and m in ("pruned", "factsstruct"):
                continue
            if a == "r1prefix" and m not in ("rba", "attr"):
                continue
            val = ("n/a" if m == "factsstruct"
                   else "0.0 [0.0–0.0]" if m == "pruned" and a != "r1"
                   else "3" if m == "runs" else "7")
            nb.put(f"t4_A_{a}_{m}", val, "x")
    nb.put("table_cells_A", "7", "x")
    nb.put("n_registered_repeats", "3", "x")
    for k in ("h3_outcome", "h3_reason", "h3_runs_not_scanned_list"):
        nb.put(k, "none", "x")
    return nb


def test_table2_prints_r1s_store_once_and_fig3_has_a_title_caption():
    """r1's store holds the same cells in every run, so Table 2(b) prints its value
    once ("18", not "18.0 [18.0–18.0]"); the per-arm scores and the oracle's
    reports are Fig. 3's, whose caption is a title since 2026-10-19 (the author's
    request); the pre-fix figures stay keys."""
    assert M._collapse_range("18.0 [18.0–18.0]") == "18"
    assert M._collapse_range("9.7 [8.0–11.0]") == "9.7 [8.0–11.0]"
    assert M._collapse_range("n/a") == "n/a"
    nb = _t3_numbers()
    nb.values["t4_A_r1_factsstruct"] = "18.0 [18.0–18.0]"
    spec = M.table3(nb, {"A": "task-a"}, [])
    assert nb["r1_store_cells_A"] == "18"
    rows = [r for r in spec.rows if isinstance(r, list) and len(r) == len(M.T3_CORPUS_HEAD)]
    assert len(rows) == 1 and rows[0][:4] == ["A*", "«table_cells_A»", "«r1_store_cells_A»",
                                              "«t4_A_r1_pruned»"]
    assert rows[0][4:] == ["n/a"] * (len(M.T3_CORPUS_HEAD) - 4)
    note = spec.notes[0]
    assert "pre-fix" not in note and "r1prefix" not in note
    assert not any("H3" in n.replace("(H3)", "") for n in spec.notes)
    cap = M.FIG3_CAPTION_TEXT
    assert cap == "Fig. 5. Role binding and coverage per corpus and private arm (RQ3)."
    assert not M.PLACEHOLDER_RE.findall(cap)
    assert [t for _, t in M.FIG3_METRICS] == [
        "(a) Attribute accuracy", "(b) RBA", "(c) coverage_stated", "(d) coverage_correct"]


def test_table2_note_reads_both_parts_in_one_short_note():
    """The author's choice of 2026-10-18 and request of 2026-10-19: one short note
    that says how to read the table; definitions are Section VI-B's."""
    nb = _folded_numbers()
    spec = M.table3(nb, {"A": "task-a", "C": "task-c"}, ["A", "C"])
    assert len(spec.notes) == 1
    note = spec.notes[0]
    for term in ("(a) Each arm's runs on corpora A–D pooled",
                 "an adversary match's N sums the run's extracted entities",
                 "(b) *Corpus A is in-sample for the RBA attribute lexicon's fix",
                 "Jaccard compares a corpus's unseeded extractions with each other",
                 "whose fact references are all absent (Section IV-C)",
                 "control: unrelated prose of the same lengths"):
        assert term in note, term
    # The leftover clauses the readability review of 2026-10-19 found are gone.
    assert "label surface" not in note and "never summed" not in note
    # Units and scopes added by the reviews of 2026-10-19 (occurrences against
    # (run, string) pairs; which columns span every extraction).
    assert "(run, string) pairs" in note and "echo, over the shared extraction" in note
    # Raised from 110 on 2026-10-19 for the label-word gloss (the author's request to
    # give H3's object one name; label words are counted in its column).
    assert "label words (constants that are ordinary English words, counted apart)" in note
    assert len(note.split()) <= 125
    assert not M.PLACEHOLDER_RE.findall(note)


def _fp_report(reset_rules: dict, session_rules: dict | None = None) -> dict:
    def mode(rules: dict) -> dict:
        return {"n": 60, "refusals": sum(rules.values()), "rate": 0.1, "wilson95": [0.05, 0.2],
                "by_stratum": {}, "rules_fired": rules}

    def linter(name: str, k: int, rules: dict) -> dict:
        return {"name": name, "n": 60, "refusals": k, "rate": k / 60, "wilson95": [0, 0.3],
                "rules_fired": rules}
    return {"deployed_chain": {"modes": {"reset_per_payload": mode(reset_rules),
                                         "session_per_stratum": mode(session_rules
                                                                     or reset_rules)},
                               "positive_control": {"caught": 4, "n_sentences": 4}},
            "linters": {"entropy": linter("entropy", 5, {"COVERT_CHANNEL_DETECTED": 5}),
                        "regex": linter("regex", 2, {"SECRET_PATTERN": 2}),
                        "shingle": linter("shingle", 0, {})},
            "n_per_stratum": 12, "corpus": {"n_payloads": 60}, "git": {"commit": "f" * 40}}


def test_table2_says_which_rules_the_deployed_refusals_fired_only_when_they_match(capsys):
    """T2-deployed-rate-composition and T2-labels."""
    rules = {"COVERT_CHANNEL_DETECTED": 5, "SECRET_PATTERN": 2}
    nb = M.Numbers()
    M.numbers_from_fp(nb, M.Src("fp_report.json", _fp_report(rules)))
    assert nb["fp_deployed_rules_standalone"] == "yes"
    spec = M.table2(nb)
    note = spec.notes[0]
    # The change set's note (2026-10-10, TAB-II-note), with the pointer to the
    # refusals the clean payloads cannot measure (re-review of 2026-10-18, A7), as
    # shortened on 2026-10-19 (the carry-kept row became its clause).
    assert note == (
        "(a) containment refuses no clean payload (a refusal would be a Bloom collision; "
        "benign text sharing a corpus run is not among these payloads, Section VII); the "
        "chain's refusals are the entropy and regex screens', and keeping the carry over each "
        "stratum gives the same «fp_deployed_session». Strata weigh equally, so the "
        "all-strata rate is not a traffic rate; refusals within one run are not independent, "
        "and the counts and the positive control get no interval.")
    # Scope §8.11: the entropy screen's standalone aggregate stays in its own row.
    assert "fp_standalone_entropy" not in note
    labels = [r[0] for r in spec.rows if isinstance(r, list)]
    assert ["Entropy screen alone", "Regex screen alone", "Containment alone"] == [
        x for x in labels if x.endswith(" alone")]
    assert not any("linter" in x for x in labels)
    assert M.T2_STRATUM_LABEL["abstract_json_compact"] == (
        "Compact abstract-payload JSON in one message string")
    # A containment rule in either mode: no such sentence, and a warning.
    nb = M.Numbers()
    M.numbers_from_fp(nb, M.Src("fp_report.json", _fp_report(
        rules, {**rules, "RAW_CONTEXT_CONTAINMENT": 1})))
    assert nb["fp_deployed_rules_standalone"] == "no"
    assert "Containment refuses none" not in M.table2(nb).notes[0]
    assert "rules_fired differ" in capsys.readouterr().err


def _doc(doc_id: str, refused: bool, masked: int, hit: int, placeholders: int,
         surfaces: int) -> dict:
    return {"doc_id": doc_id, "refused": refused, "shingle_hits_5_8": 1 if hit else 0,
            "masked_5grams": masked, "masked_5grams_hit": hit, "placeholders": placeholders,
            "surfaces_surviving": surfaces}


def _maskdoc_probe(tmp_path: Path, cid: str, docs: list[list[dict]], shares: list[float],
                   distinct: list[int], lost: bool = False, run_first: bool = True) -> dict:
    """A probe whose extractions carry per-document verdicts, the first one a
    run's extractions.json (written here, with document d2 lost when `lost`)."""
    first = (tmp_path / "runs" / f"{cid}__r5" / "extractions.json" if run_first
             else tmp_path / "probe" / f"{cid}_0.json")
    first.parent.mkdir(parents=True, exist_ok=True)
    first.write_text(json.dumps({"documents": [
        {"doc_id": "d1", "entities": [{"text": "x"}], "relations": [], "chunks": 1,
         "dropped_malformed": 0},
        {"doc_id": "d2", "entities": [] if lost else [{"text": "y"}], "relations": [],
         "chunks": 1, "dropped_malformed": 1 if lost else 0}]}), encoding="utf-8")
    pr = _probe(cid)
    base = pr["per_extraction"][0]
    pr["per_extraction"] = []
    for ds, sh, di in zip(docs, shares, distinct, strict=True):
        ext = json.loads(json.dumps(base))
        ext["masked_exposure"].update(
            masked_5gram_hit_share=sh, documents=ds,
            refused_documents=sum(d["refused"] for d in ds),
            generator_surfaces_surviving={"distinct": di, "extractor_misses": []})
        pr["per_extraction"].append(ext)
    pr["extractions"] = [{"index": 0, "files": [str(first)]}] + [
        {"index": i, "files": [f"probe/extractions/{cid}/extraction_{i}.json"]}
        for i in range(1, len(docs))]
    return pr


def test_the_masked_document_note_says_what_passes_and_what_the_extraction_lost(tmp_path):
    """T2-maskdoc-panel, S6-16, S6-05 (Table II) and C-extraction-dropped-doc."""
    corpora = {"C": "task-c", "D": "task-d"}
    refused = [_doc("d1", True, 40, 6, 3, 1), _doc("d2", True, 40, 5, 2, 1)]
    d = _maskdoc_probe(tmp_path, "task-d",
                       [[_doc("d1", True, 40, 6, 3, 1), _doc("d2", False, 30, 0, 4, 0)],
                        refused], [0.15, 0.14], [3, 2])
    c = _maskdoc_probe(tmp_path, "task-c",
                       [[_doc("d1", True, 100, 10, 5, 1), _doc("d2", True, 50, 50, 0, 9)],
                        refused], [0.4, 0.1], [8, 2], lost=True)
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/c.json", c), M.Src("probe/d.json", d)],
                          corpora, ["C", "D"])
    assert nb["maskdoc_refused_docs_D"] == "1/2; 2/2"
    assert (nb["maskdoc_passed_shared_D"], nb["maskdoc_passed_further_D"]) == ("1", "0")
    assert nb["maskdoc_shared_first"] == "yes"
    assert (nb["extraction_docs_lost_C"], nb["extraction_replies_lost_C"]) == ("1", "1")
    assert nb["maskdoc_lost_unmasked_C"] == "yes" and nb["maskdoc_upper_from_lost_C"] == "yes"
    assert nb["extraction_docs_lost_D"] == "0" and nb["maskdoc_lost_unmasked_D"] == "n/a"
    spec = M.table2(nb)
    panel = next(r[1] for r in spec.rows if isinstance(r, tuple)
                 and r[1].startswith("(b) What entity masking alone would send"))
    # Offline, and true by construction, not a test of the runs (2026-10-19); the
    # construction fact is the (b) note's first sentence since the heading is short.
    assert panel == "(b) What entity masking alone would send, offline (P2)"
    assert spec.notes[1].startswith("(b) The sealed filter refuses any masked document that "
                                    "keeps a corpus 5-gram (true by construction).")
    # One row per corpus under a sub-header naming the block's columns, as a second
    # table (its own column count) under the caption.
    i = spec.rows.index(("panel", panel))
    assert spec.rows[i + 1] == ("head", *M.T2_MASKDOC_HEAD) and len(M.T2_MASKDOC_HEAD) == 4
    assert spec.widths2 == list(M.T2_MASKDOC_WIDTHS)
    assert spec.rows[i + 2:i + 4] == [
        ["C", "«maskdoc_refused_docs_C»", "«maskdoc_5gram_pct_C»",
         "«maskdoc_surfaces_surviving_C»"],
        ["D", "«maskdoc_refused_docs_D»", "«maskdoc_5gram_pct_D»",
         "«maskdoc_surfaces_surviving_D»"]]
    assert len(spec.rows) == i + 4
    note = spec.notes[1]
    assert note.startswith("(b) The sealed filter refuses any masked document that keeps a "
                           "corpus 5-gram (true by construction). Over each corpus's "
                           "extractions, the shared one first: equal figures once, differing "
                           "document counts in that order, shares and counts as min–max.")
    assert "passes" not in note and "Corpora with runs and no probe" not in note
    assert note.endswith("On corpus C, one extractor reply did not parse; that document stays "
                         "unmasked in its masked copy and gives its upper figures.")
    # With the figures without the lost document, the row shows them, marked, and
    # the note the figures with it.
    nb.put("maskdoc_5gram_pct_nolost_C", "10.0%", "x")
    nb.put("maskdoc_surfaces_surviving_nolost_C", "2", "x")
    spec = M.table2(nb)
    assert ["C†", "«maskdoc_refused_docs_C»", "«maskdoc_5gram_pct_nolost_C»",
            "«maskdoc_surfaces_surviving_nolost_C»"] in spec.rows
    # The refused count has no parsed-only variant, so the note says it counts the
    # unmasked document (claim-preservation review of 2026-10-19, finding 4).
    assert spec.notes[1].endswith(
        "†Corpus C's shares and counts are over its parsed documents (one extractor reply did "
        "not parse); the refused count includes that document, which stays unmasked; with that "
        "document, its figures are «maskdoc_5gram_pct_C» and «maskdoc_surfaces_surviving_C».")
    # A corpus with runs and no probe is still named.
    nb.values["maskdoc_corpora_missing"] = "B"
    assert M.table2(nb).notes[1].endswith(
        "Corpora with runs and no probe: «maskdoc_corpora_missing».")
    # The upper figures are not the lost document's when its occurrences are few.
    c2 = _maskdoc_probe(tmp_path / "c2", "task-c",
                        [[_doc("d1", True, 100, 10, 5, 9), _doc("d2", True, 50, 50, 0, 1)],
                         refused], [0.4, 0.1], [8, 2], lost=True)
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/c.json", c2)], {"C": "task-c"}, ["C"])
    assert nb["maskdoc_upper_from_lost_C"] == "no"
    assert "upper figures" not in M.table2(nb).notes[1]
    assert "On corpus C, " in M.table2(nb).notes[1]
    assert "stays unmasked in its masked copy." in M.table2(nb).notes[1]


def test_a_masked_document_that_keeps_a_corpus_5gram_and_passes_stops_the_build(tmp_path):
    bad = _maskdoc_probe(tmp_path, "task-d",
                         [[_doc("d1", True, 40, 6, 3, 1), _doc("d2", False, 30, 2, 4, 0)]],
                         [0.15], [3])
    with pytest.raises(SystemExit, match="construction fact under P2"):
        M.numbers_from_probes(M.Numbers(), [M.Src("probe/d.json", bad)], {"D": "task-d"},
                              ["D"])


def test_without_a_runs_first_extraction_the_note_claims_no_order_and_no_loss(tmp_path):
    d = _maskdoc_probe(tmp_path, "task-d",
                       [[_doc("d1", True, 40, 6, 3, 1), _doc("d2", False, 30, 0, 4, 0)]],
                       [0.15], [3], lost=True, run_first=False)
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/d.json", d)], {"D": "task-d"}, ["D"])
    assert nb["maskdoc_shared_first"] == "no"
    assert nb["maskdoc_passed_shared_D"] == "n/a" and nb["extraction_docs_lost_D"] == "n/a"
    note = M.table2(nb).notes[1]
    assert "in the probe's order" in note and "shared extraction" not in note
    # The synthetic probe of the earlier tests has no per-document verdicts.
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", _probe("task-a"))], {"A": "task-a"},
                          ["A"])
    assert nb["maskdoc_passed_shared_A"] == "n/a" and nb["maskdoc_shared_first"] == "no"


def _residue(k: int, n: int) -> dict:
    return {"referents_with_residue": k, "referents_total": n}


def test_the_residue_is_split_into_the_shared_extraction_and_the_further_ones(tmp_path):
    """Follow-up F11: running text states the shared extraction's residue and the
    further extractions' as a list with 'and', not the '; '-joined key."""
    docs = [[_doc("d1", True, 40, 6, 3, 1)]] * 4
    pr = _maskdoc_probe(tmp_path, "task-a", docs, [0.1] * 4, [1] * 4)
    for ext, (k, n) in zip(pr["per_extraction"], ((5, 44), (5, 44), (6, 46), (7, 47)),
                           strict=True):
        ext["r5_residue"] = _residue(k, n)
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", pr)], {"A": "task-a"}, ["A"])
    assert nb["r5_residue_A"] == "5/44; 5/44; 6/46; 7/47"          # unchanged
    assert nb["r5_residue_shared_A"] == "5/44"
    assert nb["r5_residue_further_A"] == "5/44, 6/46 and 7/47"
    assert nb.sources["r5_residue_shared_A"]["source"] == (
        "probe/a.json#/per_extraction/0/r5_residue")
    # Two further extractions: one 'and'; none: 'none'.
    pr2 = _maskdoc_probe(tmp_path / "two", "task-a", docs[:3], [0.1] * 3, [1] * 3)
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", pr2)], {"A": "task-a"}, ["A"])
    assert nb["r5_residue_further_A"] == "0/4 and 0/4"
    pr1 = _maskdoc_probe(tmp_path / "one", "task-a", docs[:1], [0.1], [1])
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", pr1)], {"A": "task-a"}, ["A"])
    assert (nb["r5_residue_shared_A"], nb["r5_residue_further_A"]) == ("0/4", "none")
    # Without a run's extraction first, no figure is the shared one.
    nb = M.Numbers()
    M.numbers_from_probes(nb, [M.Src("probe/a.json", _probe("task-a"))], {"A": "task-a"},
                          ["A"])
    assert nb["r5_residue_A"] == "0/4"
    assert (nb["r5_residue_shared_A"], nb["r5_residue_further_A"]) == ("n/a", "n/a")


def test_the_fig2_caption_is_the_change_sets_and_names_the_corpus():
    """Change set 2026-10-10, FIG-2: the caption names the run's corpus, not its run
    directory and exchange, says what the SENT lines are, and no longer lists what
    the same request holds beyond the listing (review 2026-10-09, F2-sent-lines,
    superseded)."""
    sent = ("Accountability for the works remains under ⟦FACT:0a1b2c3d.OWNERREF⟧ and "
            "⟦FACT:4e5f6a7b.OWNERREF⟧ respectively, as the steering group recorded in its "
            "spring minutes.")
    res = ("Accountability for the works remains under Alice Brown and Bob Green "
           "respectively, as the steering group recorded in its spring minutes.")
    m = M._resolution_match(res, sent)
    ex = M.Excerpt("task-a__r1", 1, "draft:s1", [], sent, res, 1, 1, {},
                   slot_lines=SLOT_LINES, attr_line=LONG_ATTR,
                   code_lines=["- OWNERREF: owner"],
                   value_ends=[m.end(g) - m.start() for g in range(1, m.re.groups + 1)])
    short = ex.short()
    assert (short["codes"], short["codes_used"], short["others"]) == (0, 1, 1)
    cap = M.fig2_caption(ex)
    # "Wire excerpt of arm r1" since 2026-10-25: no caption starts in lower case, and
    # the caption is a title (§4 says the LOCAL line never leaves T).
    assert cap == "Fig. 2. Wire excerpt of arm r1 from a registered run on corpus «fig2_corpus»."
    assert "fig2_run" not in cap and "fig2_exchange" not in cap
    assert "the same request lists" not in cap
    full = M.fig2_caption(_excerpt(SENT, RESOLVED))
    assert full.startswith(M.FIG2_CAPTION_SHORT) and "(SENT)" not in full
    # The corpus is the run's letter under the aggregate's map, else its task id.
    nb = M.Numbers()
    M.numbers_from_fig2(nb, ex, corpora={"A": "task-a", "B": "task-b"})
    assert nb["fig2_corpus"] == "A" and nb["fig2_run"] == "task-a__r1"
    assert set(M.PLACEHOLDER_RE.findall(cap)) <= set(nb.values)
    nb2 = M.Numbers()
    M.numbers_from_fig2(nb2, ex)
    assert nb2["fig2_corpus"] == "task-a"


# ------------------------------------------------------------ change set of 2026-10-10
# The keys and builders the change set's text needs (paper/reviews/
# fix_proposals_2026-10-10.json, CODE-1 to CODE-4), added on 2026-10-15.

def test_every_run_exact_p_is_one_over_the_central_binomial():
    """LIM-4 and Table IV's note: with n runs a side, the chance under
    exchangeability that every run of one side exceeds every run of the other."""
    assert M.every_run_exact_p(3) == pytest.approx(1 / 20)
    assert [M.fmt_p(M.every_run_exact_p(n)) for n in (1, 2, 3, 4, 5)] == [
        "0.50", "0.17", "0.05", "0.014", "0.004"]
    ev = _evaluation()
    ev["hypotheses"] = {**_hypotheses()["hypotheses"], "corpora": {"A": "task-a"}}
    nb = M.Numbers()
    M.numbers_from_evaluation(nb, M.Src("evaluation.json", ev))
    assert nb["every_run_exact_p"] == "0.05" and nb["n_registered_repeats"] == "3"
    assert "1/C(2n, n)" in nb.sources["every_run_exact_p"]["format"]
    with pytest.raises(ValueError):
        M.fmt_p(0)


def test_judge_repeats_counts_each_pooled_runs_judgements(capsys):
    def ev(counts: dict[str, int], total: int | None = None) -> M.Src:
        return M.Src("evaluation.json", {
            "arms": {"r1": {"runs": sorted(counts)}},
            "utility": {"judgements": sum(counts.values()) if total is None else total},
            "per_task": [{"run": r, "judgements": [{"repeat": i} for i in range(n)]}
                         for r, n in counts.items()]
                        + [{"run": "unpooled", "judgements": [{}]}]})
    nb = M.Numbers()
    M.numbers_from_judge_repeats(nb, ev({"t__r1": 3, "t__r1__rep2": 3}))
    assert nb["judge_repeats"] == "3"
    nb = M.Numbers()
    M.numbers_from_judge_repeats(nb, ev({"t__r1": 2, "t__r1__rep2": 3}, total=9))
    assert nb["judge_repeats"] == "2–3"
    err = capsys.readouterr().err
    assert "judge_repeats is a range" in err and "utility.judgements 9" in err
    nb = M.Numbers()
    M.numbers_from_judge_repeats(nb, M.Src("e", {"arms": {}, "per_task": []}))
    assert nb["judge_repeats"] == "n/a"


def _echo_probe(tmp_path: Path, cid: str, echo: list[float], shared: bool = True) -> dict:
    pr = _probe(cid)
    base = pr["per_extraction"][0]
    pr["per_extraction"] = []
    for e in echo:
        x = json.loads(json.dumps(base))
        x["predicates"].update(echo_share=e, length_tokens={"1": 2, "2": 2})
        pr["per_extraction"].append(x)
    first = (f"eval/runs/{cid}__r5/extractions.json" if shared
             else f"eval/probe/extractions/{cid}/extraction_0.json")
    pr["extractions"] = [{"index": i, "files": [first if i == 0 else f"x/{cid}_{i}.json"]}
                         for i in range(len(echo))]
    # The probe's own control averages every extraction's lengths; the shared
    # extraction's is recomputed from the per-length shares at its lengths only.
    pr["predicate_control"] = {"expected_echo_share": 0.3,
                               "per_length": {"1": {"echo_share": 0.0},
                                              "2": {"echo_share": 0.1},
                                              "3": {"echo_share": 0.9}}}
    return pr


def test_the_pooled_predicate_echo_is_the_shared_extractions_only(tmp_path):
    """CON-1: the relations C receives are the shared extraction's; the further
    probe extractions were never sent, so pred_echo_pct_all and its control read
    the first extraction alone, and are n/a unless every measured corpus has one."""
    corpora = {"A": "task-a", "C": "task-c"}
    probes = [M.Src("probe/a.json", _echo_probe(tmp_path, "task-a", [0.4, 0.2])),
              M.Src("probe/c.json", _echo_probe(tmp_path, "task-c", [0.1, 0.5]))]
    nb = M.Numbers()
    M.numbers_from_probes(nb, probes, corpora, ["A", "C"])
    assert nb["pred_echo_pct_A"] == "20.0%–40.0%"                 # every extraction
    assert (nb["pred_echo_pct_shared_A"], nb["pred_echo_pct_shared_C"]) == ("40.0%", "10.0%")
    assert nb["pred_control_echo_pct_A"] == "30.0%"               # the probe's own
    assert nb["pred_control_echo_pct_shared_A"] == "5.0%"         # (0+0+0.1+0.1)/4
    assert nb["pred_echo_pct_all"] == "10.0%–40.0%"
    assert nb["pred_control_echo_pct_all"] == "5.0%"
    # A measured corpus without a shared first extraction: n/a, said in the source.
    nb = M.Numbers()
    M.numbers_from_probes(nb, [probes[0], M.Src("probe/c.json", _echo_probe(
        tmp_path, "task-c", [0.1], shared=False))], corpora, ["A", "C"])
    assert nb["pred_echo_pct_shared_C"] == "n/a" and nb["pred_echo_pct_all"] == "n/a"
    nb = M.Numbers()
    M.numbers_from_probes(nb, probes[:1], corpora, ["A", "C"])
    assert nb["pred_echo_pct_all"] == "n/a"
    assert "no probe for C" in nb.sources["pred_echo_pct_all"]["source"]


def test_the_monitor_and_extractor_constants_are_the_codes():
    """DES-1 and DES-5: the entropy threshold, the entropy screen's minimum token
    length, the regex rule count and the extractor's temperature are read from the
    code, which this test pins."""
    import inspect
    g = M.gateway_constants()
    assert g == {"entropy_threshold": "4.85", "entropy_min_len": 24, "regex_rules": 16}
    ent = M._load_by_path(M.ENTROPY_LINTER, "_test_entropy_linter")
    assert inspect.signature(ent.scan_payload).parameters["min_length"].default == 24
    rules = M._load_by_path(M.REGEX_RULES, "_test_regex_rules").DEFAULT_RULES
    assert len(rules) == g["regex_rules"]
    assert M.slm_constants() == {"temperature": 0.1}


def test_an_extractor_call_with_its_own_temperature_is_refused(tmp_path, monkeypatch):
    src = tmp_path / "ontology_extractor.py"
    src.write_text("async def f(c):\n    await c.chat([], think=False, temperature=0.7)\n",
                   encoding="utf-8")
    monkeypatch.setattr(M, "EXTRACTOR", src)
    with pytest.raises(SystemExit, match="passes its own temperature at line"):
        M.slm_constants()


def test_a_rate_is_written_as_a_power_of_ten():
    assert M.fmt_pow10(1e-7) == "1 in 10⁷" and M.fmt_pow10(2.5e-7) == "1 in 4,000,000"
    assert M.fmt_pow10(1e-12) == "1 in 10¹²"
    # The superscript exponent is not a numeral: the value's span covers it.
    assert [n for n, _, _ in C.numerals("a rate of 1 in 10⁷ per shingle")] == ["1", "10"]


def _deploy(tmp_path: Path, shingles: dict[str, int], target: float = 1e-7,
            m_bits: int | None = None) -> tuple[M.Src, Path]:
    from shared.shingle_filter import BloomFilter
    runs = tmp_path / "runs"
    for run, n in shingles.items():
        b = BloomFilter.sized_for(n, target, b"s")
        (runs / run).mkdir(parents=True)
        (runs / run / "manifest.json").write_text(json.dumps({
            "containment": {"m_bits": m_bits or b.m_bits, "k": b.k, "shingles": n},
            "local_model": "m:1b", "cloud_models_resolved": ["cloud-x"],
            "started_at": 1791357227.0, "finished_at": 1791357314.0,
            "runtime": {"local_model": {"details": {"quantization_level": "Q4_K_M"}}}}),
            encoding="utf-8")
    ev = M.Src("evaluation.json", {
        "hypotheses": {"corpora": {"A": "task-a", "B": "task-b"}},
        "arms": {"r1": {"runs": sorted(shingles)}},
        "per_task": [{"run": r, "task_id": r.split("__")[0]} for r in shingles]})
    return ev, runs


def test_the_bloom_filter_keys_come_from_the_runs_checked_against_the_target(tmp_path):
    ev, runs = _deploy(tmp_path, {"task-a__r1": 2208, "task-b__r1": 2256})
    fp = M.Src("fp_report.json", {"entropy_threshold": 4.85,
                                  "shingle_filter": {"target_fp_rate": 1e-07}})
    nb = M.Numbers()
    M.numbers_from_deployment(nb, ev, fp, runs)
    assert (nb["mon_bloom_bits"], nb["mon_bloom_bits_A"]) == ("74,074–75,684", "74,074")
    assert (nb["mon_bloom_k"], nb["mon_bloom_fpr"]) == ("23", "1 in 10⁷")
    assert nb["slm_quantisation"] == "Q4_K_M" and nb["slm_temperature"] == "0.1"
    assert (nb["mon_entropy_threshold"], nb["mon_entropy_min_len"],
            nb["mon_regex_rules"]) == ("4.85", "24", "16")
    # A run whose filter is not sized for its shingles at the target stops the build.
    ev2, runs2 = _deploy(tmp_path / "bad", {"task-a__r1": 2208}, m_bits=70000)
    with pytest.raises(SystemExit, match="not sized for 2208 shingles"):
        M.numbers_from_deployment(M.Numbers(), ev2, fp, runs2)
    # So does a false-positive report measured at another entropy threshold.
    with pytest.raises(SystemExit, match="overridden"):
        M.numbers_from_deployment(M.Numbers(), ev, M.Src("fp_report.json", {
            "entropy_threshold": 5.0, "shingle_filter": {"target_fp_rate": 1e-07}}), runs)
    # Without the run directories the manifest keys are n/a; the code's are not.
    nb = M.Numbers()
    M.numbers_from_deployment(nb, ev, fp, None)
    assert nb["mon_bloom_bits"] == nb["slm_quantisation"] == "n/a"
    assert nb["mon_regex_rules"] == "16"


def test_noshare_baseline_unchanged_compares_the_printed_means():
    """SET-2d says the baseline's attribute accuracy is unchanged only on 'yes'."""
    def ns(mean: float) -> M.Src:
        cell = {"runs": 3, "rba": _spr(0.5, n=1), "attribute_accuracy": _spr(mean, 0.04),
                "coverage_correct": _spr(0.09, 0.0)}
        return M.Src("review_analyses.json", {"without_share_spent": {
            "letter": "A", "facts": 11, "removed_facts": [["f", "PCT_SPENT"]],
            "cells": {"baseline": cell}}})
    for old, mean, want in (("0.50", 0.501, "yes"), ("0.50", 0.46, "no")):
        nb = M.Numbers()
        nb.put("attr_baseline_A", old, "x")
        M._review_noshare(nb, ns(mean))
        assert nb["noshare_baseline_A_unchanged"] == want
    nb = M.Numbers()
    M._review_noshare(nb, ns(0.5))
    assert "noshare_baseline_A_unchanged" not in nb.values      # no aggregate key yet


def test_the_meta_dates_must_be_their_records(tmp_path, capsys):
    """SET-4: prereg_date is the scope §8 signature, freeze_tag_date the tag
    paper-freeze's date (at freeze_commit), batch_date every pooled run's day."""
    scope = tmp_path / "SCOPE.md"
    scope.write_text("# S\n\n**Status:** signed by the author on 2026-01-01.\n\n"
                     "## 8. Amendment\n\n**Status:** signed by the author on 2026-09-27. "
                     "Drafted.\n", encoding="utf-8")
    assert M.scope_signature_date(scope) == "2026-09-27"
    ev, runs = _deploy(tmp_path, {"task-a__r1": 2208})
    tag = ("2026-10-06", "0d5e86b47447" + "0" * 28)
    good = {"prereg_date": "2026-09-27", "freeze_tag_date": "2026-10-06",
            "batch_date": "2026-10-07", "freeze_commit": "0d5e86b47447",
            "slm_model_id": "m:1b", "cloud_model_id": "cloud-x"}
    meta = lambda **o: M.Src("meta.json", {**good, **o})          # noqa: E731
    M.check_meta_dates(meta(), ev, runs, scope=scope, tag_info=tag)
    for over, match in (({"prereg_date": "2026-09-28"}, "§8 amendment was signed"),
                        ({"freeze_tag_date": "2026-10-07"}, "date of the tag"),
                        ({"freeze_commit": "abcdef1"}, "not the commit the tag"),
                        ({"batch_date": "2026-10-08"}, "started and finished on"),
                        ({"slm_model_id": "other:7b"}, "local_model"),
                        ({"cloud_model_id": "other"}, "cloud_models_resolved")):
        with pytest.raises(SystemExit, match=match):
            M.check_meta_dates(meta(**over), ev, runs, scope=scope, tag_info=tag)
    # A record that cannot be read is a warning, not a pass in silence.
    M.check_meta_dates(meta(), ev, None, scope=tmp_path / "none.md", tag_info=None)
    err = capsys.readouterr().err
    assert "prereg_date unchecked" in err and "freeze_tag_date unchecked" in err
    assert "batch_date unchecked" in err
    # A date is written as YYYY-MM-DD.
    with pytest.raises(SystemExit, match="YYYY-MM-DD"):
        M.numbers_from_meta(M.Numbers(), M.Src("meta.json", {"batch_date": "7 Oct 2026"}))
    # The repository's own records.
    assert M.scope_signature_date() == "2026-09-27"
    info = M.tag_date_and_commit()
    if info is not None:
        assert info[0] == "2026-10-06" and info[1].startswith("0d5e86b47447")


def _t4_values(**over) -> dict[str, str]:
    v = {"h1_A_outcome": "not testable", "h1_B_outcome": "not testable",
         "h1_A_rba_outcome": "not testable", "h1_B_rba_outcome": "not testable",
         "h1_A_attr_outcome": "holds", "h1_B_attr_outcome": "holds",
         "rba_baseline_A_n": "1/3", "rba_baseline_B_n": "1/3", "rba_baseline_D_n": "0/3",
         "min_rba_pool": "5",
         "h2_outcome": "holds", "h2_coverage_stated_r1_max": "0.00",
         "h2_coverage_stated_baseline_min": "0.91", "h2_coverage_correct_floor": "yes",
         "h2_coverage_correct_r1_max": "0.00", "h2_coverage_correct_baseline_min": "0.00",
         "h3_outcome": "holds", "h3_hits": "0", "h3_runs_scanned": "36",
         "h4d_outcome": "not testable", "h4d_structural_share_vs_A_outcome": "holds",
         "h4d_structural_share_vs_B_outcome": "holds",
         "h4d_rba_gain_vs_A_outcome": "not testable",
         "h4d_rba_gain_vs_B_outcome": "not testable", "h4dp_outcome": "holds",
         "h4dp_vs_A_D_max": "0.44", "h4dp_vs_A_A_min": "0.57", "h4dp_vs_B_B_min": "0.50",
         "n_registered_repeats": "3", "every_run_exact_p": "0.05"}
    v.update(over)
    return v


def _t4(**over) -> M.TableSpec:
    nb = M.Numbers()
    for k, val in _t4_values(**over).items():
        nb.put(k, val, "x")
    return M.table_hypotheses(nb, reasons=True)


def test_table3_states_each_hypothesis_with_its_rule_outcome_and_reason():
    """TAB-IV, as Table 3 since 2026-10-19: rows grouped by what each hypothesis
    checks, the registered rule verbatim, every figure a key; when each was
    registered is the note's."""
    spec = _t4()
    # The manuscript's table has no Reason column since later on 2026-10-21 (§6 gives
    # each reason); the artifact's keeps it (reasons=True).
    assert (spec.bid, spec.label) == ("t1", "Table 1")
    assert spec.header == [*M.T4_HEADER, "Reason"]
    assert M.T4_HEADER == ["Hypothesis", "Registered rule", "Outcome"]
    nb = M.Numbers()
    for k, val in _t4_values().items():
        nb.put(k, val, "x")
    # The manuscript's table has its own columns since 2026-10-27 (test below).
    short = M.table_hypotheses(nb)
    assert short.header == M.T1_HEADER and all(
        len(r) == 3 for r in short.rows if isinstance(r, list))
    assert spec.caption == "Registered hypotheses and outcomes."
    assert [r[1] for r in spec.rows if isinstance(r, tuple)] == [
        "Design checks", "Hypotheses on R1"]
    rows = {r[0]: r for r in spec.rows if isinstance(r, list)}
    # No dagger since 2026-10-25 (the author's request before submission); the note
    # names the hypothesis registered after the pilot.
    assert list(rows) == ["H3 (A–D)", "H2 (C)", "H1 (A, B)", "H4-D (D)", "H4-D′ (D)"]
    rules = {h: rule for h, _reg, rule in M.T4_RULES}
    assert all(rows[h][1] == rules[h] for h in rows)
    assert not any("†" in c for r in spec.rows for c in r[1:] if isinstance(c, str))
    # One outcome for both corpora when they agree (pre-submission checks of 2026-10-25).
    assert rows["H1 (A, B)"][2] == "«h1_A_outcome»"
    # Since 2026-10-26 (simulated review R4(5), the AI assistant at the author's request) each
    # outcome under a rule registered after the pilot says "post-pilot" in its cell, and
    # H4-D′'s that it is a score of the RBA rule.
    assert rows["H2 (C)"][2] == "«h2_outcome» (post-pilot floor rule)"
    assert rows["H4-D′ (D)"][2] == "«h4dp_outcome» (post-pilot; RBA-rule score)"
    assert rows["H4-D (D)"][1].endswith("smaller on D than on A and B")
    # The attribute half alone is Section VI-C's text since 2026-10-19 (with its
    # chance-level caveat), not the cell's.
    assert rows["H1 (A, B)"][3] == ("baseline has an RBA (at least «min_rba_pool» scored "
                                    "occurrences) in «rba_baseline_A_n» runs on corpus A and "
                                    "«rba_baseline_B_n» on corpus B (Section VI-C)")
    assert rows["H2 (C)"][3] == (
        "coverage_stated «h2_coverage_stated_r1_max» against «h2_coverage_stated_baseline_min»"
        " decides; coverage_correct is at the floor, not testable under the floor rule "
        "(registered after the pilot); under the pre-pilot reading, which needs both halves, H2 "
        "would not hold")
    assert rows["H3 (A–D)"][1] == ("No generator-written constant, in any rendered form, in "
                                   "core-sent content of any private-arm trace")
    assert rows["H4-D (D)"][3] == ("Share half holds; baseline has an RBA in "
                                   "«rba_baseline_D_n» runs on corpus D")
    # The floor rule's definition is §5.3's since later on 2026-10-21; the note says
    # that H2 holds under it (review of the same day), and since 2026-10-26 (R4(5)) what the
    # amber tint marks.
    assert spec.notes == [
        "H4-D′ was registered after the pilot, before the batch, as was H2's floor rule, under "
        "which H2 holds on its coverage_stated half alone (Section VI-B); HL1 before the L "
        "batch; the others before the pilot; Section V-C defines the test. Letters name "
        "corpora; L is track L. Shading marks the outcome, not a benefit; amber marks an "
        "outcome under a rule registered after the pilot."]
    assert len(spec.widths) == len(spec.header)
    # Typed text carries no numeral the checker would not allow in a table.
    paras = [C.Para(M.PLACEHOLDER_RE.sub("", c), "Table Paragraph", "t3", in_table=True)
             for r in spec.rows for c in (r[1:] if isinstance(r, tuple) else r)]
    assert C.check_paragraphs(paras + [C.Para("x 1", "Body Text", "p")], {}).fails[0].numeral \
        == "1"
    assert not [f for f in C.check_paragraphs(paras, {}).fails if f.kind == "numeral"]


def _t1(**over) -> M.TableSpec:
    nb = M.Numbers()
    for k, val in _t4_values(**over).items():
        nb.put(k, val, "x")
    return M.table_hypotheses(nb)


def test_the_manuscripts_table1_gives_every_result_at_a_glance():
    """Since 2026-10-29 (the author's "quá nhiều thông số trong phần Hypothesis, Evaluation"):
    results at a glance, grouped by §6's three questions; each row a label (a registered
    hypothesis's, or post hoc / measured), what was tested in plain words, and the result
    with its outcome first, set bold; each reason and verdict only where it is true."""
    spec = _t1(n_corpora_measured="4", hl1_outcome="not testable",
               l_collections_planner_refused_kn="4/6", l_runs="12", l_direct_crossed_kn="0/22",
               l_quasi_crossed_kn="1/340", chance_baseline_attr_obs="0.35",
               chance_baseline_attr="0.35", judge_private_range="2.18–2.24",
               judge_mean_exgr_oracle="5.00", fp_deployed_reset="122/500",
               fp_standalone_shingle="0/500")
    assert spec.caption == "Results at a glance."
    assert spec.header == M.T1_HEADER == ["Label", "What was tested", "Result"]
    assert spec.widths == M.T1_WIDTHS and abs(sum(M.T1_WIDTHS) - 6.5) < 1e-9
    assert spec.shade_col is None and (spec.tag_col, spec.lead_bold_col) == (1, 2)
    assert [r[1] for r in spec.rows if isinstance(r, tuple)] == [g for g, _ in M.T1_GROUPS]
    rows = [r for r in spec.rows if isinstance(r, list)]
    assert [r[0] for r in rows] == ["H3", "HL1", "post hoc", "H1", "H4-D", "H4-D′", "H2",
                                    "measured", "measured"]
    by = {r[0]: r for r in rows}
    for key, label in M.T1_LABELS.items():
        if label in by and label != "measured":
            assert by[label][1] == M.T1_CHECKS[key]
    assert by["H3"][2] == ("«h3_outcome»: «h3_hits» hits in «h3_runs_scanned» private-arm "
                           "traces, over «n_corpora_measured» extractions, one per corpus")
    assert by["HL1"][2] == ("«hl1_outcome»: no valid run on «l_collections_planner_refused_kn» "
                            "collections, whose planner request the monitor refused; in "
                            "«l_runs» admitted runs, «l_direct_crossed_kn» direct and "
                            "«l_quasi_crossed_kn» quasi-identifier strings crossed")
    assert by["H1"][2] == ("«h1_A_outcome»: baseline has an RBA in «rba_baseline_A_n» runs per "
                           "corpus")
    assert by["H4-D"][2] == ("«h4d_outcome»: baseline has an RBA in «rba_baseline_D_n» runs on D; "
                             "the store holds a smaller share of D's table facts than of A's "
                             "and B's (not a test)")
    assert by["H2"][2] == ("«h2_outcome» on coverage_stated alone (post-pilot floor rule; "
                           "fails under the pre-pilot reading)")
    assert by["post hoc"][2] == ("At chance: «chance_baseline_attr_obs» against a chance level "
                                 "of «chance_baseline_attr»")
    assert rows[-2][2] == ("Below the oracle: «judge_private_range» against "
                           "«judge_mean_exgr_oracle»; not a usability measure")
    assert rows[-1][2].startswith("Measured: «fp_deployed_reset» by the entropy and regex")
    assert spec.notes[0].startswith("H1 to H4-D′ and HL1 are the registered hypotheses")
    assert "Not testable (neither holds nor fails)" in spec.notes[0]
    # A verdict holds only where the values bear it out.
    other = _t1(h1_A_outcome="holds", h1_B_outcome="holds", h1_A_rba_outcome="holds",
                h1_B_rba_outcome="holds", h3_outcome="fails", h2_coverage_correct_floor="no",
                chance_baseline_attr_obs="0.50", chance_baseline_attr="0.35",
                judge_private_range="2.18–5.20", judge_mean_exgr_oracle="5.00")
    o = {r[0]: r for r in other.rows if isinstance(r, list)}
    assert o["H1"][2] == "«h1_A_outcome»" and o["H3"][2] == "«h3_outcome»"
    assert o["H2"][2] == "«h2_outcome»" and "HL1" not in o
    assert not o["post hoc"][2].startswith("At chance")
    assert not [r for r in other.rows if isinstance(r, list) and r[2].startswith("Below")]
    # The tag is italic; a leading outcome key, or a leading verdict and colon, is bold.
    assert M._lead_segments(M.T1_CHECKS["H4-D′ (D)"], tag=True)[0] == ("Not a test:", False,
                                                                        True)
    assert M._lead_segments("«h3_outcome»: x", bold_key=True) == [
        ("«h3_outcome»", True, False), (": x", False, False)]
    assert M._lead_segments("Below the oracle: x", bold_key=True)[0] == (
        "Below the oracle:", True, False)
    paras = [C.Para(M.PLACEHOLDER_RE.sub("", c), "Table Paragraph", "t1", in_table=True)
             for r in spec.rows for c in (r[1:] if isinstance(r, tuple) else r)]
    assert not [f for f in C.check_paragraphs(paras, {}).fails if f.kind == "numeral"]


def test_table3_carries_hl1_once_the_track_l_keys_exist():
    """Simulated review of 2026-10-18 (R1 W2): the hypotheses table gives HL1 too,
    under its own group, with the planner refusals as its reason when they made it
    not testable."""
    assert not any(isinstance(r, list) and r[0] == "HL1 (L)" for r in _t4().rows)
    spec = _t4(hl1_outcome="not testable", l_collections_planner_refused_kn="4/6",
               l_direct_crossed_kn="0/22", l_runs="12")
    assert ("panel", "Real documents") in spec.rows
    row = next(r for r in spec.rows if isinstance(r, list) and r[0] == "HL1 (L)")
    # Every figure is a key, filled later, as in the other rows.
    assert row[1] == M.T4_HL1[2] and row[2] == "«hl1_outcome»"
    assert row[3] == ("M refused the planner request on «l_collections_planner_refused_kn» "
                      "collections, which have no valid run; DIRECT «l_direct_crossed_kn» on "
                      "the «l_runs» runs of the rest")
    holds = next(r for r in _t4(hl1_outcome="holds", l_direct_crossed_kn="0/22",
                                l_runs="36").rows if isinstance(r, list) and r[0] == "HL1 (L)")
    assert holds[3] == "DIRECT «l_direct_crossed_kn» on the «l_runs» runs of the rest"


def test_table3_writes_a_reason_clause_only_for_the_outcomes_it_holds_for():
    """TAB-IV: 'Share half holds' only when both share comparisons hold, the
    pre-pilot clause only when the floor applied, and the baseline-RBA counts only
    while the RBA halves are not testable."""
    rows = {r[0].rstrip("†"): r for r in _t4(
        h2_coverage_correct_floor="no", h4d_structural_share_vs_B_outcome="fails",
        h4d_rba_gain_vs_A_outcome="holds", h1_A_rba_outcome="fails").rows
        if isinstance(r, list)}
    assert rows["H2 (C)"][3] == (
        "coverage_stated «h2_coverage_stated_r1_max» against «h2_coverage_stated_baseline_min»"
        "; coverage_correct «h2_coverage_correct_r1_max» against "
        "«h2_coverage_correct_baseline_min»")
    assert rows["H4-D (D)"][3] == (
        "Share half: «h4d_structural_share_vs_A_outcome» against corpus A, "
        "«h4d_structural_share_vs_B_outcome» against corpus B; RBA-gain half: "
        "«h4d_rba_gain_vs_A_outcome» against corpus A, «h4d_rba_gain_vs_B_outcome» against "
        "corpus B")
    assert rows["H1 (A, B)"][3] == (
        "RBA half: «h1_A_rba_outcome» on corpus A, «h1_B_rba_outcome» on corpus B; "
        "attribute half: «h1_A_attr_outcome» on corpus A (in-sample), «h1_B_attr_outcome» on "
        "corpus B")
    # A figure the build has no key for reads n/a, as in every other table.
    nb = M.Numbers()
    cells = [c for r in M.table_hypotheses(nb, reasons=True).rows
             for c in (r[1:] if isinstance(r, tuple) else r)]
    assert all("«" not in c for c in cells)
    first = next(r for r in M.table_hypotheses(nb, reasons=True).rows if isinstance(r, list))
    assert "n/a" in first[3]


@needs_docx
def test_fill_inserts_table3_at_its_marker_and_the_checker_passes_it(tmp_path):
    values = {**_t4_values(), "alpha": "0.41"}
    out = _out(tmp_path, values)
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    d = docx.Document()
    d.add_paragraph("Table 1 gives each hypothesis, at «alpha».")
    d.add_paragraph(M.BLOCKS["t1"] + " " + M.BLOCK_DESCRIPTIONS["t1"])
    d.add_paragraph("After.")
    d.save(str(src))
    M.write_tables_docx(src, [_t4()], out / "tables.docx", values)
    r = M.fill(src, dst, out)
    assert r["inserted"] == ["t1"] and r["placeholders_left"] == [] and r["text_kept"] == []
    f = docx.Document(str(dst))
    texts = [p.text for p in f.paragraphs]
    assert not any("PLACEHOLDER" in t or "Hypotheses: registered" in t for t in texts)
    assert texts[-1] == "After."
    cells = [[c.text for c in row.cells] for row in f.tables[-1].rows]
    assert cells[0] == [*M.T4_HEADER, "Reason"]
    assert cells[1][0] == "Design checks" and cells[2][0] == "H3 (A–D)"
    assert cells[2][3] == "0 hits in 36 runs, over one shared extraction per corpus"
    res = C.check(dst, out / "numbers.json")
    assert res.fails == [], [x.line() for x in res.fails]


# ------------------------------------------------------------ evaluation rewrite of 2026-10-15
# The six range keys for §6's lead (one range over corpora, or over the private
# arms, where the prose had quoted every value), and the one key that keeps its sd.

def _ranges_evaluation() -> dict:
    def arm(attr, cov, n=3):
        return {"runs": [f"r{i}" for i in range(n)],
                "faithfulness": {"rba": _spr(0.5), "attribute_accuracy": attr,
                                 "coverage_correct": _spr(0.1), "coverage_stated": cov},
                "pipeline": {"sentences_pruned": _spr(0.0), "facts_structural": _spr(0.0),
                             "not_recorded": _spr(0.0)}}
    empty = {"n": 0, "mean": None, "stdev": None, "min": None, "max": None}
    ev = _evaluation()
    ev["hypotheses"] = {**_hypotheses()["hypotheses"],
                        "corpora": {"A": "task-a", "B": "task-b", "C": "task-c"}}
    ev["by_task"] = {
        "task-a": {"key_sha256": "k1", "stale_key_runs": [], "table_data_cells": 3, "arms": {
            "baseline": arm(_spr(0.5, 0.04, 0.45, 0.54), _spr(1.0)),
            "r5": arm(_spr(0.46), _spr(1.0)), "r1": arm(_spr(0.78), _spr(0.75))}},
        "task-b": {"key_sha256": "k2", "stale_key_runs": [], "table_data_cells": 4, "arms": {
            "baseline": arm(_spr(0.34), _spr(1.0)),
            "r5": arm(_spr(0.35), _spr(1.0)), "r1": arm(_spr(0.69), _spr(0.81))}},
        "task-c": {"key_sha256": "k3", "stale_key_runs": [], "table_data_cells": 1, "arms": {
            "baseline": arm(_spr(0.35), _spr(0.91)),
            "r5": arm(_spr(0.31), _spr(0.91)), "r1": arm(empty, _spr(0.0))}}}
    ev["arms"] = {a: {"runs": [f"{t}__{a}" for t in ("task-a", "task-b", "task-c")],
                      "judge": {"mean_excluding_groundedness": {
                          "private": {"mean": m}, "oracle": {"mean": 5.0}}}}
                  for a, m in (("baseline", 2.18), ("r5", 2.24), ("r1", 2.20))}
    ev["utility"] = {"position_bias": _spr(0.51, 0.24, 0.0, 1.0, n=84), "judgements": 84,
                     "mean_excluding_groundedness": {"oracle": {"mean": 5.0}}}
    return ev


def test_judged_quality_by_presentation_order_is_keyed_for_the_private_reports():
    # Scope §8.10 registers the judge per criterion and per order; judged quality leaves
    # groundedness out (2026-10-28, critical review ANA2).
    ev = _ranges_evaluation()
    crit = lambda g, c, s, gr: {"goal_coverage": {"mean": g}, "coherence": {"mean": c},  # noqa: E731
                                "specificity": {"mean": s}, "groundedness": {"mean": gr}}
    ev["utility"]["by_order"] = {"private_first": {"private": crit(3.0, 3.0, 2.0, 1.0)},
                                 "oracle_first": {"private": crit(2.0, 2.0, 1.0, 5.0)}}
    nb = M.Numbers()
    M.numbers_from_evaluation(nb, M.Src("evaluation.json", ev))
    assert nb["judge_exgr_private_shown_first"] == "2.67"
    assert nb["judge_exgr_private_shown_second"] == "1.67"
    src = nb.sources["judge_exgr_private_shown_second"]["source"]
    assert "by_order/oracle_first/private" in src
    nb = M.Numbers()
    M.numbers_from_evaluation(nb, M.Src("evaluation.json", _ranges_evaluation()))
    assert nb["judge_exgr_private_shown_first"] == "n/a"


def test_the_range_keys_span_the_corpora_with_a_value_and_the_private_arms():
    nb = M.Numbers()
    M.numbers_from_evaluation(nb, M.Src("evaluation.json", _ranges_evaluation()))
    assert nb["attr_baseline_range"] == "0.34–0.50" and nb["attr_r5_range"] == "0.31–0.46"
    # r1 has no attribute accuracy on C: the range is over A and B, and the
    # sidecar says which corpus had no value.
    assert nb["attr_r1_C"] == "n/a" and nb["attr_r1_range"] == "0.69–0.78"
    assert nb.sources["attr_r1_range"]["raw"] == {"over": ["attr_r1_A", "attr_r1_B"],
                                                  "without_value": ["attr_r1_C"]}
    assert "attr_r1_A, attr_r1_B" in nb.sources["attr_r1_range"]["source"]
    assert "evaluation.json#/by_task/task-a/arms/r1/faithfulness/attribute_accuracy/mean" \
        in nb.sources["attr_r1_range"]["source"]
    assert nb["covstated_baseline_range"] == "0.91–1.00"
    assert nb["covstated_r1_range"] == "0.00–0.81"
    # The synthetic aggregate gives the oracle no attribute accuracy: the key exists, as n/a.
    assert nb["attr_oracle_range"] == M.NA
    assert nb.sources["attr_oracle_range"]["raw"]["over"] == []
    assert nb["judge_private_range"] == "2.18–2.24"
    assert "lo–hi" in nb.sources["judge_private_range"]["format"]
    # The per-cell keys dropped the sd; judge_position_bias keeps it, as §6
    # quotes it, and no other key carries one.
    assert nb["t4_A_baseline_attr"] == "0.50 [0.45–0.54]"
    assert nb["judge_position_bias"] == "0.51 ± 0.24 [0.00–1.00]"
    assert nb["judge_position_bias_mean"] == "0.51"
    assert [k for k, val in nb.values.items() if "±" in val] == ["judge_position_bias"]


def test_a_range_is_one_value_when_the_values_agree_and_na_without_any():
    nb = M.Numbers()
    for k, val in (("attr_baseline_A", "0.50"), ("attr_baseline_B", "0.50"),
                   ("attr_r1_A", "n/a"), ("judge_mean_exgr_r1", "2.20")):
        nb.put(k, val, "x")
    M.numbers_ranges(nb, ["A", "B"])
    assert nb["attr_baseline_range"] == "0.50" and nb["attr_r1_range"] == "n/a"
    assert nb["attr_r5_range"] == "n/a" and nb["covstated_r1_range"] == "n/a"
    assert nb["judge_private_range"] == "2.20"
    assert nb.sources["attr_r1_range"]["raw"]["without_value"] == ["attr_r1_A", "attr_r1_B"]




def test_table3_h1_reason_leaves_the_attribute_margin_to_the_text():
    """Since 2026-10-19 the attribute half on its own, B's margin and its
    chance-level caveat are Section VI-C's text (the re-review's R7 asked that the
    table not read stronger than it); the cell names the reason H1 is not testable
    and points there, with or without the margin keys."""
    for spec in (_t4(h1_B_attr_r1_min="0.48", h1_B_attr_baseline_max="0.45"), _t4()):
        row = next(r for r in spec.rows if isinstance(r, list) and r[0] == "H1 (A, B)")
        assert row[3].endswith("(Section VI-C)") and "against" not in row[3]


# ------------------------------------------------------------ track L (scope §12)

L_IDS = [f"tab-L{i}" for i in range(1, 7)]


def _lkn(k: int, n: int) -> dict:
    return {"k": k, "n": n, "rate": k / n if n else None}


def _ltally(kd: int, nd: int, kq: int, nq: int, *, types_d: dict | None = None,
            types_q: dict | None = None, labels_q: tuple[int, int] = (0, 2)) -> dict:
    """A real_l.py tally block: all, DIRECT and QUASI by entity type, labels apart."""
    return {"all": _lkn(kd + kq, nd + nq),
            "DIRECT": {**_lkn(kd, nd), "by_entity_type": types_d or {"PERSON": _lkn(kd, nd)}},
            "QUASI": {**_lkn(kq, nq), "by_entity_type": types_q or {"LOC": _lkn(kq, nq)}},
            "labels": {"DIRECT": _lkn(0, 0), "QUASI": _lkn(*labels_q)}}


def _lbound(b: dict) -> dict:
    return {**b, "rule_of_three": 3 / b["n"] if b["k"] == 0 and b["n"] else None}


def _l_arm(arm: str, *, runs: int = 12, valid: int = 12, kq: int = 2, delivered: int = 36,
           refused: int = 3, transport: int = 0, plans_refused: int = 0,
           net: tuple[int, int] = (1, 240), with_report: bool = True) -> dict:
    """One by_arm block of the scan, every string-level block over the arm's runs."""
    return {"runs": [f"{cid}__{arm}__rep{r}" for cid in L_IDS for r in (1, 2)][:runs],
            "runs_scanned": runs, "runs_valid": valid,
            "crossing_run_pairs": _ltally(0, 504, kq * 3, 2376),
            "crossing_distinct": _ltally(0, 42, kq, 198),
            "routes": [],
            "extractor_recall": {"all": _lkn(120, 240), "DIRECT": _lkn(30, 42),
                                 "QUASI": _lkn(90, 198)},
            "predicate_echo": {"all": _lkn(24, 240), "DIRECT": _lkn(4, 42),
                               "QUASI": _lkn(20, 198)},
            "utility": {"sections_delivered": delivered, "sections_refused": refused,
                        "sections_transport_failed": transport, "plans_refused": plans_refused,
                        "report_strings": ({"DATETIME": _lkn(30, 60), "QUANTITY": _lkn(6, 12),
                                            "ORG": _lkn(12, 24), "any": _lkn(48, 96)}
                                           if with_report else None)},
            "g3": {"collections": 6, "gross": _ltally(2, 42, 10, 198),
                   "net": {**_ltally(0, 42, net[0], net[1] - 42)}}}


def _l_scan(outcome: str = "holds", reason: str | None = None, *, ids: list[str] = L_IDS,
            direct: dict | None = None, pairs: dict | None = None, routes: list | None = None,
            incomplete: dict | None = None, uncovered: list | None = None,
            integrity: dict | None | str = "equal",
            attempts: dict | None = None, arms: dict | None = None,
            short: dict | None = None) -> dict:
    """scripts/real_l.py scan output of the shapes `scan` writes, synthetic throughout.
    `integrity` "equal" is the passed digest check; None, a scan without it."""
    if integrity == "equal":
        integrity = {"checked": True, "all_equal": True}
    direct = direct or _lkn(0, 42)
    pairs = pairs or _lkn(0, 120)
    n_runs = 36 if ids == L_IDS else 6 * len(ids)
    computed = outcome if sorted(ids) == sorted(L_IDS) else "holds"
    hl1 = {"label": "HL1", "section": "12.7", "rule": "...", "outcome": outcome,
           "reason": reason, "outcome_over_these_tasks": computed,
           "registered_collections": L_IDS, "collections": ids,
           "direct": _lbound({**direct, "by_entity_type": {"PERSON": direct}}),
           "direct_run_pairs": _lbound({**pairs, "by_entity_type": {"PERSON": pairs}}),
           "direct_one_call": {**_lkn(direct["k"], direct["n"]), "by_entity_type": {}},
           "direct_labels_crossed": _lkn(0, 0),
           "quasi_descriptive": {**_lkn(6, 198), "by_entity_type": {"LOC": _lkn(4, 80),
                                                                    "DATETIME": _lkn(2, 118)}},
           "routes": routes or [], "runs": n_runs, "runs_scanned": n_runs - len(incomplete or {}),
           "runs_incomplete": incomplete or {}, "cells_uncovered": uncovered or [],
           "attempts": attempts or {"scanned": [], "direct": None, "routes": []},
           "attempts_not_scanned": [], "integrity": integrity, "trace_gap": "..."}
    collections = {cid: {"corpus": f"eval/real/corpora/{cid}",
                         "identifiers_sha256": hashlib.sha256(cid.encode()).hexdigest(),
                         "strings": {"DIRECT": 7, "QUASI": 33, "labels": 2},
                         "short_of_registered_repeats": (short or {}).get(cid, {}),
                         "runs": {}, "attempts_scanned": [], "extractions": [],
                         "g3": {"control": {"present": True,
                                            "recovered": _ltally(0, 7, 1, 33)},
                                "arms": {}}} for cid in ids}
    return {"instrument": "scripts/real_l.py scan", "source": "PRIVDEEP_ICCIT_SCOPE.md §12",
            "git_commit": "0" * 40, "goals": ["g"], "definitions": {}, "health_rule": "...",
            "HL1": hl1,
            "by_arm": arms or {a: _l_arm(a) for a in ("baseline", "r5", "r1")},
            "extraction": {"distinct_extractions": 6,
                           "extractor_recall": {**_ltally(30, 42, 90, 198),
                                                "covered": _ltally(34, 42, 110, 198)},
                           "predicate_echo": _ltally(4, 42, 20, 198)},
            "collections": collections}


def _l_numbers(scan: dict, nb: M.Numbers | None = None) -> M.Numbers:
    nb = nb or M.Numbers()
    M.numbers_from_real_l(nb, M.Src("real_l_scan.json", scan))
    return nb


def _presidio(cid: str, *, kd: int = 1, nd: int = 7, kq: int = 10, nq: int = 33,
              one_call: int | None = None, sha: str | None = None,
              versions: dict | None = None) -> dict:
    """scripts/presidio_residual.py output for one collection, counts only."""
    def row(k: int, n: int) -> dict:
        return {"k": k, "k_one_call": k if one_call is None else one_call, "n": n,
                "present": n, "present_one_call": n - 1}
    return {"collection": cid, "generated_by": "scripts/presidio_residual.py",
            "configuration": {}, "manifest_sha256": "f" * 64,
            "identifiers_sha256": sha or hashlib.sha256(cid.encode()).hexdigest(),
            "documents": [{"name": f"doc{i}.md", "sha256": "0" * 64, "presidio_results": {}}
                          for i in range(1, 5)],
            "labels_excluded": 2, "rule": "...",
            "versions": versions or {"presidio-analyzer": "2.2.362",
                                     "presidio-anonymizer": "2.2.362", "en-core-web-lg": "3.8.0",
                                     "spacy": "3.8.7", "model_loaded": "en_core_web_lg 3.8.0",
                                     "python": "3.12.3"},
            "by_type": {"DIRECT": {"PERSON": row(kd, nd - 2), "CODE": row(0, 2),
                                   "total": row(kd, nd)},
                        "QUASI": {"LOC": row(kq, nq), "total": row(kq, nq)}}}


def _presidio_numbers(files: list[dict], scan: dict | None,
                      nb: M.Numbers | None = None) -> M.Numbers:
    nb = nb or M.Numbers()
    M.numbers_from_presidio(nb, [M.Src(f"presidio/presidio_{f['collection']}.json", f)
                                 for f in files],
                            M.Src("real_l_scan.json", scan) if scan else None)
    return nb


def test_track_l_keys_hold_with_the_rule_of_three_bound_and_the_paper_formats():
    nb = _l_numbers(_l_scan())
    assert nb["hl1_outcome"] == "holds" and nb["hl1_reason"] == "none"
    assert nb["hl1_computed"] == "holds" and nb["hl1_registered"] == "yes"
    assert nb["l_digest_check"] == "ok"
    assert nb["l_collections"] == "6" and nb["l_runs"] == "36"
    assert nb["l_runs_complete_kn"] == "36/36" and nb["l_runs_scanned_kn"] == "36/36"
    assert nb["l_direct_strings"] == "42" and nb["l_direct_crossed_kn"] == "0/42"
    assert nb["l_direct_bound"] == "0.071"                   # 3/42, 3 dp
    assert nb["l_direct_pairs"] == "0/120" and nb["l_direct_pairs_bound"] == "0.025"
    assert nb["l_direct_onecall_kn"] == "0/42"
    assert nb["l_direct_labels_crossed_kn"] == "0/0"
    assert nb["l_quasi_strings"] == "198" and nb["l_quasi_crossed_kn"] == "6/198"
    assert nb["l_quasi_crossed_pct"] == "3.0%"
    assert nb["l_quasi_type_LOC_kn"] == "4/80" and nb["l_quasi_type_DATETIME_kn"] == "2/118"
    assert nb["l_direct_type_PERSON_kn"] == "0/42"
    assert nb["l_cross_route_summary"] == ""
    assert nb["l_attempts_scanned"] == "0" and nb["l_attempts_direct_kn"] == "n/a"
    for arm in ("baseline", "r5", "r1"):
        assert nb[f"l_quasi_crossed_{arm}_kn"] == "2/198"
        assert nb[f"l_direct_crossed_{arm}_kn"] == "0/42"
        assert nb[f"l_g3_{arm}_net_kn"] == "1/240" and nb[f"l_g3_{arm}_gross_kn"] == "12/240"
        assert nb[f"l_g3_{arm}_collections"] == "6"
        assert nb[f"l_health_valid_{arm}_kn"] == "12/12"
        assert nb[f"l_sections_delivered_mean_{arm}"] == "3.0"
    assert nb["l_g3_control_kn"] == "6/240"
    assert nb["l_health_valid_kn"] == "36/36"
    assert nb["l_sections_delivered_mean"] == "3.0" and nb["l_sections_refused_mean"] == "0.2"
    assert nb["l_sections_delivered_total"] == "108"
    assert nb["l_sections_transport_failed_total"] == "0"
    assert nb["l_plans_refused_kn"] == "0/36"
    # 9 refused sections and no refused plan over 108 + 9 sections and 36 plans.
    assert nb["l_refused_attempts_kn"] == "9/153"
    assert nb["l_utility_strings_kn"] == "144/288" and nb["l_utility_strings_pct"] == "50.0%"
    assert nb["l_utility_ORG_pct"] == "50.0%" and nb["l_utility_strings_r1_pct"] == "50.0%"
    assert nb["l_extractions"] == "6"
    assert nb["l_recall_exact_kn"] == "120/240" and nb["l_recall_exact_pct"] == "50.0%"
    assert nb["l_recall_covered_pct"] == "60.0%" and nb["l_recall_exact_direct_pct"] == "71.4%"
    assert nb["l_echo_kn"] == "24/240" and nb["l_echo_pct"] == "10.0%"
    assert nb["l_labels_strings"] == "12" and nb["l_short_of_repeats"] == "none"
    assert nb["l_collection_list"].startswith("tab-L1, tab-L2")
    # Every key says where in the scan it came from.
    assert all(nb.sources[k]["source"].startswith("real_l_scan.json#/") for k in nb.values)
    # Nothing in the keys is a string: the values are counts, shares, words of the
    # scan's own fields or decisions.
    assert not any(re.search(r"\d{5}/\d{2}", v) for v in nb.values.values())


def test_track_l_fails_on_one_direct_crossing_with_the_routes_in_words():
    routes = [{"identifier_type": "DIRECT", "entity_type": "PERSON", "stage": "draft",
               "field": "messages", "strings": 1, "occurrences": 2}]
    att = {"scanned": ["tab-L3__r5__attempt1"], "direct": _lkn(1, 7),
           "routes": [{"identifier_type": "DIRECT", "entity_type": "CODE", "stage": "plan",
                       "field": "messages", "strings": 1, "occurrences": 1}]}
    nb = _l_numbers(_l_scan("fails", direct=_lkn(1, 42), pairs=_lkn(1, 120), routes=routes,
                            attempts=att))
    assert nb["hl1_outcome"] == "fails" and nb["hl1_reason"] == "none"
    assert nb["l_direct_crossed_kn"] == "1/42" and nb["l_direct_bound"] == "n/a"
    assert nb["l_direct_pairs"] == "1/120" and nb["l_direct_pairs_bound"] == "n/a"
    assert nb["l_cross_route_summary"] == (
        "DIRECT PERSON at the draft stage in the messages field: 1 string-run pair, "
        "2 occurrences; DIRECT CODE at the plan stage in the messages field of an archived "
        "attempt: 1 string-run pair, 1 occurrence")
    assert nb["l_attempts_scanned"] == "1" and nb["l_attempts_direct_kn"] == "1/7"


def test_track_l_is_not_testable_with_an_incomplete_run_and_says_why_in_the_scans_words():
    why = ("not scanned or incomplete: tab-L4__r1__rep2: trace incomplete, a forwarded "
           "request failed upstream and is not traced, or the run did not complete (trace_gap)")
    nb = _l_numbers(_l_scan("not testable", why, incomplete={"tab-L4__r1__rep2": "trace "
                                                              "incomplete"}))
    assert nb["hl1_outcome"] == "not testable" and nb["hl1_reason"] == why
    assert nb["l_runs_complete_kn"] == "35/36" and nb["l_runs_scanned_kn"] == "35/36"
    assert nb["l_runs_incomplete"] == "1" and nb["l_direct_bound"] == "0.071"
    # The outcome is the scan's field, never recomputed: a scan that says holds
    # over an incomplete run is reported as it says (real_l.py decides, with
    # its tests).
    assert _l_numbers(_l_scan("holds", incomplete={"x": "y"}))["hl1_outcome"] == "holds"


def test_track_l_is_not_testable_on_a_digest_mismatch_and_unrecorded_without_the_check():
    why = ("§12.8: the runs' SLM image and model digests and resolved cloud model are not all "
           "equal to the A-D batch's (slm_model_digest differs)")
    nb = _l_numbers(_l_scan("not testable", why, integrity={"checked": True,
                                                            "all_equal": False}))
    assert nb["hl1_outcome"] == "not testable" and nb["l_digest_check"] == "mismatch"
    assert nb["hl1_reason"] == why
    assert _l_numbers(_l_scan(integrity=None))["l_digest_check"] == "unrecorded"


def test_a_scan_over_other_collections_than_l1_l6_is_a_pilot_and_warns(capsys):
    nb = _l_numbers(_l_scan("not testable", "the tasks file is not the registered collections",
                            ids=["tab-L0"]))
    assert nb["hl1_outcome"] == "not testable" and nb["hl1_computed"] == "holds"
    assert nb["hl1_registered"] == "no" and nb["l_collections"] == "1"
    assert "not the registered" in capsys.readouterr().err


def test_track_l_blocks_the_scan_lacks_give_na_not_a_missing_key():
    arms = {a: _l_arm(a, with_report=False) for a in ("baseline", "r5")}
    arms["r1"] = {**_l_arm("r1", runs=0, valid=0),
                  "g3": {"collections": 0, "gross": None, "net": None}, "runs": []}
    nb = _l_numbers(_l_scan(arms=arms, short={"tab-L2": {"r1": 1}}))
    assert nb["l_g3_r1_net_kn"] == "n/a" and nb["l_utility_strings_baseline_pct"] == "n/a"
    assert nb["l_utility_strings_pct"] == "50.0%"       # r1's reports alone are pooled
    assert nb["l_short_of_repeats"] == "tab-L2 r1 1"
    assert nb["l_runs_r1"] == "0"


def test_a_file_that_is_not_the_scan_is_refused():
    with pytest.raises(SystemExit, match="not scripts/real_l.py scan output"):
        _l_numbers({**_l_scan(), "instrument": "scripts/real_l.py extraction"})
    with pytest.raises(SystemExit, match="no HL1 block"):
        _l_numbers({**_l_scan(), "HL1": {"outcome": "maybe"}})


def test_presidio_files_pool_per_string_and_single_call_readings_beside_each_other():
    scan = _l_scan()
    nb = _presidio_numbers([_presidio(cid, one_call=0) for cid in L_IDS], scan)
    assert nb["l_presidio_direct_kn"] == "6/42" and nb["l_presidio_quasi_kn"] == "60/198"
    assert nb["l_presidio_direct_onecall_kn"] == "0/42"
    assert nb["l_presidio_quasi_onecall_kn"] == "0/198"
    assert nb["l_presidio_direct_pct"] == "14.3%" and nb["l_presidio_quasi_pct"] == "30.3%"
    assert nb["l_presidio_direct_present_kn"] == "42/42"
    assert nb["l_presidio_direct_type_PERSON_kn"] == "6/30"
    assert nb["l_presidio_direct_type_CODE_kn"] == "0/12"
    assert nb["l_presidio_quasi_type_LOC_kn"] == "60/198"
    assert nb["l_presidio_version"] == "2.2.362"
    assert nb["l_presidio_model"] == "en_core_web_lg 3.8.0"
    assert nb["l_presidio_collections"] == "6" and nb["l_judgments"] == "24"
    assert nb["l_presidio_labels_excluded"] == "12"
    assert "pooled over 6 file(s)" in nb.sources["l_presidio_direct_kn"]["source"]
    # The free fallback (§12.9): Presidio alone, with no scan to check against.
    alone = _presidio_numbers([_presidio("tab-L1")], None)
    assert alone["l_presidio_direct_kn"] == "1/7" and alone["l_presidio_collections"] == "1"
    # Both instruments on one Numbers: no key is produced twice.
    both = _presidio_numbers([_presidio(cid) for cid in L_IDS], scan, _l_numbers(scan))
    assert both["l_direct_crossed_kn"] == "0/42" and both["l_presidio_direct_kn"] == "6/42"


@pytest.mark.parametrize(("files", "match"), [
    ([_presidio(cid) for cid in L_IDS[:5]], "must be over the same collections"),
    ([_presidio(cid, sha="a" * 64) if cid == "tab-L2" else _presidio(cid) for cid in L_IDS],
     "sha256 differs"),
    ([_presidio(cid, nd=8) if cid == "tab-L2" else _presidio(cid) for cid in L_IDS],
     "8 DIRECT strings scored, and the scan counts 7"),
    ([*(_presidio(cid) for cid in L_IDS), _presidio("tab-L1")], "is also"),
    ([_presidio(cid, versions={"presidio-analyzer": "2.2.360"}) if cid == "tab-L6"
      else _presidio(cid) for cid in L_IDS], "disagree on the versions"),
    ([{**_presidio(cid), "generated_by": "x"} for cid in L_IDS], "not scripts/presidio"),
])
def test_presidio_files_not_over_the_scans_strings_are_refused(files, match):
    with pytest.raises(SystemExit, match=match):
        _presidio_numbers(files, _l_scan())


def test_the_recorded_spend_is_read_from_the_budget_line_or_a_json_object(tmp_path):
    p = tmp_path / "budget.txt"
    p.write_text("recorded under eval/real/runs: private $6.130, adversary $1.500; spent "
                 "before it $0.650; cap $15.00; remaining $6.725; the next step needs at most "
                 "$0.000\n", encoding="utf-8")
    nb = M.Numbers()
    M.numbers_from_real_l_cost(nb, M.load_real_l_cost(p))
    assert nb["l_cost_usd"] == "7.63" and nb["l_cost_private_usd"] == "6.13"
    assert nb["l_cost_adversary_usd"] == "1.50" and nb["l_cost_before_usd"] == "0.65"
    assert nb["l_cost_total_usd"] == "8.28"
    assert nb.sources["l_cost_usd"]["source"] == "budget.txt#/private + /adversary (budget line)"
    j = tmp_path / "cost.json"
    j.write_text(json.dumps({"private": 6.13, "adversary": 1.5}), encoding="utf-8")
    nb = M.Numbers()
    M.numbers_from_real_l_cost(nb, M.load_real_l_cost(j))
    assert nb["l_cost_usd"] == "7.63" and nb["l_cost_before_usd"] == "0.00"
    bad = tmp_path / "bad.txt"
    bad.write_text("nothing here\n", encoding="utf-8")
    with pytest.raises(SystemExit, match="neither a JSON object"):
        M.load_real_l_cost(bad)
    with pytest.raises(SystemExit, match="missing input"):
        M.load_real_l_cost(tmp_path / "absent.txt")


def test_the_track_l_inputs_are_optional_and_reach_the_build(monkeypatch):
    got: list = []
    monkeypatch.setattr(M, "build", got.append)
    M.main([])
    assert got[-1].real_l is None and got[-1].presidio is None and got[-1].real_l_cost is None
    M.main(["--real-l", "s.json", "--presidio", "a.json", "--presidio", "b.json",
            "--real-l-cost", "c.txt"])
    a = got[-1]
    assert a.real_l == Path("s.json") and a.presidio == [Path("a.json"), Path("b.json")]
    assert a.real_l_cost == Path("c.txt")


def test_track_l_literals_are_bound_to_their_words():
    nums = {"l_direct_crossed_kn": "0/42", "l_direct_bound": "0.071", "l_quasi_crossed_pct":
            "3.0%", "l_runs": "36"}
    p = _para("A judgment is eligible with at least 4 annotators and at most 8,000 characters;"
              " a collection holds 4 judgments, and the cap is $15.")
    assert not C.check_paragraphs([p], nums).fails
    for bad in ("4 collections", "8,000 words", "15 runs", "8000 characters"):
        assert C.check_paragraphs([_para(f"It has {bad}.")], nums).fails
    # The value forms the track-L keys take pass only as filled values.
    p = _para("DIRECT strings crossed in 0/42 cases, bound 0.071, QUASI 3.0% over 36 runs.",
              ("l_direct_crossed_kn", "0/42"), ("l_direct_bound", "0.071"),
              ("l_quasi_crossed_pct", "3.0%"), ("l_runs", "36"))
    assert not C.check_paragraphs([p], nums).fails
    assert len(C.check_paragraphs([_para("DIRECT strings crossed in 0/42 cases, bound "
                                         "0.071, QUASI 3.0% over 36 runs.")], nums).fails) == 5


def _sim_arm(refused: bool, hits: int = 0, by_kind: dict | None = None,
             straddling: dict | None = None, rels_run: int = 0, max_tokens: int = 3) -> dict:
    return {"refused": refused, "reason_code": "RAW_CONTEXT_CONTAINMENT" if refused else None,
            "prompt": {"hits": hits, "by_field_kind": by_kind or {},
                       "straddling_field_kinds": straddling or {}},
            "relations": {"relations_with_run": rels_run, "max_field_tokens": max_tokens}}


def _sim_json() -> dict:
    ok = {"r5": _sim_arm(False), "baseline": _sim_arm(False), "r1": _sim_arm(False)}
    bad = {"r5": _sim_arm(True, 10, {"relations[].predicate": 9, "goal": 1},
                          {"relations[].object_type+relations[].predicate": 0}, 6, 7),
           "baseline": _sim_arm(True, 10, {"relations[].predicate": 10}),
           "r1": _sim_arm(True, 10, {"relations[].predicate": 10})}
    bad2 = {"r5": _sim_arm(True, 3, {}, {"relations[].object_type+relations[].predicate": 3},
                           2, 5),
            "baseline": _sim_arm(True, 3), "r1": _sim_arm(False)}
    return {"instrument": "scripts/real_l_refusal_sim.py",
            "collections": {"tab-L1": {"extractions": [{"arms": ok}]},
                            "tab-L2": {"extractions": [{"arms": bad}]},
                            "tab-L3": {"extractions": [{"arms": bad2}]},
                            "tab-L4": {"extractions": []}}}


def test_refusal_sim_keys_count_refused_collections_and_relation_windows():
    nb = M.Numbers()
    M.numbers_from_refusal_sim(nb, M.Src("refusal_sim.json", _sim_json()))
    assert nb["l_sim_collections"] == "3"
    assert nb["l_sim_refused_r5_kn"] == "2/3" and nb["l_sim_refused_kn"] == "2/3"
    assert nb["l_sim_refused_baseline_kn"] == "2/3" and nb["l_sim_refused_r1_kn"] == "1/3"
    # 10 + 3 windows over the two refused r5 collections; the goal window is not a relation one.
    assert nb["l_sim_windows"] == "13" and nb["l_sim_windows_relation_kn"] == "12/13"
    assert nb["l_sim_relations_with_run"] == "8" and nb["l_sim_max_field_tokens"] == "7"
    assert "arms/r5/prompt/hits" in nb.sources["l_sim_windows"]["source"]
    with pytest.raises(SystemExit, match="not scripts/real_l_refusal_sim.py"):
        M.numbers_from_refusal_sim(nb, M.Src("x.json", {"instrument": "other"}))


def test_real_l_attempt_keys_read_the_planner_refusal_notes(tmp_path: Path):
    att = tmp_path / "attempts"
    for cid, note in (("tab-L3", "planning refused by monitor: RAW_CONTEXT_CONTAINMENT: 1 hit"),
                      ("tab-L4", "planning refused by monitor: RAW_CONTEXT_CONTAINMENT: 1 hit"),
                      ("tab-L5", "plan: 8 sections")):
        d = att / f"{cid}__r5__attempt1"
        d.mkdir(parents=True)
        (d / "manifest.json").write_text(json.dumps({"notes": [note]}), encoding="utf-8")
    scan = M.Src("real_l_scan.json", {"HL1": {"registered_collections": [
        "tab-L1", "tab-L2", "tab-L3", "tab-L4", "tab-L5", "tab-L6"]}})
    nb = M.Numbers()
    M.numbers_from_real_l_attempts(nb, att, scan)
    assert nb["l_collections_planner_refused_kn"] == "2/6"
    assert nb["l_planner_refused_attempts"] == "2"
    assert nb["l_planner_refusal_code"] == "RAW_CONTEXT_CONTAINMENT"
    assert nb.sources["l_collections_planner_refused_kn"]["raw"]["collections"] == ["tab-L3",
                                                                                    "tab-L4"]
    nb2 = M.Numbers()
    M.numbers_from_real_l_attempts(nb2, att, None)
    assert nb2["l_collections_planner_refused_kn"] == "2/2"
    with pytest.raises(SystemExit, match="no attempts directory"):
        M.numbers_from_real_l_attempts(M.Numbers(), tmp_path / "none", None)


# ------------------------------------------------------------ result figures (2026-10-21)
# The monitor's refusals, what crosses, judged quality and track L are figures
# since 2026-10-21 (the author's request for an evaluation read from figures);
# they are drawn from the build's own values, parsed back from the strings the
# text prints.

RESULT_VALUES = {
    "fp_poscontrol_deployed": "33/33", "op_refused_attempts": "0/313", "op_reintegrator": "0",
    "leak_gate_drops_total": "1", "n_private_runs": "36",
    "maskdoc_refused_docs_A": "4/4", "maskdoc_refused_docs_D": "3/4; 4/4",
    "maskdoc_5gram_pct_A": "16.5%–18.4%", "maskdoc_surfaces_surviving_A": "4–5",
    "maskdoc_5gram_pct_D": "12.6%–15.3%", "maskdoc_surfaces_surviving_D": "2–5",
    "fp_deployed_reset": "122/500", "fp_deployed_reset_ci": "[0.208, 0.284]",
    "fp_deployed_reset_code_snippet": "22/100",
    "fp_deployed_reset_code_snippet_ci": "[0.150, 0.311]",
    "fp_deployed_session": "122/500",
    "fp_standalone_entropy": "103/500", "fp_standalone_entropy_ci": "[0.173, 0.244]",
    "pred_echo_pct_shared_A": "40.0%", "pred_control_echo_pct_shared_A": "4.3%",
    "judge_mean_exgr_oracle": "5.00", "judge_position_bias": "0.51 ± 0.24 [0.00–1.00]",
    "l_collections": "6", "l_collections_planner_refused_kn": "4/6",
    "l_direct_crossed_kn": "0/22", "l_quasi_crossed_kn": "1/340", "l_direct_bound": "0.136",
}
for _arm, _slots, _core, _g3 in (("baseline", "60.8 [39–73]", "0", "5/729"),
                                 ("r5", "47.8 [38–56]", "0", "13/729"),
                                 ("r1", "14.2 [11–17]", "3", "10/729")):
    RESULT_VALUES.update({
        f"obs_{_arm}_slots_adv": _slots, f"cross_a_{_arm}": "0", f"cross_a_label_{_arm}": "0",
        f"cross_a_scanned_{_arm}": "12/12", f"cross_b_core_{_arm}": _core,
        f"cross_b_scanned_{_arm}": "12/12", f"cross_b_cloud_{_arm}": "5",
        f"g3_entities_{_arm}": _g3, f"g3_5grams_{_arm}": "4/6342",
        f"judge_mean_exgr_{_arm}": "2.20", f"l_g3_{_arm}_net_kn": "2/362"})
for _c in ("goal_coverage", "coherence", "specificity"):
    RESULT_VALUES[f"judge_{_c}_private"] = "2.00"
    RESULT_VALUES[f"judge_{_c}_oracle"] = "5.00"


def test_result_values_are_parsed_only_in_the_forms_the_text_prints():
    assert M._kn("0/313") == (0, 313) and M._ci("[0.208, 0.284]") == (0.208, 0.284)
    assert M._pct_range("16.5%–18.4%") == (16.5, 18.4) and M._pct_range("4.3%") == (4.3, 4.3)
    assert M._ci("n/a") is None
    for bad in ("n/a", "0 of 313", "313"):
        with pytest.raises(ValueError):
            M._kn(bad)
    with pytest.raises(ValueError):
        M._pct_range("16.5–18.4")


def test_the_monitor_figure_puts_wanted_refusals_apart_from_false_positives():
    a, b = M.monitor_rows(RESULT_VALUES)
    wanted = [r[1] for r in a if r[0] != "head" and r[3] == "wanted"]
    assert wanted == ["fp_poscontrol_deployed", "maskdoc_refused_docs_A", "maskdoc_refused_docs_D"]
    # The pipeline's own runs carry no verdict (review of 2026-10-21).
    assert [r[1] for r in a if r[0] != "head" and r[3] == "unwanted"] == []
    assert [r[1] for r in a if r[0] != "head" and r[3] == "neutral"] == ["op_refused_attempts"]
    assert all(r[3] == "unwanted" for r in b if r[0] != "head")
    assert "kept 16.5%–18.4% of 5-grams, 4–5 constants" in [r[4] for r in a if r[0] != "head"]
    # The old table notes travel with the figure, with their values filled.
    na, nb_ = M.monitor_notes(RESULT_VALUES)
    assert "true by construction" in na and "the shared one first" in na
    assert "the all-strata rate is not a traffic rate" in nb_.replace("\n", " ")
    assert "122/500" in nb_ and "«" not in nb_


def test_crossing_cells_colour_only_what_the_core_sent():
    r1 = M.crossing_cells(RESULT_VALUES, "r1")
    assert r1["h3"] == ("0 and 0", "in 12/12 runs", M.RES_TINT_OK)
    assert r1["core"][2] == M.RES_TINT_HIT and r1["core"][0] == "3"
    # Cloud-written spans are C's own words: no good or bad colour.
    assert r1["cloud"][2] == M.RES_TINT_NEUTRAL
    assert M.crossing_cells(RESULT_VALUES, "baseline")["core"][2] == M.RES_TINT_OK


def test_the_artifact_result_figures_are_drawn_at_600_dpi_and_7_in(tmp_path):
    pytest.importorskip("matplotlib")
    for make, png, cap in ((M.make_fig_monitor, M.FIG_MONITOR_PNG, M.FIG_MONITOR_CAPTION),
                           (M.make_fig_crossings, M.FIG_CROSS_PNG, M.FIG_CROSS_CAPTION),
                           (M.make_fig_quality_real, M.FIG_QR_PNG, M.FIG_QR_CAPTION)):
        make(RESULT_VALUES, tmp_path)
        data = (tmp_path / png).read_bytes()
        w = int.from_bytes(data[16:20], "big")
        assert w == int(M.FIG3_WIDTH_IN * M.FIG_DPI), png
        caption = (tmp_path / cap).read_text(encoding="utf-8").strip()
        assert caption.startswith("Fig. ") and "«" not in caption
    # Since later on 2026-10-21 these three are the artifact's; the manuscript's one
    # data figure is the results figure.
    assert set(M.DATA_FIGURES) == {"fig3"}
    assert M.DATA_FIGURES["fig3"] == (M.FIG_RESULTS_PNG, M.FIG_RESULTS_CAPTION)


def test_the_results_figure_is_drawn_from_the_build_values_and_refuses_a_crossing(tmp_path):
    """Fig. 3 since 2026-10-29: what reaches the cloud, two panels, 600 dpi at the text
    width; (a) draws Cloze's constants as zero only where H3 holds and every arm's count is
    0; (b) needs every count as k/n and names the collection sets and the refusals."""
    pytest.importorskip("matplotlib")
    v = dict(RESULT_VALUES, h3_outcome="holds", l_presidio_direct_kn="28/65",
             l_presidio_quasi_kn="337/1001", l_direct_crossed_kn="0/22",
             l_quasi_crossed_kn="1/340", l_collections_planner_refused_kn="4/6",
             l_presidio_collections="6", l_collections_admitted="2")
    for L, kept in (("A", "5"), ("B", "6")):
        v[f"maskdoc_surfaces_surviving_shared_nolost_{L}"] = kept
    png, _ = M.make_fig_results(v, tmp_path)
    # At the text width (6.5 in since 2026-10-25), 600 dpi.
    from scripts import paper_format
    assert M.FIG_RESULTS_WIDTH_IN == paper_format.TEXT_WIDTH_IN
    assert int.from_bytes(png.read_bytes()[16:20], "big") == int(
        M.FIG_RESULTS_WIDTH_IN * M.FIG_DPI)
    assert (tmp_path / M.FIG_RESULTS_CAPTION).read_text(encoding="utf-8").strip() == \
        M.FIG_RESULTS_CAPTION_TEXT == ("Fig. 3. What reaches the cloud: Cloze against masking "
                                       "and redaction.")
    svg = (tmp_path / "fig3_results.svg").read_text(encoding="utf-8")
    assert "28/65" in svg and "0/22" in svg and "1/340" in svg and "337/1001" in svg
    assert "all 6 collections" in svg and "the 2 admitted ones" in svg and "4/6" in svg
    # Attribute accuracy and judged quality are no longer drawn.
    assert "attribute accuracy" not in svg and "judged quality" not in svg
    with pytest.raises(ValueError, match="H3 holds"):
        M.make_fig_results(dict(v, cross_a_r1="1"), tmp_path)
    with pytest.raises(ValueError, match="H3 holds"):
        M.make_fig_results(dict(v, h3_outcome="fails"), tmp_path)
    with pytest.raises(ValueError, match="k/n"):
        M.make_fig_results(dict(v, l_presidio_direct_kn="n/a"), tmp_path)


@needs_docx
def test_outcome_cells_are_shaded_by_their_least_favourable_outcome(tmp_path):
    assert M._outcome_tint("holds") == "DCE7F3"
    assert M._outcome_tint("A: not testable; B: holds") == "E8E8E4"
    assert M._outcome_tint("fails") == "F6D5D1" and M._outcome_tint("n/a") is None
    # Since 2026-10-26 (simulated review R4(5)): a post-pilot outcome is amber, apart from
    # H3's pre-pilot "holds" above; a less favourable outcome in the cell still wins.
    assert M._outcome_tint("holds (post-pilot floor rule)") == "F2E3C0"
    assert M._outcome_tint("holds (post-pilot; RBA-rule score)") == "F2E3C0"
    assert M._outcome_tint("not testable (post-pilot; RBA-rule score)") == "E8E8E4"
    src = tmp_path / "src.docx"
    docx.Document().save(str(src))
    spec = M.TableSpec("t2", "Table 2", "Cap.", ["H", "Outcome"], [["H3", "«alpha»"]],
                       shade_col=1)
    M.write_tables_docx(src, [spec], tmp_path / "tables.docx", {"alpha": "holds"})
    cell = docx.Document(str(tmp_path / "tables.docx")).tables[0].rows[1].cells[1]
    from docx.oxml.ns import qn
    shd = cell._tc.tcPr.find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "DCE7F3"
    assert cell.text == "holds" and all(r.bold for r in cell.paragraphs[0].runs if r.text)


@needs_docx
def test_a_merged_table_cell_is_as_wide_as_the_grid_columns_it_spans(tmp_path):
    """Final format check of 2026-10-28: a group-label or wide cell took the last
    column's width, which Word draws narrower than the table."""
    from docx.oxml.ns import qn
    from docx.shared import Inches
    src = tmp_path / "src.docx"
    docx.Document().save(str(src))
    spec = M.TableSpec("t2", "Table 2", "Cap.", ["A", "B", "C"],
                       [("group", "Label"), ["a", "b", "c"], ("wide", "x", "y")],
                       widths=[1.0, 2.0, 0.5])
    M.write_tables_docx(src, [spec], tmp_path / "t.docx", {})
    tbl = docx.Document(str(tmp_path / "t.docx")).tables[0]._tbl
    grid = [int(g.get(qn("w:w"))) for g in tbl.tblGrid.iter(qn("w:gridCol"))]
    spans = []
    for tr in tbl.tr_lst:
        col, row = 0, []
        for tc in tr.tc_lst:
            assert abs(tc.width.twips - sum(grid[col:col + tc.grid_span])) <= tc.grid_span
            row.append(tc.grid_span)
            col += tc.grid_span
        spans.append(row)
    assert spans[1:] == [[3], [1, 1, 1], [1, 2]]
    assert abs(tbl.tr_lst[1].tc_lst[0].width - Inches(sum(grid) / 1440)) < Inches(0.01)


def test_table1_is_the_extraction_and_r1s_store_per_corpus():
    """Table 1 since 2026-10-21: the per-corpus part of the old Table 2 without its
    echo column (Fig. 4(b)); it puts r1_store_cells_<L> as table3 does."""
    nb = _t3_numbers()
    nb.values["t4_A_r1_factsstruct"] = "18.0 [18.0–18.0]"
    spec = M.table_extraction(nb, {"A": "task-a"}, [])
    assert (spec.bid, spec.label, spec.caption) == (
        "t1", "Table 1", "Extraction and r1's store per corpus.")
    assert spec.header == list(M.T_EXTRACT_HEAD) and len(spec.header) == 7
    assert "Predicates echoing a short corpus n-gram (control)" not in spec.header
    assert nb["r1_store_cells_A"] == "18"
    assert spec.rows == [["A*", "«table_cells_A»", "«r1_store_cells_A»", "«t4_A_r1_pruned»",
                          "n/a", "n/a", "n/a"]]
    assert spec.notes[0].startswith("*Corpus A is in-sample") and "n/a: no probe" in spec.notes[0]


# Added by the AI assistant on 2026-10-27 with the --posthoc-l reader (scope §14,
# signed by the author that day: Presidio on the admitted collections, the r1
# drafter system field, the refused collections).

def _kn_block(k: int, n: int) -> dict:
    return {"k": k, "n": n, "rate": k / n if n else None}


def _tally_block(k_all, n_all, kd, nd, kq, nq) -> dict:
    return {"all": _kn_block(k_all, n_all), "DIRECT": {**_kn_block(kd, nd), "by_entity_type": {}},
            "QUASI": {**_kn_block(kq, nq), "by_entity_type": {}}}


def _posthoc_l(**over) -> dict:
    """A small synthetic posthoc_l.json in the shape scripts/posthoc_2026_10_27.py writes."""
    recall = _tally_block(4, 12, 1, 2, 3, 10)
    recall["covered"] = _tally_block(5, 12, 1, 2, 4, 10)
    ph = {
        "instrument": "scripts/posthoc_2026_10_27.py",
        "section_14_signed": True,
        "reproduced": {"l_direct_strings": "3", "l_quasi_strings": "20"},
        "inputs": {"real_l_scan.json": "0" * 64},
        "presidio_admitted": {"direct": _kn_block(1, 3), "quasi": _kn_block(6, 20),
                              "direct_CODE": _kn_block(1, 1), "direct_PERSON": _kn_block(0, 2)},
        "r1_drafter_system_field": {"same_as_A": _kn_block(7, 7)},
        "refused": {
            "code_commit": "26de996ff464a5e3d10bffc79ec91996bcd174e1",
            "recall": recall, "echo": _tally_block(0, 12, 0, 2, 0, 10), "rebuild_ok": True,
            "would_cross": _tally_block(1, 12, 0, 2, 1, 10),
            "would_cross_routes": [{"identifier_type": "QUASI", "entity_type": "DEM",
                                    "stage": "plan", "field": "system", "strings": 1,
                                    "occurrences": 1}]},
    }
    ph.update(over)
    return ph


def _posthoc_l_numbers(ph: dict, nb: M.Numbers | None = None) -> M.Numbers:
    if nb is None:
        nb = M.Numbers()
        nb.put("l_direct_strings", "3", "real_l_scan.json#/HL1/direct/n")
        nb.put("l_quasi_strings", "20", "real_l_scan.json#/HL1/quasi_descriptive/n")
    M.numbers_from_posthoc_l(nb, M.Src("posthoc_l.json", ph))
    return nb


def test_posthoc_l_analyses_become_keys_with_the_paper_formats():
    nb = _posthoc_l_numbers(_posthoc_l())
    assert nb["l_direct_strings"] == "3"
    assert nb["l_presidio_admitted_direct_kn"] == "1/3"
    assert nb["l_presidio_admitted_quasi_kn"] == "6/20"
    assert nb["l_presidio_admitted_direct_type_CODE_kn"] == "1/1"
    assert nb["l_presidio_admitted_direct_type_PERSON_kn"] == "0/2"
    assert nb["l_r1_sysfield_same_as_A_kn"] == "7/7"
    assert nb["l_refused_recall_exact_direct_kn"] == "1/2"
    assert nb["l_refused_recall_covered_quasi_kn"] == "4/10"
    assert nb["l_refused_echo_kn"] == "0/12" and nb["l_refused_echo_quasi_kn"] == "0/10"
    assert nb["l_refused_rebuild_ok"] == "yes"
    assert nb["l_refused_would_cross_direct_kn"] == "0/2"
    assert nb["l_refused_would_cross_quasi_kn"] == "1/10"
    assert nb["l_refused_would_cross_route_summary"] == (
        "QUASI DEM at the plan stage in the system field of a rebuilt r5 planner request: "
        "1 string-collection pair, 1 occurrence")
    new = [k for k in nb.values if k not in ("l_direct_strings", "l_quasi_strings")]
    assert all(nb.sources[k]["source"].startswith("posthoc_l.json#/") for k in new)


def test_posthoc_l_routes_read_none_when_nothing_would_cross_and_na_when_unrecorded():
    ph = _posthoc_l()
    ph["refused"]["would_cross_routes"] = []
    assert _posthoc_l_numbers(ph)["l_refused_would_cross_route_summary"] == "none"
    del ph["refused"]["would_cross_routes"]
    assert _posthoc_l_numbers(ph)["l_refused_would_cross_route_summary"] == "n/a"


def test_posthoc_l_a_rebuild_that_did_not_match_reports_no_would_be_crossing():
    ph = _posthoc_l()
    ph["refused"]["rebuild_ok"] = False
    nb = _posthoc_l_numbers(ph)
    assert nb["l_refused_rebuild_ok"] == "no"
    assert nb["l_refused_would_cross_direct_kn"] == "n/a"
    assert nb["l_refused_would_cross_route_summary"] == "n/a"
    # Recall and echo stay: they do not rest on the rebuild (scope §14.3's fallback).
    assert nb["l_refused_echo_kn"] == "0/12"


@pytest.mark.parametrize(("ph", "match"), [
    (_posthoc_l(instrument="scripts/posthoc_2026_10_26.py"), "not scripts/posthoc_2026_10_27.py"),
    (_posthoc_l(section_14_signed=False), "before scope §14 was signed"),
    (_posthoc_l(reproduced={}), "no reproduced frozen keys"),
    (_posthoc_l(reproduced={"l_direct_strings": "4"}), "l_direct_strings"),
    (_posthoc_l(refused=None), "not computed with commit"),
    (_posthoc_l(refused={"code_commit": "26de996ff464a5e3d10bffc79ec91996bcd174e1"}),
     "no refused/recall block"),
    (_posthoc_l(refused={"code_commit": "0" * 40}), "not computed with commit"),
])
def test_posthoc_l_analyses_that_misdescribe_themselves_are_refused(ph, match):
    with pytest.raises(SystemExit, match=match):
        _posthoc_l_numbers(ph)
